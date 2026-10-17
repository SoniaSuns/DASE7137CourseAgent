"""Read-only, per-request source inventory. Never edits the owner's originals."""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import stat
import zipfile

DEFAULT_ROOT = Path(__file__).resolve().parents[2] / 'files'
DEFAULT_INDEX = Path(__file__).resolve().parents[1] / 'course_files/library_index.json'
SOURCE_TEXT = {'.cs', '.py', '.js', '.ts', '.tsx', '.jsx', '.css', '.sql', '.java', '.c', '.cpp', '.h', '.xml', '.yaml', '.yml'}
READABLE = {'.docx', '.pptx', '.xlsx', '.xls', '.pdf', '.txt', '.md', '.csv', '.json', '.html', '.htm'} | SOURCE_TEXT
PAGE_SIZE = 16000


def sha(data):
    return hashlib.sha256(data).hexdigest()


class FileLibraryChanged(ValueError):
    """The same indexed source has new bytes; never reuse its old evidence."""
    def __init__(self,entry,actual_sha256):
        self.file_id=entry['id'];self.relative_path=entry['relative_path']
        self.expected_sha256=entry['sha256'];self.actual_sha256=actual_sha256
        super().__init__('Selected library file '+Path(self.relative_path).name+
                         ' changed during this request. A fresh source review is required.')


class FileLibrary:
    def __init__(self, root=DEFAULT_ROOT, index=DEFAULT_INDEX, *, detect_content=False):
        self.root = Path(root).absolute()
        self.index = Path(index).absolute()
        # Legacy callers retain suffix-dispatched results. Live attachment intake
        # opts into content detection for every source in its merged library.
        self.detect_content = detect_content

    def safe_path(self, relative):
        rel = Path(relative)
        if rel.is_absolute() or '..' in rel.parts or not rel.parts:
            raise ValueError('File source must be an exact library entry.')
        current = self.root
        for part in (None, *rel.parts):
            if part is not None:
                current = current / part
            info = current.lstat()
            if stat.S_ISLNK(info.st_mode) or getattr(info, 'st_file_attributes', 0) & 0x400:
                raise ValueError('Linked files and folders are not library sources.')
        if not current.resolve().is_relative_to(self.root.resolve()) or not current.is_file():
            raise ValueError('File source must remain inside the library.')
        return current

    def snapshot(self):
        if not self.root.is_dir():
            raise ValueError('The configured course file library is unavailable.')
        if self.root.is_symlink() or getattr(self.root.lstat(), 'st_file_attributes', 0) & 0x400:
            raise ValueError('The library root must not be a linked folder.')
        entries, skipped = [], []
        for folder, dirs, names in os.walk(self.root, followlinks=False):
            # Junctions must not introduce an outside subtree on Windows.
            for name in list(dirs):
                path = Path(folder) / name
                if path.is_symlink() or getattr(path.lstat(), 'st_file_attributes', 0) & 0x400:
                    dirs.remove(name)
                    skipped.append(path.relative_to(self.root).as_posix())
            for name in sorted(names):
                rel = (Path(folder) / name).relative_to(self.root).as_posix()
                try:
                    path = self.safe_path(rel)
                    # Temporary Office locks and dotfiles are not course materials.
                    if name.startswith(('~$', '.')):
                        skipped.append(rel)
                        continue
                    with path.open('rb') as source:
                        fingerprint = hashlib.file_digest(source, 'sha256').hexdigest()
                    entries.append({'id': 'file:' + sha(rel.encode())[:20], 'relative_path': rel,
                                    'format': path.suffix.lower(), 'bytes': path.stat().st_size,
                                    'sha256': fingerprint, 'text_reader': path.suffix.lower() in READABLE})
                except (OSError, ValueError):
                    skipped.append(rel)
        entries.sort(key=lambda entry: entry['relative_path'].casefold())
        result = {'root': str(self.root), 'generated_at': datetime.now(timezone.utc).isoformat(),
                  'semester': '2026-2027 Semester 1', 'files': entries, 'skipped': sorted(skipped)}
        result['snapshot_id'] = sha(json.dumps(entries, sort_keys=True).encode())
        self.index.parent.mkdir(parents=True, exist_ok=True)
        temp = self.index.with_suffix('.tmp')
        temp.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        temp.replace(self.index)
        lines = ['# Course file library', '', f'Source: `{self.root}`', '',
                 'Generated inventory; source files are unchanged. Read the JSON index before planning.',
                 'Filenames do not prove semester, class or approval. Inspect the selected contents.', '',
                 '| File | Format | Bytes |', '| --- | --- | --- |']
        lines.extend(f"| {e['relative_path'].replace('|', '&#124;')} | {e['format']} | {e['bytes']} |" for e in entries)
        self.index.with_suffix('.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
        return result

    def entry(self, snapshot, file_id):
        if snapshot.get('root') != str(self.root):
            raise ValueError('File snapshot belongs to another library.')
        matches = [e for e in snapshot['files'] if e['id'] == file_id]
        if len(matches) != 1:
            raise ValueError('Select an exact file ID from the current library index.')
        return matches[0]

    def read_bytes(self, snapshot, file_id):
        entry = self.entry(snapshot, file_id)
        path = self.safe_path(entry['relative_path'])
        if path.stat().st_size > 128 * 1024 * 1024:
            raise ValueError('Selected file exceeds the bounded reader size; use a smaller source.')
        data = path.read_bytes()
        actual=sha(data)
        if actual != entry['sha256']:
            raise FileLibraryChanged(entry,actual)
        return entry, data

    def validate_reads(self, snapshot, reads):
        for file_id in {r['file_id'] for r in reads}:
            self.read_bytes(snapshot, file_id)

    def inspect(self, snapshot, file_id, offset=0):
        if type(offset) is not int or offset < 0:
            raise ValueError('File inspection offset must be a non-negative integer.')
        entry, data = self.read_bytes(snapshot, file_id)
        result = {'file_id': file_id, 'relative_path': entry['relative_path'],
                  'sha256': entry['sha256'], 'offset': offset, 'text': '', 'next_offset': None}
        try:
            suffix=entry['format']
            if self.detect_content:
                visual=image_metadata(data)
                if visual:
                    result.update(status='image',visual=visual,
                                  text='Image supplied as visual input. Describe only visible content; do not infer unreadable text. '+
                                       ('Only the first frame is supplied.' if visual['frames']>1 else ''))
                    return result
                suffix=content_format(data,suffix)
                result['detected_format']=suffix
            # Only enough text to supply one bounded page; no OCR or formatting claims.
            text = ''
            for chunk in text_chunks(data, suffix):
                text += chunk + '\n'
                if len(text) > offset + PAGE_SIZE:
                    break
            result['text'] = text[offset:offset + PAGE_SIZE]
            if len(text) > offset + PAGE_SIZE:
                result['next_offset'] = offset + PAGE_SIZE
            result['status'] = 'read' if result['text'].strip() else 'no_extractable_text'
            if suffix=='.docx' and offset==0:
                from docx import Document
                doc=Document(io.BytesIO(data));rows=[];size=0
                for ti,table in enumerate(doc.tables):
                    for ri,row in enumerate(table.rows):
                        cells=[cell.text for cell in row.cells]
                        size+=sum(len(c) for c in cells)
                        if len(rows)>=300 or size>32000:break
                        rows.append({'row_id':f't{ti+1}:r{ri+1}','cells':cells})
                    if len(rows)>=300 or size>32000:break
                result['table_rows']=rows
        except Exception as exc:
            result.update(status='reader_unavailable', reason=type(exc).__name__ + ': cannot extract this format; do not infer its contents.')
        return result


def image_metadata(data):
    from PIL import Image, UnidentifiedImageError
    try:
        with Image.open(io.BytesIO(data)) as im:
            if im.width*im.height>25_000_000:raise ValueError('Image exceeds the 25 megapixel decoding limit.')
            result={'format':im.format,'width':im.width,'height':im.height,'frames':getattr(im,'n_frames',1)}
            im.verify()
            return result
    except UnidentifiedImageError:
        return None


def decode_text(data):
    if data.startswith((b'\xff\xfe\x00\x00',b'\x00\x00\xfe\xff')):encoding='utf-32'
    elif data.startswith((b'\xff\xfe',b'\xfe\xff')):encoding='utf-16'
    else:encoding='utf-8-sig'
    text=data.decode(encoding)
    if any(ord(c)<32 and c not in '\n\r\t\f' for c in text):
        raise ValueError('Binary source requires a specialized reader.')
    return text


def content_format(data, suffix):
    if data.lstrip().startswith(b'%PDF-'):return '.pdf'
    if zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as package:
            names=set(package.namelist())
            for name,kind in [('word/document.xml','.docx'),('ppt/presentation.xml','.pptx'),('xl/workbook.xml','.xlsx')]:
                if name in names:return kind
        return '.zip'
    if data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):return suffix
    try:text=decode_text(data)
    except (UnicodeError,ValueError):return suffix
    if text.lstrip().lower().startswith(('<!doctype html','<html','<table')):return '.html'
    return '.txt'


def text_chunks(data, suffix):
    stream = io.BytesIO(data)
    if suffix in {'.docx', '.pptx', '.xlsx'}:
        with zipfile.ZipFile(stream) as package:
            if sum(e.file_size for e in package.infolist()) > 512 * 1024 * 1024:
                raise ValueError('Expanded Office document exceeds reader limit.')
        stream.seek(0)
    if suffix == '.docx':
        from docx import Document
        doc = Document(stream)
        for paragraph in doc.paragraphs:
            yield paragraph.text
        for i, table in enumerate(doc.tables, 1):
            yield f'[Table {i}]'
            for row in table.rows:
                yield ' | '.join(cell.text for cell in row.cells)
    elif suffix == '.pptx':
        from pptx import Presentation
        for i, slide in enumerate(Presentation(stream).slides, 1):
            yield f'[Slide {i}]'
            for shape in slide.shapes:
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield ' | '.join(cell.text for cell in row.cells)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                yield f'[Sheet {sheet.title}]'
                cached_rows=cached[sheet.title].iter_rows(values_only=True)
                for index,row in enumerate(sheet.iter_rows(values_only=True),1):
                    saved=next(cached_rows,())
                    yield ' | '.join('' if v is None else str(v) for v in row)
                    formulas=[]
                    for column,value in enumerate(row):
                        if isinstance(value,str) and value.startswith('='):
                            from openpyxl.utils import get_column_letter
                            result=saved[column] if column<len(saved) else None
                            formulas.append(f'{get_column_letter(column+1)}{index}='+str(result))
                    if formulas:yield '[Saved formula results; not recalculated] '+'; '.join(formulas)
        finally:
            workbook.close()
            cached.close()
    elif suffix in {'.html', '.htm'}:
        # Parse local bytes only: never execute scripts, load images or follow links.
        from lxml import html
        try:source=data.decode('utf-8-sig')
        except UnicodeDecodeError:source=data
        document = html.fromstring(source, parser=html.HTMLParser(no_network=True))
        for node in document.xpath('//script|//style|//template|//noscript'):
            node.drop_tree()
        title=document.xpath('string(//title)')
        if title:yield '[Title] '+title.strip()
        for node in document.iter():
            if node.tag=='tr':
                yield ' | '.join(' '.join(cell.itertext()).strip() for cell in node if cell.tag in {'th','td'})
            elif node.tag in {'h1','h2','h3','h4','p','li','pre'} and not any(p.tag in {'td','th'} for p in node.iterancestors()):
                yield ' '.join(node.itertext()).strip()
    elif suffix == '.xls':
        import xlrd
        book = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            for sheet in book.sheets():
                yield f'[Sheet {sheet.name}]'
                for i in range(sheet.nrows):
                    yield ' | '.join(str(v) for v in sheet.row_values(i))
        finally:
            book.release_resources()
    elif suffix == '.pdf':
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(stream).pages, 1):
            content = page.extract_text() or ''
            if content.strip():
                yield f'[Page {i}]\n' + content
    elif suffix == '.zip':
        # Read members in memory, never extract paths or execute packaged code.
        with zipfile.ZipFile(stream) as package:
            entries=package.infolist()
            if len(entries)>512 or sum(e.file_size for e in entries)>32*1024*1024:
                raise ValueError('Archive exceeds bounded inspection limits.')
            for entry in entries:
                if entry.is_dir():continue
                yield '[Archive member] '+entry.filename
                if entry.flag_bits & 1:
                    yield '[Encrypted member; password is required]';continue
                member=package.read(entry)
                kind=content_format(member,Path(entry.filename).suffix.lower())
                if kind=='.zip':
                    yield '[Nested archive; further inspection required]';continue
                try:
                    yield from text_chunks(member,kind)
                except Exception:
                    yield '[Member has no extracted text; further inspection required]'
    elif suffix in {'.txt', '.md', '.csv', '.json'} | SOURCE_TEXT:
        yield decode_text(data)
    else:
        raise ValueError('No text reader for this format.')


if __name__ == '__main__':
    config_path = Path(__file__).with_name('chrome_operator_config.json')
    config = json.loads(config_path.read_text(encoding='utf-8'))
    library = FileLibrary(config.get('file_library_root', DEFAULT_ROOT))
    snapshot = library.snapshot()
    print(f"Indexed {len(snapshot['files'])} files; skipped {len(snapshot['skipped'])}. Index: {library.index}")
