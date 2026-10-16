from __future__ import annotations

import hashlib
import json
import os
import random
import re
import shutil
import tempfile
from copy import copy, deepcopy
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any


SUPPORTED_SUFFIXES = {".docx", ".xlsx", ".pptx"}
CHANGE_TYPES = {
    "anchored_text_edit",
    "scalar_fact_update",
    "repeatable_record_rebuild",
}
ORDERING_POLICIES = {
    "preserve_existing_then_append_new",
    "seeded_random",
    "explicit_order",
}
OVERFLOW_POLICIES = {"block_on_capacity_overflow"}
RULE_POLICIES = {"preserve_existing_rules"}
ACCEPTED_CONFIDENCE = {"high", "medium"}
FORBIDDEN_PLAN_KEYS = {
    "command",
    "shell",
    "python",
    "python_code",
    "javascript",
    "javascript_code",
    "raw_code",
    "script",
    "script_text",
    "xpath",
    "xml",
    "local_path",
    "file_path",
}


class DocumentChangeBlocked(RuntimeError):
    pass


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _canonical(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=_json_value,
    ).encode("utf-8")


def fingerprint(value: Any) -> str:
    return hashlib.sha256(_canonical(value)).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _with_fingerprint(item: dict[str, Any]) -> dict[str, Any]:
    normalized = dict(item)
    normalized["fingerprint"] = fingerprint(normalized)
    return normalized


def _run_style(run: Any) -> dict[str, Any]:
    font = getattr(run, "font", None)
    try:
        color = getattr(getattr(font, "color", None), "rgb", None)
    except (AttributeError, ValueError):
        color = None
    try:
        size_pt = float(font.size.pt) if getattr(font, "size", None) is not None else None
    except (AttributeError, TypeError, ValueError):
        size_pt = None
    return {
        "text": str(getattr(run, "text", "") or ""),
        "bold": getattr(run, "bold", getattr(font, "bold", None)),
        "italic": getattr(run, "italic", getattr(font, "italic", None)),
        "underline": str(getattr(run, "underline", getattr(font, "underline", None)) or ""),
        "font_name": str(getattr(font, "name", "") or ""),
        "font_size_pt": size_pt,
        "font_color": str(color or ""),
    }


def _docx_cell_structure(cell: Any) -> dict[str, Any]:
    properties = cell._tc.tcPr
    grid_span = getattr(getattr(properties, "gridSpan", None), "val", None)
    vertical_merge = getattr(getattr(properties, "vMerge", None), "val", None)
    return {
        "style": str(getattr(getattr(cell, "style", None), "name", "") or ""),
        "vertical_alignment": str(getattr(cell, "vertical_alignment", None) or ""),
        "width_emu": int(cell.width) if cell.width is not None else None,
        "grid_span": int(grid_span) if grid_span is not None else 1,
        "vertical_merge": str(vertical_merge or ""),
        "paragraph_styles": [
            {
                "style": str(getattr(paragraph.style, "name", "") or ""),
                "runs": [_run_style(run) for run in paragraph.runs],
            }
            for paragraph in cell.paragraphs
        ],
    }


def _docx_model(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    structures: list[dict[str, Any]] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        structures.append(
            _with_fingerprint(
                {
                    "id": f"docx:p{paragraph_index:04d}",
                    "kind": "paragraph",
                    "text": paragraph.text,
                    "style": str(getattr(paragraph.style, "name", "") or ""),
                    "position": {"paragraph_index": paragraph_index - 1},
                    "runs": [_run_style(run) for run in paragraph.runs],
                }
            )
        )
    for table_index, table in enumerate(document.tables, start=1):
        table_id = f"docx:t{table_index:04d}"
        rows: list[list[str]] = []
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text for cell in row.cells]
            rows.append(values)
            for column_index, value in enumerate(values, start=1):
                structures.append(
                    _with_fingerprint(
                        {
                            "id": f"{table_id}:r{row_index:04d}:c{column_index:04d}",
                            "kind": "table_cell",
                            "parent_id": table_id,
                            "row_index": row_index - 1,
                            "column_index": column_index - 1,
                            "text": value,
                            "display_value": value,
                            "position": {
                                "table_index": table_index - 1,
                                "row_index": row_index - 1,
                                "column_index": column_index - 1,
                            },
                            **_docx_cell_structure(row.cells[column_index - 1]),
                        }
                    )
                )
        structures.append(
            _with_fingerprint(
                {
                    "id": table_id,
                    "kind": "table",
                    "row_count": len(rows),
                    "column_count": max((len(row) for row in rows), default=0),
                    "rows": rows,
                    "position": {"table_index": table_index - 1},
                }
            )
        )
    return {"format": "docx", "structures": structures}


def _xlsx_model(path: Path) -> dict[str, Any]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False)
    display_workbook = openpyxl.load_workbook(path, data_only=True)
    structures: list[dict[str, Any]] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        sheet_id = f"xlsx:s{sheet_index:04d}"
        min_row, min_column, max_row, max_column = 1, 1, worksheet.max_row, worksheet.max_column
        rows: list[list[Any]] = []
        display_rows: list[list[Any]] = []
        display_worksheet = display_workbook.worksheets[sheet_index - 1]
        for row_index in range(min_row, max_row + 1):
            values: list[Any] = []
            display_values: list[Any] = []
            for column_index in range(min_column, max_column + 1):
                cell = worksheet.cell(row_index, column_index)
                value = _json_value(cell.value)
                display_value = _json_value(display_worksheet.cell(row_index, column_index).value)
                values.append(value)
                display_values.append(display_value)
                if cell.value not in (None, ""):
                    structures.append(
                        _with_fingerprint(
                            {
                                "id": f"{sheet_id}:c{cell.coordinate}",
                                "kind": "worksheet_cell",
                                "parent_id": sheet_id,
                                "row_index": row_index - 1,
                                "column_index": column_index - 1,
                                "value": value,
                                "display_value": display_value,
                                "number_format": cell.number_format,
                                "is_formula": isinstance(cell.value, str) and cell.value.startswith("="),
                                "formula": cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else "",
                                "style_id": int(cell.style_id),
                                "style_name": str(cell.style or ""),
                                "position": {
                                    "sheet_index": sheet_index - 1,
                                    "row_index": row_index - 1,
                                    "column_index": column_index - 1,
                                    "coordinate": cell.coordinate,
                                },
                            }
                        )
                    )
            rows.append(values)
            display_rows.append(display_values)
        structures.append(
            _with_fingerprint(
                {
                    "id": sheet_id,
                    "kind": "table",
                    "sheet_title": worksheet.title,
                    "range": {
                        "min_row": min_row,
                        "min_column": min_column,
                        "max_row": max_row,
                        "max_column": max_column,
                    },
                    "row_count": len(rows),
                    "column_count": max_column,
                    "rows": rows,
                    "display_rows": display_rows,
                    "merged_ranges": sorted(str(item) for item in worksheet.merged_cells.ranges),
                    "position": {"sheet_index": sheet_index - 1},
                }
            )
        )
    return {"format": "xlsx", "structures": structures}


def _pptx_model(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    presentation = Presentation(path)
    structures: list[dict[str, Any]] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        slide_id = f"pptx:s{slide_index:04d}"
        structures.append(
            _with_fingerprint(
                {
                    "id": slide_id,
                    "kind": "slide",
                    "position": {"slide_index": slide_index - 1},
                    "shape_count": len(slide.shapes),
                }
            )
        )
        for shape_index, shape in enumerate(slide.shapes, start=1):
            shape_id = f"{slide_id}:sh{shape_index:04d}"
            if bool(getattr(shape, "has_table", False)):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                table_id = f"{shape_id}:t0001"
                for row_index, row in enumerate(rows, start=1):
                    for column_index, value in enumerate(row, start=1):
                        structures.append(
                            _with_fingerprint(
                                {
                                    "id": f"{table_id}:r{row_index:04d}:c{column_index:04d}",
                                    "kind": "table_cell",
                                    "parent_id": table_id,
                                    "row_index": row_index - 1,
                                    "column_index": column_index - 1,
                                    "text": value,
                                    "display_value": value,
                                    "position": {
                                        "slide_index": slide_index - 1,
                                        "shape_index": shape_index - 1,
                                        "row_index": row_index - 1,
                                        "column_index": column_index - 1,
                                    },
                                    "is_merge_origin": bool(getattr(shape.table.cell(row_index - 1, column_index - 1), "is_merge_origin", False)),
                                    "is_spanned": bool(getattr(shape.table.cell(row_index - 1, column_index - 1), "is_spanned", False)),
                                }
                            )
                        )
                structures.append(
                    _with_fingerprint(
                        {
                            "id": table_id,
                            "kind": "table",
                            "slide_index": slide_index - 1,
                            "shape_index": shape_index - 1,
                            "row_count": len(rows),
                            "column_count": max((len(row) for row in rows), default=0),
                            "rows": rows,
                            "merged_cells": [
                                {
                                    "row_index": row_index,
                                    "column_index": column_index,
                                    "is_merge_origin": bool(getattr(cell, "is_merge_origin", False)),
                                    "is_spanned": bool(getattr(cell, "is_spanned", False)),
                                }
                                for row_index, row in enumerate(shape.table.rows)
                                for column_index, cell in enumerate(row.cells)
                                if bool(getattr(cell, "is_merge_origin", False))
                                or bool(getattr(cell, "is_spanned", False))
                            ],
                            "position": {
                                "slide_index": slide_index - 1,
                                "shape_index": shape_index - 1,
                                "left_emu": int(shape.left),
                                "top_emu": int(shape.top),
                                "width_emu": int(shape.width),
                                "height_emu": int(shape.height),
                            },
                        }
                    )
                )
            elif bool(getattr(shape, "has_text_frame", False)):
                structures.append(
                    _with_fingerprint(
                        {
                            "id": shape_id,
                            "kind": "text_box",
                            "slide_index": slide_index - 1,
                            "shape_index": shape_index - 1,
                            "text": shape.text,
                            "display_value": shape.text,
                            "position": {
                                "slide_index": slide_index - 1,
                                "shape_index": shape_index - 1,
                                "left_emu": int(shape.left),
                                "top_emu": int(shape.top),
                                "width_emu": int(shape.width),
                                "height_emu": int(shape.height),
                            },
                            "paragraphs": [
                                {
                                    "level": int(paragraph.level),
                                    "runs": [_run_style(run) for run in paragraph.runs],
                                }
                                for paragraph in shape.text_frame.paragraphs
                            ],
                        }
                    )
                )
            else:
                structures.append(
                    _with_fingerprint(
                        {
                            "id": shape_id,
                            "kind": "unsupported_shape",
                            "shape_type": str(getattr(shape, "shape_type", "") or ""),
                            "has_chart": bool(getattr(shape, "has_chart", False)),
                            "position": {
                                "slide_index": slide_index - 1,
                                "shape_index": shape_index - 1,
                                "left_emu": int(shape.left),
                                "top_emu": int(shape.top),
                                "width_emu": int(shape.width),
                                "height_emu": int(shape.height),
                            },
                            "execution_support": "preserve_or_block",
                        }
                    )
                )
    return {"format": "pptx", "structures": structures}


def build_document_model(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentChangeBlocked(f"unsupported editable Office format: {suffix or '(none)'}")
    if suffix == ".docx":
        model = _docx_model(source)
    elif suffix == ".xlsx":
        model = _xlsx_model(source)
    else:
        model = _pptx_model(source)
    model["source_name"] = source.name
    model["source_sha256"] = file_sha256(source)
    model["document_fingerprint"] = fingerprint(
        {
            "format": model.get("format"),
            "source_sha256": model["source_sha256"],
            "structures": model.get("structures") or [],
        }
    )
    return model


def planner_safe_document_model(
    model: dict[str, Any],
    *,
    max_structures: int = 1000,
    max_chars: int = 280_000,
) -> dict[str, Any]:
    def compact_run(run: Any) -> dict[str, Any]:
        if not isinstance(run, dict):
            return {}
        return {key: value for key, value in run.items() if key != "text"}

    structures = []
    for item in model.get("structures") or []:
        if not isinstance(item, dict):
            continue
        copy_item = dict(item)
        rows = copy_item.get("rows")
        if isinstance(rows, list) and len(rows) > 80:
            copy_item["rows"] = rows[:80]
            copy_item["rows_truncated"] = len(rows) - 80
        display_rows = copy_item.get("display_rows")
        if isinstance(display_rows, list) and len(display_rows) > 80:
            copy_item["display_rows"] = display_rows[:80]
        for text_key in ("text", "display_value", "value"):
            value = copy_item.get(text_key)
            if isinstance(value, str) and len(value) > 8_000:
                copy_item[text_key] = value[:8_000]
                copy_item[f"{text_key}_truncated"] = len(value) - 8_000
        runs = copy_item.get("runs")
        if isinstance(runs, list):
            copy_item["run_count"] = len(runs)
            copy_item["runs"] = [compact_run(run) for run in runs[:8]]
        paragraph_styles = copy_item.get("paragraph_styles")
        if isinstance(paragraph_styles, list):
            copy_item["paragraph_styles"] = [
                {
                    **paragraph,
                    "runs": [compact_run(run) for run in (paragraph.get("runs") or [])[:4]],
                    "run_count": len(paragraph.get("runs") or []),
                }
                for paragraph in paragraph_styles[:8]
                if isinstance(paragraph, dict)
            ]
        paragraphs = copy_item.get("paragraphs")
        if isinstance(paragraphs, list):
            copy_item["paragraphs"] = [
                {
                    **paragraph,
                    "runs": [compact_run(run) for run in (paragraph.get("runs") or [])[:4]],
                    "run_count": len(paragraph.get("runs") or []),
                }
                for paragraph in paragraphs[:8]
                if isinstance(paragraph, dict)
            ]
        projected = {
            "format": model.get("format"),
            "source_name": model.get("source_name"),
            "source_sha256": model.get("source_sha256"),
            "document_fingerprint": model.get("document_fingerprint"),
            "structures": [*structures, copy_item],
        }
        if structures and len(json.dumps(projected, ensure_ascii=False, default=_json_value)) > max_chars:
            break
        structures.append(copy_item)
        if len(structures) >= max_structures:
            break
    return {
        "format": model.get("format"),
        "source_name": model.get("source_name"),
        "source_sha256": model.get("source_sha256"),
        "document_fingerprint": model.get("document_fingerprint"),
        "structures": structures,
        "structures_truncated": max(0, len(model.get("structures") or []) - len(structures)),
    }


def structure_index(model: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {
        str(item.get("id")): item
        for item in model.get("structures") or []
        if isinstance(item, dict) and str(item.get("id") or "")
    }


def _forbidden_plan_paths(value: Any, prefix: str = "") -> list[str]:
    found: list[str] = []
    if isinstance(value, dict):
        for key, nested in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            if str(key).casefold() in FORBIDDEN_PLAN_KEYS:
                found.append(path)
            found.extend(_forbidden_plan_paths(nested, path))
    elif isinstance(value, list):
        for index, nested in enumerate(value):
            found.extend(_forbidden_plan_paths(nested, f"{prefix}[{index}]"))
    return found


def validate_document_change_plan(plan: dict[str, Any], model: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    forbidden = _forbidden_plan_paths(plan)
    if forbidden:
        errors.append("plan contains forbidden executable/location fields: " + ", ".join(forbidden[:10]))
    if str(plan.get("source_document_fingerprint") or "") != str(model.get("document_fingerprint") or ""):
        errors.append("plan source_document_fingerprint does not match the current source")
    targets = structure_index(model)
    operations = plan.get("operations")
    if not isinstance(operations, list):
        errors.append("plan.operations must be a list")
        return errors
    for index, operation in enumerate(operations, start=1):
        label = f"operations[{index}]"
        if not isinstance(operation, dict):
            errors.append(f"{label} must be an object")
            continue
        change_type = str(operation.get("type") or "")
        if change_type not in CHANGE_TYPES:
            errors.append(f"{label}.type is unsupported: {change_type!r}")
        target_id = str(operation.get("target_id") or "")
        target = targets.get(target_id)
        if target is None:
            errors.append(f"{label}.target_id does not resolve exactly: {target_id!r}")
        elif str(operation.get("target_fingerprint") or "") != str(target.get("fingerprint") or ""):
            errors.append(f"{label}.target_fingerprint does not match the current structure")
        evidence = operation.get("evidence")
        if not isinstance(evidence, list) or not evidence or not all(isinstance(item, dict) for item in evidence):
            errors.append(f"{label}.evidence must be a non-empty list of provenance objects")
        if str(operation.get("confidence") or "").lower() not in ACCEPTED_CONFIDENCE:
            errors.append(f"{label}.confidence must be high or medium")
        if not isinstance(operation.get("readback"), dict):
            errors.append(f"{label}.readback must be an object")
        if change_type in {"anchored_text_edit", "scalar_fact_update"}:
            before_value = operation.get("expected_before")
            after_value = operation.get("replacement_text")
            before = "" if before_value is None else str(before_value)
            after = "" if after_value is None else str(after_value)
            if not before or "replacement_text" not in operation or after_value is None or before == after:
                errors.append(f"{label} requires distinct expected_before and replacement_text")
            if target is not None:
                if target.get("kind") not in {"paragraph", "text_box", "table_cell", "worksheet_cell"}:
                    errors.append(f"{label} target is not a supported text-bearing structure")
                current = str(target.get("text") if "text" in target else target.get("value") or "")
                if before not in current:
                    errors.append(f"{label}.expected_before is not present in the target")
        if change_type == "repeatable_record_rebuild":
            if target is not None and target.get("kind") != "table":
                errors.append(f"{label} target must be a table")
            roles = operation.get("column_roles")
            if not isinstance(roles, dict) or not roles:
                errors.append(f"{label}.column_roles must be a non-empty object")
            else:
                try:
                    role_columns = [int(value) for value in roles.values()]
                except (TypeError, ValueError):
                    errors.append(f"{label}.column_roles values must be zero-based integers")
                    role_columns = []
                if role_columns and len(set(role_columns)) != len(role_columns):
                    errors.append(f"{label}.column_roles must use unique columns")
                column_count = int((target or {}).get("column_count") or 0)
                if any(value < 0 or value >= column_count for value in role_columns):
                    errors.append(f"{label}.column_roles contains a column outside the target table")
            if not isinstance(operation.get("record_source"), dict):
                errors.append(f"{label}.record_source must be an object")
            policy = str(operation.get("ordering_policy") or "")
            if policy not in ORDERING_POLICIES:
                errors.append(f"{label}.ordering_policy is unsupported: {policy!r}")
            if policy == "seeded_random":
                if str(operation.get("randomization_evidence_ref") or "") != "current_prompt":
                    errors.append(f"{label}.seeded_random requires randomization_evidence_ref=current_prompt")
                if not str(operation.get("explicit_prompt_basis") or "").strip():
                    errors.append(f"{label}.seeded_random requires explicit_prompt_basis")
            rule_policy = str(operation.get("rule_policy") or "")
            if rule_policy and rule_policy not in RULE_POLICIES:
                errors.append(f"{label}.rule_policy is unsupported: {rule_policy!r}")
            if rule_policy == "preserve_existing_rules":
                rule_refs = [str(value) for value in operation.get("preserved_rule_refs") or []]
                evidence_refs = {
                    str(item.get("source_ref") or "")
                    for item in operation.get("evidence") or []
                    if isinstance(item, dict)
                }
                if not rule_refs or any(value not in targets for value in rule_refs):
                    errors.append(f"{label}.preserved_rule_refs must contain exact current structure IDs")
                elif any(value not in evidence_refs for value in rule_refs):
                    errors.append(f"{label} must cite every preserved rule structure in evidence")
            if str((operation.get("schedule") or {}).get("overflow_policy") or "block_on_capacity_overflow") not in OVERFLOW_POLICIES:
                errors.append(f"{label}.schedule.overflow_policy is unsupported")
    return errors


def _group_snapshot_hash(snapshot: dict[str, Any]) -> str:
    canonical = {
        "course_id": str(snapshot.get("course_id") or ""),
        "groups": sorted(
            [
                {
                    "name": str(item.get("name") or ""),
                    "id": str(item.get("id") or ""),
                    "grouping_names": sorted(str(value) for value in item.get("grouping_names") or []),
                }
                for item in snapshot.get("groups") or []
                if isinstance(item, dict) and str(item.get("name") or "")
            ],
            key=lambda item: (item["name"].casefold(), item["id"]),
        ),
        "groupings": sorted(
            [
                {
                    "name": str(item.get("name") or ""),
                    "group_names": sorted(str(value) for value in item.get("group_names") or []),
                }
                for item in snapshot.get("groupings") or []
                if isinstance(item, dict) and str(item.get("name") or "")
            ],
            key=lambda item: item["name"].casefold(),
        ),
    }
    return fingerprint(canonical)


def sanitized_group_snapshot_hash(snapshot: dict[str, Any]) -> str:
    return _group_snapshot_hash(snapshot)


def _selected_group_names(source: dict[str, Any], snapshot: dict[str, Any]) -> list[str]:
    all_names = [
        str(item.get("name") or "")
        for item in snapshot.get("groups") or []
        if isinstance(item, dict) and str(item.get("name") or "")
    ]
    all_names = sorted(dict.fromkeys(all_names), key=str.casefold)
    if str(source.get("kind") or "") != "moodle_groups":
        raise DocumentChangeBlocked("repeatable record source must be moodle_groups")
    selection = source.get("selection") if isinstance(source.get("selection"), dict) else {}
    mode = str(selection.get("mode") or "all_groups")
    if mode == "all_groups":
        return all_names
    if mode == "explicit_names":
        requested = [str(value) for value in selection.get("group_names") or []]
        if not requested or len(requested) != len(set(requested)):
            raise DocumentChangeBlocked("explicit Group selection must be non-empty and unique")
        unknown = [value for value in requested if value not in all_names]
        if unknown:
            raise DocumentChangeBlocked(f"planner selected unknown Groups: {unknown}")
        return requested
    if mode == "grouping":
        grouping_name = str(selection.get("grouping_name") or "")
        matches = [
            item for item in snapshot.get("groupings") or []
            if isinstance(item, dict) and str(item.get("name") or "") == grouping_name
        ]
        if len(matches) != 1:
            raise DocumentChangeBlocked("planner-selected Grouping does not resolve uniquely")
        names = [str(value) for value in matches[0].get("group_names") or []]
        if not names:
            raise DocumentChangeBlocked("selected Grouping has no detailed Group membership readback")
        unknown = [value for value in names if value not in all_names]
        if unknown:
            raise DocumentChangeBlocked("Grouping readback contains unknown Groups")
        return names
    raise DocumentChangeBlocked(f"unsupported Group selection mode: {mode!r}")


def _parse_hhmm(value: str) -> datetime:
    if not re.fullmatch(r"(?:[01]?\d|2[0-3]):[0-5]\d", value):
        raise DocumentChangeBlocked(f"invalid HH:MM time: {value!r}")
    parsed = datetime.strptime(value, "%H:%M")
    return parsed


def _display_hhmm(value: datetime) -> str:
    return value.strftime("%H:%M")


def _materialize_records(
    operation: dict[str, Any],
    target: dict[str, Any],
    group_snapshot: dict[str, Any],
    *,
    run_id: str,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    names = _selected_group_names(operation.get("record_source") or {}, group_snapshot)
    roles = {str(key): int(value) for key, value in (operation.get("column_roles") or {}).items()}
    group_column = roles.get("group_name")
    if group_column is None:
        raise DocumentChangeBlocked("repeatable Group record rebuild requires a group_name column role")
    header_rows = int(operation.get("header_row_count") or 0)
    rows = target.get("rows") if isinstance(target.get("rows"), list) else []
    existing = []
    for row in rows[header_rows:]:
        if isinstance(row, list) and 0 <= group_column < len(row):
            value = str(row[group_column] or "")
            if value in names and value not in existing:
                existing.append(value)
    policy = str(operation.get("ordering_policy") or "")
    if policy == "preserve_existing_then_append_new":
        ordered = existing + [name for name in names if name not in existing]
    elif policy == "explicit_order":
        ordered = [str(value) for value in operation.get("explicit_order") or []]
        if len(ordered) != len(names) or set(ordered) != set(names):
            raise DocumentChangeBlocked("explicit Group order must contain every selected Group exactly once")
    elif policy == "seeded_random":
        basis = str(operation.get("explicit_prompt_basis") or "").strip()
        if not basis or str(operation.get("randomization_evidence_ref") or "") != "current_prompt":
            raise DocumentChangeBlocked("seeded_random requires planner evidence from the current staff prompt")
        ordered = list(names)
        seed_text = f"{run_id}|{_group_snapshot_hash(group_snapshot)}|{operation.get('id') or operation.get('target_id')}"
        seed = int(hashlib.sha256(seed_text.encode("utf-8")).hexdigest()[:16], 16)
        random.Random(seed).shuffle(ordered)
    else:
        raise DocumentChangeBlocked(f"unsupported ordering policy: {policy!r}")

    schedule = operation.get("schedule") if isinstance(operation.get("schedule"), dict) else {}
    duration = schedule.get("duration_minutes")
    start_text = str(schedule.get("window_start") or "")
    end_text = str(schedule.get("window_end") or "")
    start = _parse_hhmm(start_text) if start_text else None
    end = _parse_hhmm(end_text) if end_text else None
    if any(role in roles for role in {"start_time", "end_time", "prepare_start_time", "prepare_end_time"}):
        if start is None or duration in (None, ""):
            raise DocumentChangeBlocked("time column roles require planner-grounded window_start and duration_minutes")
    try:
        duration_int = int(duration) if duration not in (None, "") else 0
    except (TypeError, ValueError) as exc:
        raise DocumentChangeBlocked("duration_minutes must be an integer") from exc
    if duration_int < 0 or (start is not None and duration_int == 0):
        raise DocumentChangeBlocked("duration_minutes must be positive for a timed schedule")
    required_minutes = len(ordered) * duration_int
    if start is not None and end is not None and start + timedelta(minutes=required_minutes) > end:
        raise DocumentChangeBlocked(
            f"Group schedule needs {required_minutes} minutes but the planner-grounded window "
            f"{start_text}-{end_text} is too short"
        )

    static_values = operation.get("static_values") if isinstance(operation.get("static_values"), dict) else {}
    prepare_offset = int(schedule.get("prepare_offset_minutes") or 0)
    records: list[dict[str, Any]] = []
    for index, name in enumerate(ordered, start=1):
        record = {str(key): _json_value(value) for key, value in static_values.items()}
        record.update({"group_name": name, "ordinal": index})
        if start is not None:
            item_start = start + timedelta(minutes=(index - 1) * duration_int)
            item_end = item_start + timedelta(minutes=duration_int)
            record.update(
                {
                    "start_time": _display_hhmm(item_start),
                    "end_time": _display_hhmm(item_end),
                    "duration_minutes": duration_int,
                }
            )
            if prepare_offset:
                record["prepare_start_time"] = _display_hhmm(item_start - timedelta(minutes=prepare_offset))
                record["prepare_end_time"] = _display_hhmm(item_start)
        records.append(record)
    return records, {
        "group_snapshot_hash": _group_snapshot_hash(group_snapshot),
        "ordered_group_names": ordered,
        "required_minutes": required_minutes,
    }


def materialize_document_change_plan(
    plan: dict[str, Any],
    model: dict[str, Any],
    group_snapshot: dict[str, Any],
    *,
    run_id: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    materialized = deepcopy(plan)
    targets = structure_index(model)
    metadata: dict[str, Any] = {"group_snapshot_hash": _group_snapshot_hash(group_snapshot), "operations": []}
    for operation in materialized.get("operations") or []:
        if not isinstance(operation, dict) or operation.get("type") != "repeatable_record_rebuild":
            continue
        target = targets.get(str(operation.get("target_id") or ""))
        if target is None:
            continue
        records, item_metadata = _materialize_records(operation, target, group_snapshot, run_id=run_id)
        operation["records"] = records
        metadata["operations"].append({"id": operation.get("id"), **item_metadata})
    return materialized, metadata


def _replace_text_nodes(nodes: list[Any], before: str, after: str) -> bool:
    parts = [str(node.text or "") for node in nodes]
    joined = "".join(parts)
    start = joined.find(before)
    if start < 0:
        return False
    end = start + len(before)
    spans: list[tuple[int, int]] = []
    cursor = 0
    for part in parts:
        spans.append((cursor, cursor + len(part)))
        cursor += len(part)
    start_index = next((i for i, (left, right) in enumerate(spans) if left <= start < right), -1)
    end_index = next((i for i, (left, right) in enumerate(spans) if left < end <= right), -1)
    if start_index < 0 or end_index < 0:
        return False
    left = spans[start_index][0]
    right = spans[end_index][1]
    prefix = parts[start_index][: start - left]
    suffix = parts[end_index][len(parts[end_index]) - (right - end) :]
    nodes[start_index].text = prefix + after + (suffix if start_index == end_index else "")
    for index in range(start_index + 1, end_index):
        nodes[index].text = ""
    if end_index != start_index:
        nodes[end_index].text = suffix
    return True


def _set_text_nodes_preserving_structure(nodes: list[Any], value: str) -> bool:
    if not nodes:
        return False
    nodes[0].text = value
    for node in nodes[1:]:
        node.text = ""
    return True


def _docx_text_nodes(element: Any) -> list[Any]:
    return [node for node in element.iter() if str(node.tag).rsplit("}", 1)[-1] in {"t", "instrText"}]


def _apply_docx(path: Path, operations: list[dict[str, Any]]) -> None:
    from docx import Document

    document = Document(path)
    for operation in operations:
        target_id = str(operation.get("target_id") or "")
        change_type = str(operation.get("type") or "")
        paragraph_match = re.fullmatch(r"docx:p(\d{4})", target_id)
        cell_match = re.fullmatch(r"docx:t(\d{4}):r(\d{4}):c(\d{4})", target_id)
        table_match = re.fullmatch(r"docx:t(\d{4})", target_id)
        if change_type in {"anchored_text_edit", "scalar_fact_update"}:
            if paragraph_match:
                element = document.paragraphs[int(paragraph_match.group(1)) - 1]._p
            elif cell_match:
                table = document.tables[int(cell_match.group(1)) - 1]
                element = table.rows[int(cell_match.group(2)) - 1].cells[int(cell_match.group(3)) - 1]._tc
            else:
                raise DocumentChangeBlocked(f"unsupported DOCX text target: {target_id}")
            if not _replace_text_nodes(
                _docx_text_nodes(element),
                str(operation.get("expected_before") or ""),
                str(operation.get("replacement_text") or ""),
            ):
                raise DocumentChangeBlocked(f"DOCX target changed before write: {target_id}")
        elif change_type == "repeatable_record_rebuild":
            if not table_match:
                raise DocumentChangeBlocked(f"unsupported DOCX table target: {target_id}")
            table = document.tables[int(table_match.group(1)) - 1]
            header_rows = int(operation.get("header_row_count") or 0)
            if any(
                _docx_cell_structure(cell).get("grid_span") != 1
                or bool(_docx_cell_structure(cell).get("vertical_merge"))
                for row in table.rows[header_rows:]
                for cell in row.cells
            ):
                raise DocumentChangeBlocked("DOCX repeatable table rebuild does not allow merged data rows")
            template_index = int(operation.get("template_row_index") or header_rows)
            if not (0 <= template_index < len(table.rows)):
                raise DocumentChangeBlocked("DOCX template_row_index is out of range")
            template = deepcopy(table.rows[template_index]._tr)
            for row in list(table.rows[header_rows:]):
                table._tbl.remove(row._tr)
            roles = {str(key): int(value) for key, value in (operation.get("column_roles") or {}).items()}
            for record in operation.get("records") or []:
                table._tbl.append(deepcopy(template))
                row = table.rows[-1]
                for role, column_index in roles.items():
                    if role in record:
                        cell = row.cells[column_index]
                        if not _set_text_nodes_preserving_structure(
                            _docx_text_nodes(cell._tc), str(record[role])
                        ):
                            cell.text = str(record[role])
    document.save(path)


def _apply_xlsx(path: Path, operations: list[dict[str, Any]]) -> None:
    import openpyxl
    from openpyxl.formula.translate import Translator

    workbook = openpyxl.load_workbook(path, data_only=False)
    for operation in operations:
        target_id = str(operation.get("target_id") or "")
        cell_match = re.fullmatch(r"xlsx:s(\d{4}):c([A-Z]+\d+)", target_id)
        table_match = re.fullmatch(r"xlsx:s(\d{4})", target_id)
        change_type = str(operation.get("type") or "")
        if change_type in {"anchored_text_edit", "scalar_fact_update"}:
            if not cell_match:
                raise DocumentChangeBlocked(f"unsupported XLSX cell target: {target_id}")
            worksheet = workbook.worksheets[int(cell_match.group(1)) - 1]
            cell = worksheet[cell_match.group(2)]
            current = str(cell.value or "")
            before = str(operation.get("expected_before") or "")
            if before not in current:
                raise DocumentChangeBlocked(f"XLSX target changed before write: {target_id}")
            cell.value = current.replace(before, str(operation.get("replacement_text") or ""), 1)
        elif change_type == "repeatable_record_rebuild":
            if not table_match:
                raise DocumentChangeBlocked(f"unsupported XLSX table target: {target_id}")
            worksheet = workbook.worksheets[int(table_match.group(1)) - 1]
            if worksheet.merged_cells.ranges:
                raise DocumentChangeBlocked("XLSX repeatable table rebuild does not allow merged ranges")
            header_rows = int(operation.get("header_row_count") or 0)
            template_index = int(operation.get("template_row_index") or header_rows)
            template_row = template_index + 1
            if not (1 <= template_row <= worksheet.max_row):
                raise DocumentChangeBlocked("XLSX template_row_index is out of range")
            template_cells = [copy(worksheet.cell(template_row, col)) for col in range(1, worksheet.max_column + 1)]
            template_height = worksheet.row_dimensions[template_row].height
            data_start = header_rows + 1
            old_count = max(0, worksheet.max_row - header_rows)
            if old_count:
                worksheet.delete_rows(data_start, old_count)
            records = operation.get("records") or []
            if records:
                worksheet.insert_rows(data_start, len(records))
            roles = {str(key): int(value) for key, value in (operation.get("column_roles") or {}).items()}
            for record_index, record in enumerate(records):
                row_number = data_start + record_index
                worksheet.row_dimensions[row_number].height = template_height
                for column_index, source_cell in enumerate(template_cells, start=1):
                    target_cell = worksheet.cell(row_number, column_index)
                    if source_cell.has_style:
                        target_cell._style = copy(source_cell._style)
                    target_cell.number_format = source_cell.number_format
                    target_cell.font = copy(source_cell.font)
                    target_cell.fill = copy(source_cell.fill)
                    target_cell.border = copy(source_cell.border)
                    target_cell.alignment = copy(source_cell.alignment)
                    target_cell.protection = copy(source_cell.protection)
                    if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                        target_cell.value = Translator(
                            source_cell.value,
                            origin=source_cell.coordinate,
                        ).translate_formula(target_cell.coordinate)
                    else:
                        target_cell.value = source_cell.value
                for role, zero_based_column in roles.items():
                    if role in record:
                        worksheet.cell(row_number, zero_based_column + 1).value = record[role]
    workbook.save(path)


def _pptx_text_nodes(element: Any) -> list[Any]:
    return [node for node in element.iter() if str(node.tag).rsplit("}", 1)[-1] == "t"]


def _apply_pptx(path: Path, operations: list[dict[str, Any]]) -> None:
    from pptx import Presentation

    presentation = Presentation(path)
    for operation in operations:
        target_id = str(operation.get("target_id") or "")
        shape_match = re.fullmatch(r"pptx:s(\d{4}):sh(\d{4})", target_id)
        cell_match = re.fullmatch(r"pptx:s(\d{4}):sh(\d{4}):t0001:r(\d{4}):c(\d{4})", target_id)
        table_match = re.fullmatch(r"pptx:s(\d{4}):sh(\d{4}):t0001", target_id)
        change_type = str(operation.get("type") or "")
        if change_type in {"anchored_text_edit", "scalar_fact_update"}:
            if shape_match:
                element = presentation.slides[int(shape_match.group(1)) - 1].shapes[int(shape_match.group(2)) - 1]._element
            elif cell_match:
                shape = presentation.slides[int(cell_match.group(1)) - 1].shapes[int(cell_match.group(2)) - 1]
                element = shape.table.rows[int(cell_match.group(3)) - 1].cells[int(cell_match.group(4)) - 1]._tc
            else:
                raise DocumentChangeBlocked(f"unsupported PPTX text target: {target_id}")
            if not _replace_text_nodes(
                _pptx_text_nodes(element),
                str(operation.get("expected_before") or ""),
                str(operation.get("replacement_text") or ""),
            ):
                raise DocumentChangeBlocked(f"PPTX target changed before write: {target_id}")
        elif change_type == "repeatable_record_rebuild":
            if not table_match:
                raise DocumentChangeBlocked(f"unsupported PPTX table target: {target_id}")
            shape = presentation.slides[int(table_match.group(1)) - 1].shapes[int(table_match.group(2)) - 1]
            table = shape.table
            header_rows = int(operation.get("header_row_count") or 0)
            if any(
                bool(getattr(cell, "is_merge_origin", False)) or bool(getattr(cell, "is_spanned", False))
                for row in list(table.rows)[header_rows:]
                for cell in row.cells
            ):
                raise DocumentChangeBlocked("PPTX repeatable table rebuild does not allow merged data rows")
            template_index = int(operation.get("template_row_index") or header_rows)
            rows = list(table.rows)
            if not (0 <= template_index < len(rows)):
                raise DocumentChangeBlocked("PPTX template_row_index is out of range")
            template = deepcopy(rows[template_index]._tr)
            for row in list(table.rows)[header_rows:]:
                table._tbl.remove(row._tr)
            roles = {str(key): int(value) for key, value in (operation.get("column_roles") or {}).items()}
            for record in operation.get("records") or []:
                table._tbl.append(deepcopy(template))
                row = list(table.rows)[-1]
                for role, column_index in roles.items():
                    if role in record:
                        cell = row.cells[column_index]
                        if not _set_text_nodes_preserving_structure(
                            _pptx_text_nodes(cell._tc), str(record[role])
                        ):
                            cell.text = str(record[role])
    presentation.save(path)


def _verify_readback(
    path: Path,
    operations: list[dict[str, Any]],
    *,
    model: dict[str, Any] | None = None,
) -> list[str]:
    model = model or build_document_model(path)
    targets = structure_index(model)
    errors: list[str] = []
    for operation in operations:
        target_id = str(operation.get("target_id") or "")
        target = targets.get(target_id)
        change_type = str(operation.get("type") or "")
        readback = operation.get("readback") if isinstance(operation.get("readback"), dict) else {}
        replacement_value = operation.get("replacement_text")
        expected_value = readback.get("expected_text") if "expected_text" in readback else replacement_value
        expected = "" if expected_value is None else str(expected_value)
        if target is None:
            # Spreadsheet models omit empty cells. A precisely anchored clear therefore
            # removes the cell structure while the parent worksheet remains available
            # for scope/readback validation.
            if change_type in {"anchored_text_edit", "scalar_fact_update"} and expected == "":
                continue
            errors.append(f"readback target disappeared: {target_id}")
            continue
        if change_type in {"anchored_text_edit", "scalar_fact_update"}:
            text = str(target.get("text") if "text" in target else target.get("value") or "")
            if expected not in text:
                errors.append(f"readback text missing at {target_id}: {expected!r}")
        elif change_type == "repeatable_record_rebuild":
            header_rows = int(operation.get("header_row_count") or 0)
            records = operation.get("records") or []
            expected_rows = header_rows + len(records)
            if int(target.get("row_count") or 0) != expected_rows:
                errors.append(
                    f"readback row count mismatch at {target_id}: "
                    f"expected {expected_rows}, got {target.get('row_count')}"
                )
                continue
            rows = target.get("rows") if isinstance(target.get("rows"), list) else []
            roles = {str(key): int(value) for key, value in (operation.get("column_roles") or {}).items()}
            for record_index, record in enumerate(records):
                row = rows[header_rows + record_index] if header_rows + record_index < len(rows) else []
                for role, column_index in roles.items():
                    if role not in record:
                        continue
                    observed = row[column_index] if isinstance(row, list) and column_index < len(row) else None
                    if str(observed if observed is not None else "") != str(record[role]):
                        errors.append(
                            f"readback value mismatch at {target_id} row {record_index + 1} role {role}"
                        )
    return errors


def _unplanned_structure_changes(
    before_model: dict[str, Any],
    after_model: dict[str, Any],
    operations: list[dict[str, Any]],
) -> list[str]:
    before = structure_index(before_model)
    after = structure_index(after_model)
    allowed: set[str] = set()
    for operation in operations:
        target_id = str(operation.get("target_id") or "")
        target = before.get(target_id) or {}
        allowed.add(target_id)
        parent_id = str(target.get("parent_id") or "")
        if parent_id:
            allowed.add(parent_id)
        if target.get("kind") == "table":
            allowed.update(
                item_id
                for item_id, item in before.items()
                if str(item.get("parent_id") or "") == target_id
            )
            allowed.update(
                item_id
                for item_id, item in after.items()
                if str(item.get("parent_id") or "") == target_id
            )
    errors: list[str] = []
    for item_id in sorted((set(before) | set(after)) - allowed):
        old = before.get(item_id)
        new = after.get(item_id)
        if old is None or new is None or old.get("fingerprint") != new.get("fingerprint"):
            errors.append(f"unplanned structure changed during execution: {item_id}")
    return errors


def apply_document_change_plan(
    path: str | Path,
    plan: dict[str, Any],
    *,
    group_snapshot: dict[str, Any],
    run_id: str,
) -> dict[str, Any]:
    target_path = Path(path)
    before_model = build_document_model(target_path)
    errors = validate_document_change_plan(plan, before_model)
    has_group_rebuild = any(
        isinstance(item, dict) and item.get("type") == "repeatable_record_rebuild"
        for item in plan.get("operations") or []
    )
    if has_group_rebuild and str(plan.get("group_snapshot_hash") or "") != _group_snapshot_hash(group_snapshot):
        errors.append("plan group_snapshot_hash does not match the current sanitized Group snapshot")
    if errors:
        return {"ok": False, "status": "blocked_invalid_plan", "blockers": errors}
    temporary_path: Path | None = None
    try:
        materialized, metadata = materialize_document_change_plan(
            plan,
            before_model,
            group_snapshot,
            run_id=run_id,
        )
        operations = [item for item in materialized.get("operations") or [] if isinstance(item, dict)]
        suffix = target_path.suffix.lower()
        with tempfile.NamedTemporaryFile(
            prefix=f".{target_path.stem}.document-change-",
            suffix=suffix,
            dir=target_path.parent,
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
        shutil.copy2(target_path, temporary_path)
        if suffix == ".docx":
            _apply_docx(temporary_path, operations)
        elif suffix == ".xlsx":
            _apply_xlsx(temporary_path, operations)
        elif suffix == ".pptx":
            _apply_pptx(temporary_path, operations)
        else:
            raise DocumentChangeBlocked(f"unsupported editable Office format: {suffix}")
        after_model = build_document_model(temporary_path)
        readback_errors = _verify_readback(temporary_path, operations, model=after_model)
        readback_errors.extend(_unplanned_structure_changes(before_model, after_model, operations))
        if readback_errors:
            return {"ok": False, "status": "blocked_readback_failed", "blockers": readback_errors}
        output_fingerprint = after_model.get("document_fingerprint")
        os.replace(temporary_path, target_path)
        temporary_path = None
        return {
            "ok": True,
            "status": "modified_by_document_change_plan" if operations else "unchanged_empty_plan",
            "operation_count": len(operations),
            "materialization": metadata,
            "source_document_fingerprint": before_model.get("document_fingerprint"),
            "output_document_fingerprint": output_fingerprint,
        }
    except DocumentChangeBlocked as exc:
        return {"ok": False, "status": "blocked_document_change", "blockers": [str(exc)]}
    except Exception as exc:
        return {"ok": False, "status": "blocked_document_change_error", "blockers": [str(exc)]}
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
