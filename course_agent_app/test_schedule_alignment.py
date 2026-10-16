import copy
import tempfile
import unittest
from concurrent.futures import Future
from pathlib import Path
from unittest.mock import Mock
from docx import Document
from chrome_file_library import FileLibrary
from chrome_task_memory import continuation
from chrome_approval import ChromeApproval
from chrome_schedule_alignment import validate,check_review,verify_current,review_context,apply_review,partition_review,normalize_bindings,inherit_citation_sources,resolve_missing_destinations,normalize_week_columns,AlignmentReviewRequired
from run_shared_chrome_agent import Runtime


class AlignmentTests(unittest.TestCase):
    def setUp(self):
        self.browser=Mock();self.browser.forms={};self.browser.overviews={}
        self.browser.targets={
            'a:m1':{'key':'a:m1','course_id':'a','kind':'module','title':'Storyboard Development','section_key':'a:s6'},
            'a:m2':{'key':'a:m2','course_id':'a','kind':'module','title':'XRCC tutorial 3','section_key':'a:s5'},
            'a:m3':{'key':'a:m3','course_id':'a','kind':'module','title':'XRCC tutorial 4','section_key':'a:s6'},
            **{f'a:s{n}':{'key':f'a:s{n}','course_id':'a','kind':'section','title':f'Week {n}','visibility':'shown'} for n in (5,6,8)}}
        self.evidence=[{'file_id':'syllabus-a','status':'read','sha256':'fixture','table_rows':[
            {'row_id':'t1:r5','cells':['5','Storyboard Development']},
            {'row_id':'t1:r6','cells':['6','XRCC tutorial 3']},
            {'row_id':'t1:r8','cells':['8','XRCC tutorial 4']}]}]
        self.contract={'required':True,'course_ids':['a'],'sources':[{'course_id':'a','file_id':'syllabus-a'}],
            'items':[{'target':f'a:m{i}','target_section':f'a:s{w}','file_id':'syllabus-a','row_id':f't1:r{w}',
                      'week_column':0,'disposition':'place','visibility':'preserve','reason':'Topic matches the cited row'}
                     for i,w in enumerate((5,6,8),1)]}
        self.ops=[{'id':str(i),'capability':'move_module','target':r['target'],'params':{'target_section':r['target_section']}}
                  for i,r in enumerate(self.contract['items'])]

    def test_actual_failure_wrong_destination_rejected_before_any_write(self):
        self.contract['items'][0]['target_section']='a:s6';self.ops[0]['params']['target_section']='a:s6'
        with self.assertRaisesRegex(ValueError,'row week'):validate(self.contract,self.ops,self.browser,self.evidence)
        self.browser.execute.assert_not_called()

    def test_auxiliary_documents_do_not_become_authoritative_syllabuses(self):
        self.contract['sources'].append({'course_id':'a','file_id':'slides'})
        self.evidence.append({'file_id':'slides','status':'read','text':'Workshop slides'})
        normalized,notes=normalize_bindings(self.contract,self.evidence)
        self.assertEqual(normalized['sources'],[{'course_id':'a','file_id':'syllabus-a'}])
        self.assertEqual(notes[0]['supplemental_file_ids'],['slides'])
        validate(normalized,self.ops,self.browser,self.evidence)
        self.assertEqual(len(self.contract['sources']),2)
        self.assertEqual(len(self.evidence),2)

    def test_normalization_never_chooses_between_competing_structured_sources(self):
        self.contract['sources'].append({'course_id':'a','file_id':'other-syllabus'})
        self.evidence.append({**self.evidence[0],'file_id':'other-syllabus'})
        normalized,notes=normalize_bindings(self.contract,self.evidence)
        self.assertFalse(notes)
        with self.assertRaisesRegex(ValueError,'exactly one'):validate(normalized,self.ops,self.browser,self.evidence)

    def test_unread_or_cited_auxiliary_binding_is_not_silently_dropped(self):
        self.contract['sources'].append({'course_id':'a','file_id':'slides'})
        normalized,notes=normalize_bindings(self.contract,self.evidence)
        self.assertFalse(notes)
        self.evidence.append({'file_id':'slides','status':'read','text':'Slides'})
        self.contract['items'][0]['file_id']='slides'
        normalized,notes=normalize_bindings(self.contract,self.evidence)
        self.assertFalse(notes)
        with self.assertRaises(ValueError):validate(normalized,self.ops,self.browser,self.evidence)

    def test_real_six_binding_failure_retains_full_source_evidence_and_23_moves(self):
        import json
        data=json.loads((Path(__file__).parent/'eval_fixtures/source_binding_recovery.json').read_text(encoding='utf-8'))
        browser=Mock();browser.targets=data['targets']
        normalized,notes=normalize_bindings(data['alignment'],data['evidence'])
        self.assertEqual(len(data['alignment']['sources']),6)
        self.assertEqual(len(normalized['sources']),1)
        self.assertEqual(len(notes[0]['supplemental_file_ids']),5)
        normalized,columns=normalize_week_columns(normalized,browser,data['evidence'])
        self.assertEqual(len(columns),28)
        contract=validate(normalized,data['operations'],browser,data['evidence'])
        self.assertEqual(len(contract['items']),59)
        self.assertEqual(len(data['operations']),23)
        browser.execute.assert_not_called()

    def test_week_column_normalization_requires_header_and_never_changes_destination(self):
        self.contract['items'][0]['week_column']=1
        unchanged,notes=normalize_week_columns(self.contract,self.browser,self.evidence)
        self.assertFalse(notes)
        self.evidence[0]['table_rows'].insert(0,{'row_id':'t1:r1','cells':['Week','Topic']})
        fixed,notes=normalize_week_columns(self.contract,self.browser,self.evidence)
        self.assertEqual(fixed['items'][0]['week_column'],0)
        self.assertEqual(fixed['items'][0]['target_section'],'a:s5')
        self.contract['items'][0]['target_section']='a:s6'
        unchanged,notes=normalize_week_columns(self.contract,self.browser,self.evidence)
        self.assertFalse(notes)
        with self.assertRaises(ValueError):validate(unchanged,self.ops,self.browser,self.evidence)

    def test_blank_citation_inherits_only_observed_course_binding(self):
        for row in self.contract['items']:row['file_id']=''
        fixed,notes=inherit_citation_sources(self.contract,self.browser,self.evidence)
        self.assertEqual(len(notes),3)
        self.assertTrue(all(r['file_id']=='syllabus-a' for r in fixed['items']))
        validate(fixed,self.ops,self.browser,self.evidence)
        self.assertEqual(self.contract['items'][0]['file_id'],'')
        twice,notes=inherit_citation_sources(fixed,self.browser,self.evidence)
        self.assertEqual(twice,fixed);self.assertFalse(notes)

    def test_explicit_wrong_source_is_not_replaced_by_inheritance(self):
        self.contract['items'][0]['file_id']='syllabus-b'
        fixed,notes=inherit_citation_sources(self.contract,self.browser,self.evidence)
        self.assertFalse(notes)
        self.assertEqual(fixed['items'][0]['file_id'],'syllabus-b')
        with self.assertRaisesRegex(ValueError,'explicit source citation'):validate(fixed,self.ops,self.browser,self.evidence)

    def test_inheritance_rejects_missing_evidence_ambiguous_binding_and_unknown_row(self):
        self.contract['items'][0]['file_id']=''
        for case in ('unread','ambiguous','unknown_row','unknown_target'):
            with self.subTest(case=case):
                contract=copy.deepcopy(self.contract);evidence=copy.deepcopy(self.evidence)
                if case=='unread':evidence=[]
                if case=='ambiguous':contract['sources'].append({'course_id':'a','file_id':'other'})
                if case=='unknown_row':contract['items'][0]['row_id']='t1:r999'
                if case=='unknown_target':contract['items'][0]['target']='b:unobserved'
                fixed,notes=inherit_citation_sources(contract,self.browser,evidence)
                self.assertFalse(notes);self.assertEqual(fixed['items'][0]['file_id'],'')
                with self.assertRaises(ValueError):validate(fixed,self.ops,self.browser,evidence)

    def test_two_courses_inherit_their_own_sources_even_when_row_ids_match(self):
        self.contract['course_ids'].append('b');self.contract['sources'].append({'course_id':'b','file_id':'syllabus-b'})
        self.evidence.append({**self.evidence[0],'file_id':'syllabus-b'})
        self.browser.targets['b:m1']={'kind':'module','course_id':'b','title':'B material'}
        self.browser.targets['b:s5']={'kind':'section','course_id':'b','title':'Week 5'}
        self.contract['items'].append({**self.contract['items'][0],'target':'b:m1','target_section':'b:s5','file_id':''})
        self.contract['items'][0]['file_id']=''
        fixed,notes=inherit_citation_sources(self.contract,self.browser,self.evidence)
        self.assertEqual(fixed['items'][0]['file_id'],'syllabus-a')
        self.assertEqual(fixed['items'][-1]['file_id'],'syllabus-b')
        validate(fixed,self.ops,self.browser,self.evidence)

    def test_real_empty_citation_failures_pass_structural_checks_after_compilation(self):
        import json
        cases=json.loads((Path(__file__).parent/'eval_fixtures/citation_inheritance.json').read_text(encoding='utf-8'))
        for case in cases:
            with self.subTest(plan=case['plan_artifact']):
                browser=Mock();browser.targets=case['targets']
                fixed,notes=inherit_citation_sources(case['alignment'],browser,case['evidence'])
                self.assertEqual(len(notes),case['blank_citations'])
                fixed,_=normalize_week_columns(fixed,browser,case['evidence'])
                contract=validate(fixed,case['operations'],browser,case['evidence'])
                self.assertEqual(len(contract['items']),59)
                browser.execute.assert_not_called()

    def test_unknown_destination_compiles_only_unique_cited_week_in_same_course(self):
        self.evidence[0]['table_rows'].insert(0,{'row_id':'t1:r1','cells':['Week','Topic']})
        self.contract['items'][0]['target_section']='a:section:invented'
        self.ops[0]['params']['target_section']='a:section:invented'
        fixed,ops,notes=resolve_missing_destinations(self.contract,self.ops,self.browser,self.evidence)
        self.assertEqual(fixed['items'][0]['target_section'],'a:s5')
        self.assertEqual(ops[0]['params']['target_section'],'a:s5')
        self.assertEqual(ops[0]['target'],'a:m1')
        self.assertEqual(len(notes),1)
        validate(fixed,ops,self.browser,self.evidence)
        self.assertEqual(self.ops[0]['params']['target_section'],'a:section:invented')

    def test_destination_compiler_does_not_guess_conflicts_or_module_identity(self):
        self.evidence[0]['table_rows'].insert(0,{'row_id':'t1:r1','cells':['Week','Topic']})
        for case in ('existing_wrong_week','other_course','unknown_module','duplicate_week','no_header'):
            with self.subTest(case=case):
                c=copy.deepcopy(self.contract);b=Mock();b.targets=copy.deepcopy(self.browser.targets);e=copy.deepcopy(self.evidence)
                c['items'][0]['target_section']='a:section:invented'
                if case=='existing_wrong_week':c['items'][0]['target_section']='a:s6'
                if case=='other_course':c['items'][0]['target_section']='b:section:5'
                if case=='unknown_module':c['items'][0]['target']='a:missing'
                if case=='duplicate_week':b.targets['a:second5']={**b.targets['a:s5'],'key':'a:second5'}
                if case=='no_header':e[0]['table_rows'].pop(0)
                fixed,ops,notes=resolve_missing_destinations(c,self.ops,b,e)
                self.assertFalse(notes);self.assertEqual(fixed,c)
                with self.assertRaises(ValueError):validate(fixed,ops,b,e)

    def test_real_invented_week10_section_resolves_without_rewriting_materials(self):
        import json
        d=json.loads((Path(__file__).parent/'eval_fixtures/inventory_grounding.json').read_text(encoding='utf-8'))
        b=Mock();b.targets=d['before_targets']
        fixed,ops,notes=resolve_missing_destinations(d['alignment'],d['operations'],b,d['evidence'])
        self.assertEqual(len(notes),3)
        self.assertTrue(all(n['section']=='140147:section:1607935' for n in notes))
        self.assertEqual(ops,d['operations'])
        validate(fixed,ops,b,d['evidence'])

    def test_valid_week_5_6_8_mapping_and_post_readback(self):
        validated=validate(self.contract,self.ops,self.browser,self.evidence)
        self.assertEqual(len(verify_current(validated,self.browser)),3)
        for r in validated['items']:self.browser.targets[r['target']]['section_key']=r['target_section']
        self.assertEqual(verify_current(validated,self.browser),[])

    def test_history_only_and_missing_sources_cannot_authorize_schedule_plan(self):
        with self.assertRaisesRegex(ValueError,'source-grounded'):validate({},self.ops,self.browser,[],required=True)
        with self.assertRaisesRegex(ValueError,'Read the authoritative'):validate(self.contract,self.ops,self.browser,[])

    def test_cross_course_citation_and_invented_row_rejected(self):
        for key,value in [('file_id','syllabus-b'),('row_id','t99:r99')]:
            with self.subTest(key=key):
                bad=copy.deepcopy(self.contract);bad['items'][0][key]=value
                with self.assertRaises(ValueError):validate(bad,self.ops,self.browser,self.evidence)

    def test_missing_coverage_and_uncited_move_rejected(self):
        bad=copy.deepcopy(self.contract);bad['items'].pop()
        with self.assertRaisesRegex(ValueError,'every current activity'):validate(bad,self.ops,self.browser,self.evidence)
        self.ops[0]['params']['target_section']='a:s6'
        with self.assertRaisesRegex(ValueError,'Every move'):validate(self.contract,self.ops,self.browser,self.evidence)

    def test_independent_review_must_accept_every_exact_target(self):
        review={'accepted':True,'findings':[],'items':[{'target':r['target'],'accepted':True,'reason':'Checked'} for r in self.contract['items']]}
        check_review(review,self.contract)
        review['items'][0].update(accepted=False,reason='Storyboard belongs to Week 5, not the cited tutorial row.')
        with self.assertRaisesRegex(ValueError,'Storyboard'):check_review(review,self.contract)
        review['items'].pop(0)
        with self.assertRaises(ValueError):check_review(review,self.contract)

    def test_future_uncertainty_does_not_block_an_independent_accepted_move(self):
        review={'accepted':False,'findings':[],'items':[
            {'target':r['target'],'accepted':r['target']!='a:m3','reason':'Need the material contents'} for r in self.contract['items']]}
        revised,notes=apply_review(review,self.contract,[self.ops[0]],self.browser)
        self.assertEqual(revised['items'][2]['disposition'],'unresolved')
        self.assertEqual(revised['items'][2]['target_section'],'');self.assertTrue(notes)
        self.assertEqual(self.contract['items'][2]['disposition'],'place')
        with self.assertRaises(AlignmentReviewRequired):apply_review(review,self.contract,[self.ops[2]],self.browser)
        with self.assertRaises(AlignmentReviewRequired):apply_review(review,self.contract,[],self.browser)

    def test_section_prerequisite_cannot_change_rejected_descendants(self):
        review={'accepted':False,'findings':[],'items':[
            {'target':r['target'],'accepted':r['target']!='a:m3','reason':'Unclear material'} for r in self.contract['items']]}
        safe={'capability':'set_section_visibility','target':'a:s5'}
        revised,notes=apply_review(review,self.contract,[safe],self.browser)
        self.assertTrue(notes)
        with self.assertRaises(AlignmentReviewRequired):
            apply_review(review,self.contract,[{'capability':'set_section_visibility','target':'a:s6'}],self.browser)
        review['items'].pop()
        with self.assertRaises(AlignmentReviewRequired):apply_review(review,self.contract,[safe],self.browser)

    def test_partition_holds_rejected_writes_and_transitive_dependencies(self):
        review={'accepted':False,'items':[{'target':r['target'],'accepted':r['target']!='a:m3','reason':'Ambiguous'} for r in self.contract['items']]}
        contract,ops,notes=partition_review(review,self.contract,self.ops,self.browser)
        self.assertEqual([o['target'] for o in ops],['a:m1','a:m2'])
        self.assertEqual(len(contract['items']),3)
        self.assertEqual(contract['items'][2]['disposition'],'unresolved')
        self.assertEqual(len(self.ops),3)
        self.ops[1]['depends_on']=['2']
        contract,ops,notes=partition_review(review,self.contract,self.ops,self.browser)
        self.assertEqual([o['target'] for o in ops],['a:m1'])
        self.assertEqual(contract['items'][1]['disposition'],'unresolved')
        self.ops[0]['depends_on']=['1']
        with self.assertRaises(AlignmentReviewRequired):partition_review(review,self.contract,self.ops,self.browser)

    def test_partition_holds_ordering_anchor_and_section_effects(self):
        review={'accepted':False,'items':[{'target':r['target'],'accepted':r['target']!='a:m3','reason':'Ambiguous'} for r in self.contract['items']]}
        self.ops[0]['params']['before_module']='a:m3'
        section={'id':'show','capability':'set_section_visibility','target':'a:s6'}
        contract,ops,_=partition_review(review,self.contract,self.ops+[section],self.browser)
        self.assertEqual([o['target'] for o in ops],['a:m2'])

    def test_unresolved_material_prevents_whole_goal_success(self):
        for r in self.contract['items']:self.browser.targets[r['target']]['section_key']=r['target_section']
        self.contract['items'][0].update(disposition='unresolved',reason='Source material unavailable')
        self.assertEqual(verify_current(self.contract,self.browser),['Source material unavailable'])

    def test_changed_inventory_and_visibility_are_in_final_review(self):
        for r in self.contract['items']:self.browser.targets[r['target']]['section_key']=r['target_section']
        self.contract['items'][0]['visibility']='shown'
        self.browser.inspect.return_value={'fields':[{'key':'id_visible','value':'0'}]}
        self.assertIn('not shown',verify_current(self.contract,self.browser)[0])
        self.browser.targets['a:new']={'kind':'module','course_id':'a','title':'New item'}
        self.assertIn('inventory changed',verify_current(self.contract,self.browser)[0])

    def test_docx_rows_preserve_the_original_week_cells(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);doc=Document();t=doc.add_table(rows=2,cols=2)
            t.cell(0,0).text='Week';t.cell(0,1).text='Topic';t.cell(1,0).text='5';t.cell(1,1).text='Storyboard Development'
            doc.save(root/'source.docx');lib=FileLibrary(root,root/'index.json');snap=lib.snapshot()
            result=lib.inspect(snap,snap['files'][0]['id'])
            self.assertEqual(result['table_rows'][1],{'row_id':'t1:r2','cells':['5','Storyboard Development']})

    def test_section_renamed_after_approval_is_not_verified_by_id_alone(self):
        validated=validate(self.contract,self.ops,self.browser,self.evidence)
        for r in validated['items']:self.browser.targets[r['target']]['section_key']=r['target_section']
        self.browser.targets['a:s5']['title']='Week 9'
        self.assertIn('week label changed',verify_current(validated,self.browser)[0])
        with self.assertRaisesRegex(ValueError,'row week'):
            validate(validated,self.ops,self.browser,self.evidence,required=True)

    def test_reviewer_cannot_confuse_current_form_section_with_proposed_week(self):
        self.browser.overviews={'a':{'course_id':'a','label':'Class A'}}
        self.browser.forms={'a:m1':{'resource_files':['Storyboard_Development.pdf'],'fields':[
            {'label':'Section','type':'select-one','value':'a:s6'},
            {'label':'Description','type':'textarea','value':'Storyboard exercise content'}]}}
        context=review_context(self.contract,self.browser,self.evidence,'Align syllabus',[])
        item=context['alignment']['items'][0]
        self.assertEqual(item['current_section_title'],'Week 6')
        self.assertEqual(item['proposed_section_title'],'Week 5')
        self.assertEqual(item['cited_row_cells'],['5','Storyboard Development'])
        self.assertEqual(item['observed_content']['resource_files'],['Storyboard_Development.pdf'])
        self.assertEqual(item['observed_content']['text_fields'],[{'label':'Description','value':'Storyboard exercise content'}])
        self.assertNotIn('forms',context)
        self.contract['items'][0]['disposition']='unresolved'
        context=review_context(self.contract,self.browser,self.evidence,'Align syllabus',[])
        item=context['alignment']['items'][0]
        self.assertEqual(item['proposed_section_title'],'')
        self.assertEqual(item['target_section'],'')
        self.assertNotIn('cited_row_cells',item)


class GoalLifecycleTests(unittest.TestCase):
    def test_runtime_inherits_citations_before_independent_review_not_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            fixture=AlignmentTests();fixture.setUp()
            runtime=Runtime(Mock(),fixture.browser,['chat'],[],Mock(),Path(tmp),Mock())
            flow=runtime.flows['chat'];flow.begin('Align syllabus','staff','msg')
            runtime.file_evidence=fixture.evidence
            candidate={'decision':'plan','alignment':fixture.contract,'operations':fixture.ops}
            for row in candidate['alignment']['items']:row['file_id']=''
            runtime.submit=Mock(side_effect=lambda chat,context,*args:runtime.jobs.update({chat:{'context':context}}))
            self.assertFalse(runtime.ensure_alignment_review('chat',candidate,{'rounds':1,'repairs':1,'context':{}}))
            context=runtime.submit.call_args.args[1]
            self.assertEqual(context['review_kind'],'schedule_alignment')
            self.assertTrue(all(r['file_id']=='syllabus-a' for r in context['alignment']['items']))
            self.assertEqual(flow.state['phase'],'PLANNING')
            self.assertFalse(flow.state.get('approval_sent'))
            fixture.browser.execute.assert_not_called()

    def review_runtime(self,tmp,*,touch_rejected=False):
        fixture=AlignmentTests();fixture.setUp();browser=fixture.browser
        runtime=Runtime(Mock(),browser,['chat'],[],Mock(),Path(tmp),Mock());runtime.send=Mock()
        flow=runtime.flows['chat'];flow.begin('Align syllabus','staff','msg')
        candidate={'decision':'plan','alignment':fixture.contract,'operations':[fixture.ops[2 if touch_rejected else 0]],
                   'impact_review':{'unresolved_dependencies':[]}}
        future=Future();future.set_result({'accepted':False,'findings':[],'items':[
            {'target':r['target'],'accepted':r['target']!='a:m3','reason':'Need source contents'} for r in fixture.contract['items']]})
        runtime.jobs['chat']={'future':future,'context':{'review_kind':'schedule_alignment'},'alignment_candidate':candidate,
            'review_contract':fixture.contract,'origin_job':{'context':{},'rounds':1,'repairs':1},'rounds':1,'repairs':1}
        return runtime

    def test_runtime_holds_future_item_and_offers_only_reviewed_batch(self):
        from chrome_approval import digest
        with tempfile.TemporaryDirectory() as tmp:
            runtime=self.review_runtime(tmp);runtime.accept_decision=Mock()
            runtime.flows['chat'].state['alignment_repairs']=2
            runtime.complete_jobs()
            _,candidate,job=runtime.accept_decision.call_args.args
            self.assertEqual(candidate['alignment']['items'][2]['disposition'],'unresolved')
            self.assertEqual(candidate['operations'][0]['target'],'a:m1')
            self.assertTrue(candidate['impact_review']['unresolved_dependencies'])
            self.assertEqual(job['alignment_checked'],digest(candidate))

    def test_unclassified_review_failure_requires_technical_repair_not_generic_question(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime=self.review_runtime(tmp,touch_rejected=True);runtime.accept_decision=Mock()
            runtime.flows['chat'].state['alignment_repairs']=2
            runtime.complete_jobs()
            self.assertEqual(runtime.flows['chat'].state['phase'],'FAILED')
            runtime.accept_decision.assert_not_called()
            self.assertIn('technical repair',runtime.send.call_args.args[1])
            self.assertNotIn('Please specify',runtime.send.call_args.args[1])
            self.assertNotIn('Rejected',runtime.send.call_args.args[1])

    def test_runtime_partitions_after_bounded_semantic_repair(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime=self.review_runtime(tmp,touch_rejected=True);runtime.accept_decision=Mock()
            runtime.flows['chat'].state['alignment_repairs']=2
            fixture=AlignmentTests();fixture.setUp()
            runtime.jobs['chat']['alignment_candidate']['operations']=fixture.ops
            runtime.complete_jobs()
            _,candidate,_=runtime.accept_decision.call_args.args
            self.assertEqual([o['target'] for o in candidate['operations']],['a:m1','a:m2'])
            self.assertEqual(len(candidate['alignment']['items']),3)
            runtime.send.assert_not_called()

    def test_coordinate_repair_does_not_consume_semantic_repair_budget(self):
        with tempfile.TemporaryDirectory() as tmp:
            runtime=self.review_runtime(tmp,touch_rejected=True)
            runtime.accept_decision=Mock();runtime.submit=Mock()
            runtime.complete_jobs()
            self.assertEqual(runtime.flows['chat'].state['phase'],'PLANNING')
            self.assertEqual(runtime.flows['chat'].state['alignment_repairs'],1)
            self.assertEqual(runtime.submit.call_args.args[-1],1)
            self.assertIn('Need source contents',runtime.submit.call_args.args[1]['feedback'])
            runtime.send.assert_not_called();runtime.accept_decision.assert_not_called()

    def test_real_rejected_batch_keeps_twenty_moves_and_all_59_goal_items(self):
        import json
        fixture=json.loads((Path(__file__).parent/'eval_fixtures/review_recovery.json').read_text(encoding='utf-8'))
        browser=Mock();browser.targets=fixture['targets']
        contract,ops,notes=partition_review(fixture['review'],fixture['contract'],fixture['operations'],browser)
        self.assertEqual(len(fixture['operations']),21)
        self.assertEqual(len(ops),20);self.assertEqual(len(contract['items']),59)
        self.assertNotIn('140147:module:4254911',{o['target'] for o in ops})
        self.assertTrue(all(r['disposition']=='unresolved' for r in contract['items'] if r['target']=='140147:module:4254911'))
        self.assertEqual(len(notes),7)
        browser.execute.assert_not_called()

    def test_operator_correction_preserves_receipts_without_reusing_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            b=Mock();b.forms={};b.overviews={};b.targets={}
            runtime=Runtime(Mock(),b,['chat'],[],Mock(),Path(tmp),Mock())
            flow=runtime.flows['chat'];flow.begin('Align syllabus','staff','msg')
            flow.state.update(phase='DONE',mention_authorized=True,conversation=[],
                selected=[{'id':'move','target':'a:m1','capability':'move_module','params':{'target_section':'a:s6'}}],
                results=[{'verified':True}],write_started='move',approval_sent=True)
            runtime.recheck_task(flow.state['task_id'],'Correct Class A against its internal syllabus.')
            queued=runtime.queue[0]
            self.assertEqual(len(queued['completed_operation_memory']),1)
            self.assertTrue(queued['schedule_required'])
            self.assertNotIn('approval_sent',queued);self.assertNotIn('selected',queued)
            runtime.queue.clear();flow.state['phase']='NEEDS_READBACK'
            with self.assertRaises(ValueError):runtime.recheck_task(flow.state['task_id'],'Correct it')

    def test_move_renewal_replans_instead_of_reusing_old_manifest(self):
        from chrome_approval import digest
        with tempfile.TemporaryDirectory() as tmp:
            b=Mock();b.forms={};b.overviews={};b.targets={}
            runtime=Runtime(Mock(),b,['chat'],[],Mock(),Path(tmp),Mock());runtime.start=Mock()
            flow=runtime.flows['chat'];flow.begin('Align syllabus','staff','msg')
            selected=[{'id':'move','target':'a:m1','capability':'move_module','params':{'target_section':'a:s6'}}]
            flow.state.update(phase='EXPIRED',mention_authorized=True,selected=selected,results=[],manifest_hash=digest(selected))
            runtime.renew_approval(flow.state['run'])
            self.assertTrue(runtime.start.called);b.inspect.assert_not_called()
            self.assertNotIn('manifest_hash',runtime.start.call_args.args[1])

    def test_expired_section_prerequisite_keeps_all_44_parent_items(self):
        scope=[{'id':str(n),'capability':'move_module','target':str(n),'params':{'target_section':'w5'},'fields':[]} for n in range(44)]
        state={'run':'a'*12,'task':'Align the course','phase':'EXPIRED','results':[],
               'previously_selected_scope':scope,'selected':[{'id':'show','capability':'set_section_visibility','target':'s6','params':{'action':'show'},'fields':[]}]}
        first=continuation(state)
        self.assertEqual(first['previously_selected_scope'],scope)
        state.update(phase='DONE',results=[{'verified':True}],write_started='show',**{'goal_scope':first['goal_scope']})
        self.assertEqual(continuation(state)['previously_selected_scope'],scope)

    def test_batch_receipt_does_not_emit_done_before_goal_verification(self):
        with tempfile.TemporaryDirectory() as tmp:
            flow=ChromeApproval(Path(tmp)/'state.json','chat');flow.begin('Align syllabus','staff','msg')
            op={'id':'one','capability':'move_module','target':'x','scope':'requested','summary':'Move','depends_on':[]}
            flow.state['requires_goal_verification']=True;flow.offer({},[op]);flow.mark_sent()
            flow.approve('OK '+flow.state['code'],'staff','approval');flow.before_save(op);flow.receipt({'verified':True})
            self.assertEqual(flow.state['phase'],'INTERRUPTED')

    def test_incomplete_batch_queues_followup_and_keeps_unresolved_scope(self):
        with tempfile.TemporaryDirectory() as tmp:
            b=Mock();b.forms={};b.overviews={};b.targets={}
            runtime=Runtime(Mock(),b,['chat'],[],Mock(),Path(tmp),Mock());runtime.send=Mock()
            flow=runtime.flows['chat'];flow.begin('Align syllabus','staff','msg')
            op={'id':'one','capability':'update_moodle_form','target':'x','params':{},'fields':[]}
            flow.state.update(phase='INTERRUPTED',selected=[op],results=[{'verified':True}],write_started='one',
                plan={'impact_review':{'unresolved_dependencies':['Need missing material']}})
            self.assertTrue(runtime.finish_goal_review('chat'))
            self.assertEqual(flow.state['phase'],'INTERRUPTED');self.assertEqual(len(runtime.queue),1)
            self.assertNotIn('Status: Done',runtime.send.call_args.args[1])
