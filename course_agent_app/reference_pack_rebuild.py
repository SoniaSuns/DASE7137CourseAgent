from __future__ import annotations

import csv
import copy
import hashlib
import json
import re
import shutil
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OPERATION_KIND = "rebuild_reference_pack"
EXECUTION_MODE = "local_file_action"

REFERENCE_ROOT = PROJECT_ROOT / "course_files" / "reference"
REPORT_ROOT = PROJECT_ROOT / "course_files" / "evidence" / "process" / "course_agent_app" / "reference_pack_rebuild"
DEFAULT_REFERENCE_PACK_SEMESTER = "2026-2027-sem1"
DEFAULT_TIMEZONE = "Asia/Hong_Kong"
WEEKDAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
REFERENCE_SOURCE_ROLES = {
    "timeline_build_report",
    "calendar_grounding_report",
    "timeline_reference",
    "calendar_reference",
    "source_calendar_reference",
    "target_calendar_reference",
    "syllabus_reference",
    "refresh_instructions",
    "timeline_json",
    "calendar_json",
    "reference",
}
REFERENCE_SOURCE_ROLE_ALIASES = {
    "build_report": "timeline_build_report",
    "timeline_report": "timeline_build_report",
    "calendar_report": "calendar_grounding_report",
    "calendar_grounding": "calendar_grounding_report",
    "source_calendar": "source_calendar_reference",
    "source_calendar_pdf": "source_calendar_reference",
    "old_calendar_reference": "source_calendar_reference",
    "target_calendar": "target_calendar_reference",
    "target_calendar_pdf": "target_calendar_reference",
    "existing_timeline_json": "timeline_json",
    "existing_calendar_json": "calendar_json",
    "instructions": "refresh_instructions",
    "rules": "refresh_instructions",
}
FORBIDDEN_JOB_KEYS = {
    "command",
    "shell",
    "python",
    "python_code",
    "javascript",
    "javascript_code",
    "raw_code",
    "script_text",
}


def now_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")


def safe_name(value: str, fallback: str = "reference") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "-", str(value or "")).strip(" .-_")
    cleaned = re.sub(r"\s+", "-", cleaned)
    return cleaned[:120] or fallback


def resolve_local_path(path_text: str | Path) -> Path:
    path = Path(str(path_text or "")).expanduser()
    if path.is_absolute():
        return path
    for base in (PROJECT_ROOT, Path.cwd()):
        candidate = base / path
        if candidate.exists():
            return candidate
    return PROJECT_ROOT / path


def repo_rel(path: str | Path) -> str:
    resolved = resolve_local_path(path)
    try:
        return str(resolved.resolve(strict=False).relative_to(PROJECT_ROOT))
    except ValueError:
        return str(resolved)


def is_rebuild_reference_pack_operation(operation: dict[str, Any]) -> bool:
    return str(operation.get("operation_kind", "")).strip() == OPERATION_KIND


def _path_error(label: str, value: str) -> str:
    if value.startswith(("http://", "https://")):
        return f"{label} must be a local path"
    if ".env" in Path(value).parts:
        return f"{label} must not point at a secret env file"
    return ""


def _params(operation: dict[str, Any]) -> dict[str, Any]:
    return operation.get("params") if isinstance(operation.get("params"), dict) else {}


def _reference_file_records(params: dict[str, Any]) -> list[dict[str, str]]:
    records: list[dict[str, str]] = []

    def add(path_text: Any, role: str) -> None:
        text = str(path_text or "").strip()
        if text:
            records.append({"path": text, "role": normalize_source_role(role)})

    for item in params.get("reference_files") or params.get("reference_file_paths") or []:
        if isinstance(item, dict):
            add(item.get("path") or item.get("local_path") or item.get("file_path"), str(item.get("role") or ""))
        else:
            add(item, "")

    explicit_keys = {
        "build_report_path": "timeline_build_report",
        "timeline_report_path": "timeline_build_report",
        "calendar_report_path": "calendar_grounding_report",
        "calendar_grounding_report_path": "calendar_grounding_report",
        "timeline_reference_path": "timeline_reference",
        "calendar_reference_path": "calendar_reference",
        "source_calendar_reference_path": "source_calendar_reference",
        "target_calendar_reference_path": "target_calendar_reference",
        "syllabus_reference_path": "syllabus_reference",
        "instructions_path": "refresh_instructions",
    }
    for key, role in explicit_keys.items():
        add(params.get(key), role)

    seen: set[str] = set()
    unique: list[dict[str, str]] = []
    for record in records:
        key = record["path"]
        if key in seen:
            continue
        seen.add(key)
        unique.append(record)
    return unique


def class_weekdays_from_params(params: dict[str, Any]) -> list[str]:
    raw = params.get("class_weekdays") or params.get("teaching_weekdays") or []
    if isinstance(raw, str):
        raw = [item for item in re.split(r"[,;/\s]+", raw) if item]
    if not isinstance(raw, (list, tuple)):
        return []
    weekdays: list[str] = []
    for item in raw:
        value = str(item or "").strip().title()
        if value in WEEKDAY_NAMES and value not in weekdays:
            weekdays.append(value)
    return sorted(weekdays, key=WEEKDAY_NAMES.index)


def validate_rebuild_reference_pack_shape(operation: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    params = _params(operation)
    forbidden = sorted(FORBIDDEN_JOB_KEYS & set(params))
    if forbidden:
        errors.append("rebuild_reference_pack.params contains forbidden executable fields: " + ", ".join(forbidden))
    records = _reference_file_records(params)
    if not records:
        errors.append("rebuild_reference_pack requires params.reference_files or explicit reference paths")
    raw_weekdays = params.get("class_weekdays") or params.get("teaching_weekdays")
    class_weekdays = class_weekdays_from_params(params)
    if not raw_weekdays:
        errors.append(
            "rebuild_reference_pack requires params.class_weekdays with the staff-provided recurring class weekdays"
        )
    elif not class_weekdays:
        errors.append(
            "rebuild_reference_pack.params.class_weekdays must contain canonical weekday names such as Tuesday"
        )
    elif isinstance(raw_weekdays, (list, tuple)) and len(class_weekdays) != len(
        {str(item or "").strip().title() for item in raw_weekdays if str(item or "").strip()}
    ):
        errors.append(
            "rebuild_reference_pack.params.class_weekdays contains unsupported weekday values"
        )
    for index, record in enumerate(records, start=1):
        path_text = record["path"]
        label = f"rebuild_reference_pack.params.reference_files[{index}]"
        role = normalize_source_role(record.get("role", ""))
        if not role:
            errors.append(
                f"{label} requires a planner-selected role; use a reference_files object "
                "with role chosen from the file content, not a bare path or filename guess"
            )
        elif role not in REFERENCE_SOURCE_ROLES:
            errors.append(
                f"{label}.role is unsupported: {role}; allowed roles are "
                + ", ".join(sorted(REFERENCE_SOURCE_ROLES))
            )
        path_error = _path_error(label, path_text)
        if path_error:
            errors.append(path_error)
            continue
        path = resolve_local_path(path_text)
        if not path.exists():
            errors.append(f"{label} does not exist: {path_text}")
    target_path = str(params.get("reference_pack_path") or "").strip()
    if target_path:
        path_error = _path_error("rebuild_reference_pack.params.reference_pack_path", target_path)
        if path_error:
            errors.append(path_error)
    return errors


def reference_pack_path(params: dict[str, Any], env: dict[str, str]) -> Path:
    course_code = safe_name(str(params.get("course_code") or env.get("COURSE_AGENT_DEFAULT_COURSE_CODE") or "DASE7137"))
    default_semester = str(env.get("COURSE_AGENT_DEFAULT_SEMESTER") or DEFAULT_REFERENCE_PACK_SEMESTER)
    semester = canonical_reference_pack_semester(
        str(params.get("semester") or default_semester),
        default_semester=default_semester,
    )
    canonical_target = REFERENCE_ROOT / course_code / semester
    explicit = str(params.get("reference_pack_path") or env.get("COURSE_REFERENCE_PACK_PATH") or "").strip()
    if explicit:
        resolved = resolve_local_path(explicit)
        try:
            resolved.resolve(strict=False).relative_to(REFERENCE_ROOT.resolve(strict=False))
        except ValueError:
            return resolved
    return canonical_target


def canonical_reference_pack_semester(
    value: str,
    *,
    default_semester: str = DEFAULT_REFERENCE_PACK_SEMESTER,
) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    default_text = str(default_semester or DEFAULT_REFERENCE_PACK_SEMESTER).strip()
    default_match = re.search(r"\b(20\d{2})\s*[-/]\s*(20\d{2})\b", default_text)
    default_start_year = int(default_match.group(1)) if default_match else 2026
    default_term = "sem2" if _semester_term(default_text) == "sem2" else "sem1"
    term = _semester_term(text) or default_term
    match = re.search(r"\b(20\d{2})\s*[-/]\s*(20\d{2})\b", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}-{term}"
    if (
        "new academic" in lowered
        or "academic-year" in lowered
        or "新学年" in text
    ):
        start_year = target_start_year(text, default=default_start_year)
        return f"{start_year}-{start_year + 1}-{term}"
    if _semester_term(text):
        return f"{default_start_year}-{default_start_year + 1}-{term}"
    return safe_name(text or "reference-pack")


def _semester_term(value: str) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    compact = re.sub(r"[^a-z0-9]+", "", lowered)
    if (
        "sem2" in compact
        or "semester2" in compact
        or "secondsemester" in compact
        or "第二学期" in text
    ):
        return "sem2"
    if (
        "sem1" in compact
        or "semester1" in compact
        or "firstsemester" in compact
        or "第一学期" in text
    ):
        return "sem1"
    return ""


def load_json(path: Path) -> dict[str, Any] | None:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError):
        return None
    return value if isinstance(value, dict) else None


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def copy_to_sources(path: Path, sources_root: Path) -> Path:
    sources_root.mkdir(parents=True, exist_ok=True)
    target = sources_root / safe_name(path.name, "reference")
    if target.exists() and target.resolve(strict=False) != path.resolve(strict=False):
        stem = target.stem
        suffix = target.suffix
        counter = 2
        while target.exists():
            target = sources_root / f"{stem}-{counter}{suffix}"
            counter += 1
    if target.resolve(strict=False) != path.resolve(strict=False):
        shutil.copy2(path, target)
    return target


def normalize_source_role(value: str) -> str:
    role = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return REFERENCE_SOURCE_ROLE_ALIASES.get(role, role)


def classify_source(_path: Path, declared_role: str) -> str:
    role = normalize_source_role(declared_role)
    return role if role in REFERENCE_SOURCE_ROLES else "reference"


def detect_reference_container(path: Path) -> str:
    try:
        with path.open("rb") as handle:
            prefix = handle.read(4096)
    except OSError:
        return "unreadable"
    if prefix.startswith(b"%PDF-"):
        return "pdf"
    if zipfile.is_zipfile(path):
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
        except (OSError, zipfile.BadZipFile):
            return "unreadable"
        if "word/document.xml" in names:
            return "word"
        if "xl/workbook.xml" in names:
            return "workbook"
        return "archive"
    try:
        text = prefix.decode("utf-8-sig")
    except UnicodeDecodeError:
        return "binary"
    stripped = text.lstrip()
    if stripped.startswith(("{", "[")):
        try:
            json.loads(path.read_text(encoding="utf-8-sig"))
            return "json"
        except (OSError, UnicodeDecodeError, json.JSONDecodeError):
            pass
    return "text"


def source_container_error(path: Path) -> str:
    container = detect_reference_container(path)
    if container == "unreadable":
        return "attachment content could not be read"
    if container == "archive":
        return "attachment is an archive without a supported document structure"
    if container == "binary":
        return "attachment is an unsupported binary container"
    return ""


def target_start_year(semester: str, default: int = 2026) -> int:
    match = re.search(r"\b(20\d{2})\s*[-/]\s*20\d{2}\b", semester)
    if match:
        return int(match.group(1))
    match = re.search(r"\b(20\d{2})\b", semester)
    if match:
        return int(match.group(1))
    return default


def iso_date(value: Any, year: int) -> str:
    text = str(value or "").strip()
    if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", text):
        return text
    for fmt in ("%b %d %Y", "%B %d %Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(f"{text} {year}", fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def clean_cell_text(value: Any) -> str:
    text = str(value or "").replace("\n", " ")
    text = text.replace("（", "(").replace("）", ")")
    text = re.sub(r"\s*/\s*", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" /")


NON_TEACHING_TIMELINE_MARKERS = (
    "holiday",
    "reading week",
    "reading/field trip",
    "revision period",
    "assessment period",
    "class suspension",
)


def timeline_row_counts_as_teaching(row: dict[str, Any]) -> bool:
    explicit = row.get("counts_as_teaching_week")
    if explicit is not None:
        return bool(explicit)
    marker = " ".join(
        clean_cell_text(row.get(key)).lower()
        for key in ("lecture", "workshop", "tutorial", "tutorial_workshop", "notes")
    )
    return not any(token in marker for token in NON_TEACHING_TIMELINE_MARKERS)


def parse_iso(value: str) -> datetime:
    return datetime.strptime(value, "%Y-%m-%d")


def month_date(month: str, day: str | int, year: str | int) -> str:
    parsed = datetime.strptime(f"{str(month).title()} {int(day)} {int(year)}", "%b %d %Y")
    return parsed.strftime("%Y-%m-%d")


def parse_single_dates(text: str, label: str) -> list[str]:
    pattern = re.compile(
        rf"{re.escape(label)}:\s*([A-Za-z]{{3}})\s+(\d{{1,2}}),\s*(20\d{{2}})",
        re.IGNORECASE,
    )
    return [month_date(month, day, year) for month, day, year in pattern.findall(text)]


def parse_date_ranges(text: str, label: str) -> list[tuple[str, str]]:
    months = "Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec"
    pattern = re.compile(
        rf"{re.escape(label)}:.*?\b({months})\s+(\d{{1,2}})\s*-\s*(?:({months})\s+)?(\d{{1,2}}),\s*(20\d{{2}})",
        re.IGNORECASE | re.DOTALL,
    )
    ranges: list[tuple[str, str]] = []
    for start_month, start_day, end_month, end_day, year in pattern.findall(text):
        ranges.append((month_date(start_month, start_day, year), month_date(end_month or start_month, end_day, year)))
    return ranges


def parse_semester_range(text: str, label: str) -> tuple[str, str] | None:
    pattern = re.compile(
        rf"{label}\s+SEMESTER:\s*([A-Za-z]{{3}})\s+(\d{{1,2}})\s*-\s*([A-Za-z]{{3}})\s+(\d{{1,2}}),\s*(20\d{{2}})",
        re.IGNORECASE,
    )
    match = pattern.search(text)
    if not match:
        return None
    start_month, start_day, end_month, end_day, year = match.groups()
    return month_date(start_month, start_day, year), month_date(end_month, end_day, year)


MONTH_INDEX = {
    "JAN": 1,
    "FEB": 2,
    "MAR": 3,
    "APR": 4,
    "MAY": 5,
    "JUN": 6,
    "JUL": 7,
    "AUG": 8,
    "SEP": 9,
    "OCT": 10,
    "NOV": 11,
    "DEC": 12,
}
HOLIDAY_MARKERS = {
    "[": {
        "type": "general_holiday",
        "name": "General Holiday",
        "classes_suspended": True,
        "full_day": True,
    },
    "(": {
        "type": "university_holiday",
        "name": "University Holiday",
        "classes_suspended": True,
        "full_day": True,
    },
    "<": {
        "type": "university_holiday_afternoon",
        "name": "University Holiday (afternoon only)",
        "classes_suspended": False,
        "full_day": False,
    },
}


def next_month(month: int, year: int) -> tuple[int, int]:
    if month == 12:
        return 1, year + 1
    return month + 1, year


def marked_calendar_days_from_text(text: str) -> list[dict[str, Any]]:
    current_month: int | None = None
    current_year: int | None = None
    last_day: int | None = None
    holidays: dict[str, dict[str, Any]] = {}
    token_pattern = re.compile(r"^(?P<open>[\[\(<])?(?P<day>\d{1,2})(?P<close>[\]\)>])?$")
    month_pattern = re.compile(r"^(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)-(\d{2})$", re.IGNORECASE)
    stop_words = {
        "FIRST",
        "SECOND",
        "OPTIONAL",
        "First",
        "Last",
        "Reading/Field",
        "Revision",
        "Assessment",
        "Class",
        "Break",
        "Notes:",
        "THE",
        "Calendar",
        "SUN",
    }
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("[ ]") or line.startswith("( )") or line.startswith("< >"):
            break
        tokens = line.split()
        for index, token in enumerate(tokens):
            month_match = month_pattern.match(token)
            if month_match:
                current_month = MONTH_INDEX[month_match.group(1).upper()]
                current_year = 2000 + int(month_match.group(2))
                last_day = None
                continue
            if any(token.startswith(word) for word in stop_words):
                break
            match = token_pattern.match(token)
            if not match or current_month is None or current_year is None:
                continue
            marker = match.group("open") or ""
            day = int(match.group("day"))
            if day > 31:
                continue
            if last_day is not None and day < last_day:
                if marker or day == 1 or index + 1 < len(tokens):
                    current_month, current_year = next_month(current_month, current_year)
                else:
                    continue
            last_day = day
            if marker not in HOLIDAY_MARKERS:
                continue
            try:
                day_date = datetime(current_year, current_month, day).strftime("%Y-%m-%d")
            except ValueError:
                continue
            basis = HOLIDAY_MARKERS[marker]
            close_marker = {"[": "]", "(": ")", "<": ">"}[marker]
            holidays[day_date] = {
                "date": day_date,
                "name": basis["name"],
                "type": basis["type"],
                "classes_suspended": basis["classes_suspended"],
                "full_day": basis["full_day"],
                "source_marker": f"{marker}{day}{close_marker}",
            }
    return [holidays[key] for key in sorted(holidays)]


def period_record(
    name: str,
    period_type: str,
    start: str | None,
    end: str | None,
    *,
    classes_suspended: bool,
    counts_as_teaching_week: bool = True,
) -> dict[str, Any] | None:
    if not start or not end:
        return None
    return {
        "name": name,
        "type": period_type,
        "start": start,
        "end": end,
        "classes_suspended": classes_suspended,
        "counts_as_teaching_week": counts_as_teaching_week,
    }


def extract_pdf_text(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber is required to extract calendar PDF text") from exc
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def build_calendar_from_text(
    text: str,
    *,
    course_code: str,
    source_basis: Path,
    source_docs: list[str],
) -> dict[str, Any] | None:
    normalized = re.sub(r"\s+", " ", text)
    label_match = re.search(r"Academic Year\s+(20\d{2}\s*-\s*20\d{2})", normalized, re.IGNORECASE)
    first_range = parse_semester_range(normalized, "FIRST")
    second_range = parse_semester_range(normalized, "SECOND")
    first_teaching = parse_single_dates(normalized, "First Day of Teaching")
    last_teaching = parse_single_dates(normalized, "Last Day of Teaching")
    reading = parse_date_ranges(normalized, "Reading/Field Trip Week")
    revision = parse_date_ranges(normalized, "Revision Period")
    assessment = parse_date_ranges(normalized, "Assessment Period")
    suspension = parse_date_ranges(normalized, "Class Suspension Period for the Lunar New Year")
    if not (first_range and second_range and len(first_teaching) >= 2 and len(last_teaching) >= 2):
        return None
    if len(reading) < 2 or len(revision) < 2 or len(assessment) < 2:
        return None
    non_teaching_days = marked_calendar_days_from_text(text)
    teaching_suspensions = [
        item
        for item in (
            period_record(
                "Reading/Field Trip Week",
                "reading_or_field_trip_week",
                reading[0][0],
                reading[0][1],
                classes_suspended=True,
                counts_as_teaching_week=False,
            ),
            period_record("Revision Period", "revision_period", revision[0][0], revision[0][1], classes_suspended=True, counts_as_teaching_week=False),
            period_record("Assessment Period", "assessment_period", assessment[0][0], assessment[0][1], classes_suspended=True, counts_as_teaching_week=False),
            period_record(
                "Class Suspension Period for the Lunar New Year",
                "class_suspension",
                suspension[0][0] if suspension else None,
                suspension[0][1] if suspension else None,
                classes_suspended=True,
                counts_as_teaching_week=False,
            ),
            period_record(
                "Reading/Field Trip Week",
                "reading_or_field_trip_week",
                reading[1][0],
                reading[1][1],
                classes_suspended=True,
                counts_as_teaching_week=False,
            ),
            period_record("Revision Period", "revision_period", revision[1][0], revision[1][1], classes_suspended=True, counts_as_teaching_week=False),
            period_record("Assessment Period", "assessment_period", assessment[1][0], assessment[1][1], classes_suspended=True, counts_as_teaching_week=False),
        )
        if item
    ]
    return {
        "course_code": course_code,
        "academic_year": (label_match.group(1).replace(" ", "") if label_match else ""),
        "target_semester": "first",
        "status": "source-confirmed",
        "source_basis": [repo_rel(source_basis)],
        "source_documents": source_docs,
        "non_teaching_days": non_teaching_days,
        "teaching_suspensions": teaching_suspensions,
        "semesters": {
            "first": {
                "name": "first",
                "start": first_range[0],
                "end": first_range[1],
                "first_teaching": first_teaching[0],
                "last_teaching": last_teaching[0],
                "reading_start": reading[0][0],
                "reading_end": reading[0][1],
                "revision_start": revision[0][0],
                "revision_end": revision[0][1],
                "assessment_start": assessment[0][0],
                "assessment_end": assessment[0][1],
                "class_suspension_start": None,
                "class_suspension_end": None,
            },
            "second": {
                "name": "second",
                "start": second_range[0],
                "end": second_range[1],
                "first_teaching": first_teaching[1],
                "last_teaching": last_teaching[1],
                "reading_start": reading[1][0],
                "reading_end": reading[1][1],
                "revision_start": revision[1][0],
                "revision_end": revision[1][1],
                "assessment_start": assessment[1][0],
                "assessment_end": assessment[1][1],
                "class_suspension_start": suspension[0][0] if suspension else None,
                "class_suspension_end": suspension[0][1] if suspension else None,
            },
        },
        "refresh_policy": {
            "unmapped_dates": "needs_review",
        },
    }


def build_calendar_from_pdf(path: Path, *, course_code: str, source_docs: list[str]) -> dict[str, Any] | None:
    return build_calendar_from_text(
        extract_pdf_text(path),
        course_code=course_code,
        source_basis=path,
        source_docs=source_docs,
    )


def reference_visible_text(path: Path) -> str:
    container = detect_reference_container(path)
    if container == "pdf":
        return extract_pdf_text(path)
    if container == "word":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to read Word references") from exc
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            "\t".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return "\n".join(parts)
    if container == "workbook":
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to read workbook references") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return "\n".join(
                "\t".join("" if value is None else str(value) for value in row)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
            )
        finally:
            workbook.close()
    if container == "json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        return json.dumps(payload, ensure_ascii=False)
    if container == "text":
        return path.read_text(encoding="utf-8-sig", errors="replace")
    return ""


def build_calendar_from_reference(
    path: Path,
    *,
    course_code: str,
    source_docs: list[str],
) -> dict[str, Any] | None:
    if detect_reference_container(path) == "json":
        payload = load_json(path) or {}
        if isinstance(payload.get("target_calendar"), dict):
            return build_calendar(path, course_code=course_code, source_docs=source_docs)
        if payload.get("semesters"):
            return normalize_existing_calendar(
                path,
                course_code=course_code,
                source_docs=source_docs,
            )
    text = reference_visible_text(path)
    if not text:
        return None
    return build_calendar_from_text(
        text,
        course_code=course_code,
        source_basis=path,
        source_docs=source_docs,
    )


def calendar_payload_matches_target(calendar_payload: dict[str, Any], semester: str) -> bool:
    target_year = str(target_start_year(semester))
    next_year = str(int(target_year) + 1)
    academic_year = str(calendar_payload.get("academic_year") or "")
    if target_year in academic_year and next_year in academic_year:
        return True
    facts = first_semester_facts(calendar_payload)
    if not facts:
        return False
    first_teaching = str(facts.get("first_teaching") or facts.get("start") or "")
    return first_teaching.startswith(f"{target_year}-")


def calendar_target_mismatch_message(calendar_payload: dict[str, Any], semester: str) -> str:
    target_year = target_start_year(semester)
    academic_year = str(calendar_payload.get("academic_year") or "unknown academic year")
    return (
        f"target university calendar PDF for {target_year}-{target_year + 1} "
        f"(uploaded calendar resolved to {academic_year})"
    )


def first_semester_facts(calendar_payload: dict[str, Any]) -> dict[str, Any] | None:
    semesters = calendar_payload.get("semesters")
    if isinstance(semesters, dict) and isinstance(semesters.get("first"), dict):
        return semesters["first"]
    if isinstance(semesters, list):
        for item in semesters:
            name = str(item.get("name", "")).strip().lower() if isinstance(item, dict) else ""
            if isinstance(item, dict) and name in {"first", "first semester", "semester 1", "sem 1"}:
                return item
    return None


def calendar_semester_records(
    calendar_payload: dict[str, Any],
) -> list[tuple[str, dict[str, Any]]]:
    semesters = calendar_payload.get("semesters")
    records: list[tuple[str, dict[str, Any]]] = []
    if isinstance(semesters, dict):
        for key, item in semesters.items():
            if isinstance(item, dict):
                records.append((str(key), item))
        return records
    if isinstance(semesters, list):
        for index, item in enumerate(semesters, start=1):
            if not isinstance(item, dict):
                continue
            name = str(item.get("name") or f"semester-{index}")
            records.append((name, item))
    return records


def calendar_semester_for_date(
    calendar_payload: dict[str, Any],
    day: datetime,
) -> tuple[str, dict[str, Any]] | None:
    for name, facts in calendar_semester_records(calendar_payload):
        start_text = str(facts.get("start") or facts.get("first_teaching") or "")
        end_text = str(facts.get("end") or facts.get("assessment_end") or facts.get("last_teaching") or "")
        if not start_text or not end_text:
            continue
        try:
            if parse_iso(start_text) <= day <= parse_iso(end_text):
                return name, facts
        except ValueError:
            continue
    return None


def calendar_date_from_source_text(
    value: Any,
    calendar_payload: dict[str, Any],
) -> tuple[datetime, str, dict[str, Any]] | None:
    text = clean_cell_text(value)
    if not text:
        return None
    if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", text):
        try:
            parsed = parse_iso(text)
        except ValueError:
            return None
        semester_record = calendar_semester_for_date(calendar_payload, parsed)
        if semester_record:
            return parsed, semester_record[0], semester_record[1]
        return None

    normalized = re.sub(r"(?i)(\d)(st|nd|rd|th)\b", r"\1", text)
    normalized = normalized.replace(",", " ")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    candidate_years: set[int] = set()
    academic_year = str(calendar_payload.get("academic_year") or "")
    for match in re.finditer(r"\b20\d{2}\b", academic_year):
        candidate_years.add(int(match.group(0)))
    for _name, facts in calendar_semester_records(calendar_payload):
        for key in ("start", "end", "first_teaching", "last_teaching"):
            match = re.match(r"(20\d{2})-", str(facts.get(key) or ""))
            if match:
                candidate_years.add(int(match.group(1)))

    candidates: list[tuple[datetime, str, dict[str, Any]]] = []
    for year in sorted(candidate_years):
        iso_value = iso_date(normalized, year)
        if not re.fullmatch(r"20\d{2}-\d{2}-\d{2}", iso_value):
            continue
        parsed = parse_iso(iso_value)
        semester_record = calendar_semester_for_date(calendar_payload, parsed)
        if semester_record:
            candidates.append((parsed, semester_record[0], semester_record[1]))
    unique = {
        (item[0].strftime("%Y-%m-%d"), item[1]): item
        for item in candidates
    }
    return next(iter(unique.values())) if len(unique) == 1 else None


def source_calendar_week_number(
    semester_facts: dict[str, Any],
    day: datetime,
) -> int | None:
    first_text = str(
        semester_facts.get("first_teaching")
        or semester_facts.get("start")
        or ""
    )
    try:
        first = parse_iso(first_text)
    except ValueError:
        return None
    if day < first:
        return None
    first_week_start = first - timedelta(days=first.weekday())
    day_week_start = day - timedelta(days=day.weekday())
    return ((day_week_start.date() - first_week_start.date()).days // 7) + 1


def source_date_context(
    value: Any,
    source_calendar_payload: dict[str, Any] | None,
) -> dict[str, Any]:
    if not source_calendar_payload:
        return {}
    resolved = calendar_date_from_source_text(value, source_calendar_payload)
    if not resolved:
        return {}
    parsed, semester_name, semester_facts = resolved
    return {
        "source_date_iso": parsed.strftime("%Y-%m-%d"),
        "source_weekday": WEEKDAY_NAMES[parsed.weekday()],
        "source_calendar_week": source_calendar_week_number(semester_facts, parsed),
        "source_semester": semester_name,
        "source_academic_year": str(source_calendar_payload.get("academic_year") or ""),
        "source_calendar_status": (
            calendar_anchor_non_teaching_reason(source_calendar_payload, parsed)
            or "teaching_date"
        ),
    }


def timeline_rows_from_matrix(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    if not matrix:
        return []
    headers = [clean_cell_text(value).lower() for value in matrix[0]]

    def column(*markers: str) -> int | None:
        for index, header in enumerate(headers):
            if any(marker in header for marker in markers):
                return index
        return None

    week_index = column("week")
    lecture_index = column("lecture", "theme")
    if week_index is None or lecture_index is None:
        return []
    date_index = column("date")
    workshop_index = column("workshop", "tutorial", "lab")
    submission_index = column("submission", "deadline", "assessment")

    def value(cells: list[Any], index: int | None) -> str:
        return clean_cell_text(cells[index]) if index is not None and index < len(cells) else ""

    rows: list[dict[str, Any]] = []
    for cells in matrix[1:]:
        week_text = value(cells, week_index)
        match = re.fullmatch(r"\s*(\d+)\s*", week_text)
        if not match:
            continue
        rows.append(
            {
                "week": int(match.group(1)),
                "date": value(cells, date_index),
                "lecture": value(cells, lecture_index),
                "workshop": value(cells, workshop_index),
                "submission": value(cells, submission_index),
            }
        )
    return rows


def extract_timeline_rows_from_reference(path: Path) -> list[dict[str, Any]]:
    container = detect_reference_container(path)
    if container == "word":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to read Word reference tables") from exc
        document = Document(path)
        for table in document.tables:
            rows = timeline_rows_from_matrix(
                [[cell.text for cell in row.cells] for row in table.rows]
            )
            if rows:
                return rows
        return []
    if container == "workbook":
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to read workbook reference tables") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = timeline_rows_from_matrix(
                    [list(row) for row in sheet.iter_rows(values_only=True)]
                )
                if rows:
                    return rows
        finally:
            workbook.close()
        return []
    if container == "pdf":
        try:
            import pdfplumber
        except ImportError as exc:
            raise RuntimeError("pdfplumber is required to read PDF reference tables") from exc
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables() or []:
                    rows = timeline_rows_from_matrix(
                        [[value or "" for value in row] for row in table or []]
                    )
                    if rows:
                        return rows
        return []
    if container == "json":
        payload = load_json(path) or {}
        raw_rows = payload.get("timeline_rows") or payload.get("week_calendar") or []
        if isinstance(raw_rows, list):
            return [
                {
                    "week": int(row.get("week") or row.get("teaching_week") or 0),
                    "date": clean_cell_text(row.get("date")),
                    "lecture": clean_cell_text(row.get("lecture")),
                    "workshop": clean_cell_text(
                        row.get("workshop") or row.get("tutorial") or row.get("lab")
                    ),
                    "submission": clean_cell_text(
                        row.get("submission") or row.get("deadline")
                    ),
                }
                for row in raw_rows
                if isinstance(row, dict)
                and str(row.get("week") or row.get("teaching_week") or "").isdigit()
            ]
        return []
    if container == "text":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        for delimiter in (",", "\t", "|", ";"):
            matrix = [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
            rows = timeline_rows_from_matrix(matrix)
            if rows:
                return rows
    return []


def extract_timeline_rows_from_docx(path: Path) -> list[dict[str, Any]]:
    return extract_timeline_rows_from_reference(path)


def date_range_contains(start: str | None, end: str | None, day: datetime) -> bool:
    if not start or not end:
        return False
    try:
        return parse_iso(start) <= day <= parse_iso(end)
    except ValueError:
        return False


def calendar_anchor_non_teaching_reason(calendar_payload: dict[str, Any], anchor: datetime) -> str:
    for item in calendar_payload.get("teaching_suspensions") if isinstance(calendar_payload.get("teaching_suspensions"), list) else []:
        if not isinstance(item, dict):
            continue
        counts = item.get("counts_as_teaching_week")
        should_exclude = counts is False or bool(item.get("classes_suspended")) or str(item.get("type") or "") == "reading_or_field_trip_week"
        if should_exclude and date_range_contains(str(item.get("start") or ""), str(item.get("end") or ""), anchor):
            return str(item.get("name") or item.get("type") or "calendar non-teaching period")
    facts = first_semester_facts(calendar_payload) or {}
    legacy_periods = (
        ("Reading/Field Trip Week", facts.get("reading_start"), facts.get("reading_end")),
        ("Revision Period", facts.get("revision_start"), facts.get("revision_end")),
        ("Assessment Period", facts.get("assessment_start"), facts.get("assessment_end")),
        ("Class Suspension Period", facts.get("class_suspension_start"), facts.get("class_suspension_end")),
    )
    for label, start, end in legacy_periods:
        if date_range_contains(str(start or ""), str(end or ""), anchor):
            return label
    for item in calendar_payload.get("non_teaching_days") if isinstance(calendar_payload.get("non_teaching_days"), list) else []:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        if str(item.get("date") or "") == anchor.strftime("%Y-%m-%d"):
            return str(item.get("name") or item.get("type") or "calendar non-teaching day")
    return ""


def target_calendar_week_number(calendar_payload: dict[str, Any], day: datetime) -> int | None:
    facts = first_semester_facts(calendar_payload)
    if not facts:
        return None
    first_teaching = str(facts.get("first_teaching") or facts.get("start") or "")
    try:
        first = parse_iso(first_teaching)
    except ValueError:
        return None
    if day < first:
        return None
    return ((day.date() - first.date()).days // 7) + 1


def target_teaching_week_anchors(calendar_payload: dict[str, Any]) -> tuple[list[datetime], list[dict[str, Any]]]:
    facts = first_semester_facts(calendar_payload)
    if not facts:
        return [], []
    first_teaching = str(facts.get("first_teaching") or facts.get("start") or "")
    last_teaching = str(facts.get("last_teaching") or facts.get("revision_start") or facts.get("end") or "")
    if not first_teaching or not last_teaching:
        return [], []
    try:
        anchor = parse_iso(first_teaching)
        end = parse_iso(last_teaching)
    except ValueError:
        return [], []
    anchors: list[datetime] = []
    gaps: list[dict[str, Any]] = []
    while anchor <= end:
        reason = calendar_anchor_non_teaching_reason(calendar_payload, anchor)
        calendar_week = target_calendar_week_number(calendar_payload, anchor)
        if reason:
            gaps.append(
                {
                    "date": anchor.strftime("%Y-%m-%d"),
                    "calendar_week": calendar_week,
                    "reason": reason,
                    "week_numbering_effect": "does_not_consume_teaching_week",
                }
            )
        else:
            anchors.append(anchor)
        anchor += timedelta(days=7)
    return anchors, gaps


def build_timeline_from_reference(
    path: Path,
    *,
    course_code: str,
    semester: str,
    calendar_payload: dict[str, Any],
    source_calendar_payload: dict[str, Any] | None = None,
    source_docs: list[str],
) -> dict[str, Any] | None:
    facts = first_semester_facts(calendar_payload)
    if not facts:
        return None
    raw_rows = extract_timeline_rows_from_docx(path)
    if not raw_rows:
        return None
    content_rows = []
    excluded_source_rows: list[dict[str, Any]] = []
    source_calendar_overrides: list[dict[str, Any]] = []
    source_calendar_unresolved_rows: list[dict[str, Any]] = []
    for row in raw_rows:
        lecture = clean_cell_text(row.get("lecture"))
        workshop = clean_cell_text(row.get("workshop"))
        source_week = int(row.get("week") or 0)
        source_date = clean_cell_text(row.get("date"))
        normalized_row = {
            "week": source_week,
            "date": source_date,
            "lecture": lecture,
            "workshop": workshop,
            "submission": clean_cell_text(row.get("submission")),
        }
        date_context = source_date_context(source_date, source_calendar_payload)
        text_counts_as_teaching = timeline_row_counts_as_teaching(normalized_row)
        calendar_status = str(date_context.get("source_calendar_status") or "")
        calendar_blocks_teaching = bool(
            calendar_status and calendar_status != "teaching_date"
        )
        if source_calendar_payload and source_date and not date_context:
            source_calendar_unresolved_rows.append(
                {
                    "source_week": source_week,
                    "source_date": source_date,
                    "reason": (
                        "source date could not be resolved uniquely inside the declared "
                        "source academic calendar; Timeline ordering remains authoritative"
                    ),
                }
            )
        if text_counts_as_teaching and calendar_blocks_teaching:
            source_calendar_overrides.append(
                {
                    "source_week": source_week,
                    "source_date": source_date,
                    "source_date_iso": date_context.get("source_date_iso"),
                    "lecture": lecture,
                    "workshop": workshop,
                    "calendar_reason": calendar_status,
                    "resolution": "timeline_overrides_source_calendar",
                    "reason": "Timeline explicitly schedules teaching content on this source-calendar suspension.",
                }
            )
        if not text_counts_as_teaching:
            excluded_source_rows.append(
                {
                    "source_week": source_week,
                    "source_date": source_date,
                    "lecture": lecture,
                    "workshop": workshop,
                    **date_context,
                    "reason": (
                        f"source calendar marks this row as {calendar_status}"
                        if calendar_blocks_teaching
                        else "source row is non-teaching and does not consume a teaching week"
                    ),
                }
            )
            continue
        source_teaching_week = len(content_rows) + 1
        content_rows.append(
            {
                "source_week": source_week,
                "source_visible_week": source_week,
                "source_teaching_week": source_teaching_week,
                "source_date": source_date,
                **date_context,
                "lecture": lecture,
                "workshop": workshop,
                "submission": clean_cell_text(row.get("submission")),
            }
        )
    anchors, calendar_gaps = target_teaching_week_anchors(calendar_payload)
    rows: list[dict[str, Any]] = []
    for index, item in enumerate(content_rows, start=1):
        week_date = anchors[index - 1] if index <= len(anchors) else None
        rows.append(
            {
                "week": index,
                "teaching_week": index,
                "calendar_week": (
                    target_calendar_week_number(calendar_payload, week_date)
                    if week_date
                    else None
                ),
                "date": week_date.strftime("%Y-%m-%d") if week_date else "",
                "lecture": item["lecture"],
                "workshop": item["workshop"],
                "submission": item["submission"],
                "source_week": item["source_week"],
                "source_visible_week": item["source_visible_week"],
                "source_teaching_week": item["source_teaching_week"],
                "source_date": item["source_date"],
                "source_date_iso": item.get("source_date_iso", ""),
                "source_weekday": item.get("source_weekday", ""),
                "source_calendar_week": item.get("source_calendar_week"),
                "source_semester": item.get("source_semester", ""),
                "source_academic_year": item.get("source_academic_year", ""),
                "counts_as_teaching_week": True,
            }
        )
    return {
        "course_code": course_code,
        "semester": semester,
        "status": "source-confirmed",
        "source_basis": [repo_rel(path), *calendar_payload.get("source_basis", [])],
        "source_documents": source_docs,
        "source_academic_year": (
            str(source_calendar_payload.get("academic_year") or "")
            if source_calendar_payload
            else ""
        ),
        "source_calendar_basis": (
            list(source_calendar_payload.get("source_basis") or [])
            if source_calendar_payload
            else []
        ),
        "source_week_numbering": {
            "visible_week_field": "source_visible_week",
            "calendar_week_field": "source_calendar_week",
            "teaching_week_field": "source_teaching_week",
            "weekday_field": "source_weekday",
            "rule": (
                "Source calendar weeks and weekdays are resolved from source_calendar.json; "
                "source teaching weeks exclude source rows marked non-teaching or blocked by "
                "the source calendar."
            ),
        },
        "week_numbering": {
            "mode": "dual",
            "mapping_field": "teaching_week",
            "display_field": "date_planner_selected",
            "available_display_fields": ["calendar_week", "teaching_week"],
            "legacy_week_field": "teaching_week",
            "rule": (
                "Mappings use teaching_week. File refresh Date Planner selects calendar_week or teaching_week "
                "for each visible Week column from the source table semantics."
            ),
            "source_non_teaching_rows_excluded": True,
            "target_non_teaching_calendar_gaps_excluded": True,
        },
        "week_count": len(rows),
        "teaching_week_count": len(rows),
        "calendar_week_count": max(
            [
                int(item.get("calendar_week") or 0)
                for item in [*rows, *calendar_gaps]
                if isinstance(item, dict)
            ]
            or [0]
        ),
        "week_calendar": rows,
        "excluded_source_rows": excluded_source_rows,
        "source_calendar_precedence": "timeline_overrides_source_calendar",
        "source_calendar_overrides": source_calendar_overrides,
        "source_calendar_unresolved_rows": source_calendar_unresolved_rows,
        "non_teaching_calendar_gaps": calendar_gaps,
    }


def build_timeline_from_docx(
    path: Path,
    *,
    course_code: str,
    semester: str,
    calendar_payload: dict[str, Any],
    source_calendar_payload: dict[str, Any] | None = None,
    source_docs: list[str],
) -> dict[str, Any] | None:
    return build_timeline_from_reference(
        path,
        course_code=course_code,
        semester=semester,
        calendar_payload=calendar_payload,
        source_calendar_payload=source_calendar_payload,
        source_docs=source_docs,
    )


def build_timeline(report_path: Path, *, course_code: str, semester: str, source_docs: list[str]) -> dict[str, Any]:
    report = load_json(report_path) or {}
    year = target_start_year(semester)
    rows = []
    for row in report.get("timeline_rows") or []:
        if not isinstance(row, dict):
            continue
        rows.append(
            {
                "week": int(row["week"]),
                "date": iso_date(row.get("date"), year),
                "lecture": str(row.get("lecture") or ""),
                "workshop": str(row.get("workshop") or ""),
                "submission": str(row.get("submission") or ""),
            }
        )
    return {
        "course_code": course_code,
        "semester": semester,
        "status": "source-confirmed",
        "source_basis": [repo_rel(report_path)],
        "source_documents": source_docs,
        "week_calendar": rows,
    }


def normalize_existing_timeline(path: Path, *, course_code: str, semester: str, source_docs: list[str]) -> dict[str, Any] | None:
    data = load_json(path)
    if not data or not isinstance(data.get("week_calendar"), list):
        return None
    normalized = dict(data)
    normalized["course_code"] = str(normalized.get("course_code") or course_code)
    normalized["semester"] = str(normalized.get("semester") or semester)
    normalized.setdefault("status", "source-confirmed")
    normalized["source_basis"] = [repo_rel(path)]
    if source_docs:
        normalized["source_documents"] = source_docs
    return normalized


def teaching_content_rows(timeline_payload: dict[str, Any]) -> list[dict[str, Any]]:
    rows = timeline_payload.get("week_calendar")
    if not isinstance(rows, list):
        return []
    return [
        dict(row)
        for row in rows
        if isinstance(row, dict) and timeline_row_counts_as_teaching(row)
    ]


def build_schedule_from_class_weekdays(
    timeline_payload: dict[str, Any],
    calendar_payload: dict[str, Any],
    *,
    class_weekdays: list[str],
    course_code: str,
    semester: str,
    source_basis: list[str],
) -> tuple[dict[str, Any] | None, str]:
    try:
        from semester_schedule_update import (
            base_schedule_payload,
            recurring_teaching_week_dates,
            validate_schedule_payload,
        )
    except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
        from course_agent_app.semester_schedule_update import (
            base_schedule_payload,
            recurring_teaching_week_dates,
            validate_schedule_payload,
        )

    content_rows = teaching_content_rows(timeline_payload)
    required_weeks = len(content_rows)
    if not required_weeks:
        return None, "timeline reference does not contain any teaching rows"
    reference_payloads = {
        "calendar.json": calendar_payload,
        "timeline.json": timeline_payload,
    }
    dates_by_week = recurring_teaching_week_dates(
        reference_payloads,
        class_weekdays,
        through_week=required_weeks,
    )
    missing_week = next(
        (week for week in range(1, required_weeks + 1) if not dates_by_week.get(week)),
        None,
    )
    if missing_week is not None:
        facts = first_semester_facts(calendar_payload) or {}
        last_teaching = str(
            facts.get("last_teaching")
            or facts.get("revision_start")
            or facts.get("end")
            or ""
        ).strip()
        return None, (
            f"The recurring {', '.join(class_weekdays)} arrangement needs {required_weeks} teaching weeks, "
            f"but only {len(dates_by_week)} eligible teaching weeks fit before the teaching period ends"
            + (f" ({last_teaching})" if last_teaching else "")
            + f"; Teaching Week {missing_week} has no valid class date after holidays and class-suspension periods are excluded."
        )
    weeks: list[dict[str, Any]] = []
    for teaching_week in range(1, required_weeks + 1):
        teaching_dates = dates_by_week[teaching_week]
        first_date = parse_iso(teaching_dates[0])
        weeks.append(
            {
                "week": teaching_week,
                "teaching_week": teaching_week,
                "calendar_week": target_calendar_week_number(calendar_payload, first_date),
                "label": f"Teaching Week {teaching_week}",
                "teaching_dates": teaching_dates,
                "allowed_session_weekdays": class_weekdays,
                "notes": (
                    "Generated during reference-pack rebuild from the staff-provided recurring class weekdays; "
                    "holiday and class-suspension dates do not consume teaching-week numbers."
                ),
            }
        )
    schedule = base_schedule_payload(
        course_code,
        canonical_reference_pack_semester(semester),
        source_basis,
        weeks,
    )
    schedule["provenance"] = {
        "origin": "reference_rebuild_staff_weekdays",
        "user_supplied": True,
        "input": "rebuild_reference_pack.params.class_weekdays",
        "merge_mode": "replace_all",
    }
    schedule["mapping_rules"].update(
        {
            "lab_visit": "must_map_to_teaching_dates_in_same_teaching_week",
            "lecture": "must_map_to_teaching_dates_in_same_teaching_week",
            "tutorial": "must_map_to_teaching_dates_in_same_teaching_week",
            "session_date": "must_map_to_teaching_dates_in_same_teaching_week",
            "class_visit": "must_map_to_teaching_dates_in_same_teaching_week",
            "assignment_deadline": "preserve_old_teaching_week_weekday_and_time",
            "submission_deadline": "preserve_old_teaching_week_weekday_and_time",
        }
    )
    validation_errors = validate_schedule_payload(schedule, reference_payloads)
    if validation_errors:
        return None, "; ".join(validation_errors)
    return schedule, ""


def non_teaching_gaps_from_schedule(
    schedule_payload: dict[str, Any],
    calendar_payload: dict[str, Any],
) -> list[dict[str, Any]]:
    weeks = [
        item
        for item in schedule_payload.get("weeks") or []
        if isinstance(item, dict) and item.get("calendar_week")
    ]
    if not weeks:
        return []
    facts = first_semester_facts(calendar_payload) or {}
    try:
        first_teaching = parse_iso(str(facts.get("first_teaching") or facts.get("start") or ""))
    except ValueError:
        return []
    present = {int(item["calendar_week"]) for item in weeks}
    last_calendar_week = max(present)
    weekdays = {
        str(day)
        for item in weeks
        for day in item.get("allowed_session_weekdays") or []
        if str(day) in WEEKDAY_NAMES
    }
    gaps: list[dict[str, Any]] = []
    for calendar_week in range(1, last_calendar_week + 1):
        if calendar_week in present:
            continue
        window_start = first_teaching + timedelta(days=7 * (calendar_week - 1))
        candidates = [
            window_start + timedelta(days=offset)
            for offset in range(7)
            if WEEKDAY_NAMES[(window_start + timedelta(days=offset)).weekday()] in weekdays
        ]
        reason = ""
        gap_date = window_start
        for candidate in candidates:
            candidate_reason = calendar_anchor_non_teaching_reason(calendar_payload, candidate)
            if candidate_reason:
                reason = candidate_reason
                gap_date = candidate
                break
        if not reason:
            for offset in range(7):
                candidate = window_start + timedelta(days=offset)
                candidate_reason = calendar_anchor_non_teaching_reason(calendar_payload, candidate)
                if candidate_reason:
                    reason = candidate_reason
                    gap_date = candidate
                    break
        gaps.append(
            {
                "date": gap_date.strftime("%Y-%m-%d"),
                "calendar_week": calendar_week,
                "reason": reason or "No eligible class date for the staff-provided recurring weekdays",
                "date_semantics": "cancelled_recurring_course_session_date",
                "class_weekdays": sorted(weekdays, key=WEEKDAY_NAMES.index),
                "week_numbering_effect": "does_not_consume_teaching_week",
            }
        )
    return gaps


def reflow_timeline_to_schedule(
    timeline_payload: dict[str, Any],
    schedule_payload: dict[str, Any],
    calendar_payload: dict[str, Any],
) -> dict[str, Any]:
    content_rows = teaching_content_rows(timeline_payload)
    schedule_weeks = [
        item for item in schedule_payload.get("weeks") or [] if isinstance(item, dict)
    ]
    if len(content_rows) != len(schedule_weeks):
        raise ValueError(
            f"timeline teaching row count {len(content_rows)} does not match schedule week count {len(schedule_weeks)}"
        )
    rows: list[dict[str, Any]] = []
    existing_mappings = [
        copy.deepcopy(item)
        for item in timeline_payload.get("content_week_mappings") or []
        if isinstance(item, dict)
    ]
    preserve_existing_mappings = bool(existing_mappings)
    content_week_mappings: list[dict[str, Any]] = existing_mappings
    for source, target in zip(content_rows, schedule_weeks):
        teaching_week = int(target.get("teaching_week") or target.get("week"))
        teaching_dates = [str(item) for item in target.get("teaching_dates") or []]
        source_teaching_week = int(
            source.get("source_teaching_week")
            or source.get("teaching_week")
            or source.get("week")
            or teaching_week
        )
        source_visible_week = int(
            source.get("source_visible_week")
            or source.get("source_week")
            or source_teaching_week
        )
        row = copy.deepcopy(source)
        content_item_ids = [
            str(item)
            for item in row.get("content_item_ids") or []
            if str(item or "").strip()
        ]
        if preserve_existing_mappings and not content_item_ids:
            for mapping in existing_mappings:
                try:
                    target_week = int(mapping.get("target_teaching_week"))
                except (TypeError, ValueError):
                    continue
                if target_week != teaching_week or str(mapping.get("status") or "") == "deleted":
                    continue
                item_id = str(mapping.get("item_id") or "").strip()
                if item_id:
                    content_item_ids.append(item_id)
        if not preserve_existing_mappings and not content_item_ids:
            for content_kind in ("lecture", "workshop", "submission"):
                content_text = clean_cell_text(source.get(content_kind))
                if not content_text:
                    continue
                item_id = f"tw{source_teaching_week}-{content_kind}"
                content_item_ids.append(item_id)
                content_week_mappings.append(
                    {
                        "item_id": item_id,
                        "content_kind": content_kind,
                        "content_text": content_text,
                        "source_teaching_week": source_teaching_week,
                        "source_visible_week": source_visible_week,
                        "source_calendar_week": source.get("source_calendar_week"),
                        "source_date": clean_cell_text(source.get("source_date") or source.get("date")),
                        "source_date_iso": clean_cell_text(source.get("source_date_iso")),
                        "source_weekday": clean_cell_text(source.get("source_weekday")),
                        "source_semester": clean_cell_text(source.get("source_semester")),
                        "source_academic_year": clean_cell_text(source.get("source_academic_year")),
                        "target_teaching_week": teaching_week,
                        "status": (
                            "unchanged"
                            if source_teaching_week == teaching_week
                            else "moved"
                        ),
                    }
                )
        row.setdefault("source_week", source_visible_week)
        row.setdefault("source_visible_week", source_visible_week)
        row.setdefault("source_teaching_week", source_teaching_week)
        row.setdefault("source_calendar_week", source.get("source_calendar_week"))
        row.setdefault("source_date", clean_cell_text(source.get("source_date") or source.get("date")))
        row.setdefault("source_date_iso", clean_cell_text(source.get("source_date_iso")))
        row.setdefault("source_weekday", clean_cell_text(source.get("source_weekday")))
        row.setdefault("source_semester", clean_cell_text(source.get("source_semester")))
        row.setdefault("source_academic_year", clean_cell_text(source.get("source_academic_year")))
        row.update(
            {
                "week": teaching_week,
                "teaching_week": teaching_week,
                "calendar_week": int(target["calendar_week"]),
                "date": teaching_dates[0],
                "teaching_dates": teaching_dates,
                "allowed_session_weekdays": list(target.get("allowed_session_weekdays") or []),
                "counts_as_teaching_week": True,
            }
        )
        if content_item_ids:
            row["content_item_ids"] = content_item_ids
        rows.append(row)
    gaps = non_teaching_gaps_from_schedule(schedule_payload, calendar_payload)
    rebuilt = dict(timeline_payload)
    rebuilt.update(
        {
            "status": "source-confirmed",
            "week_numbering": {
                "mode": "dual",
                "mapping_field": "teaching_week",
                "display_field": "date_planner_selected",
                "available_display_fields": ["calendar_week", "teaching_week"],
                "legacy_week_field": "teaching_week",
                "rule": (
                    "Mappings use teaching_week. File refresh Date Planner selects calendar_week or teaching_week "
                    "for each visible Week column from the source table semantics."
                ),
                "source_non_teaching_rows_excluded": True,
                "target_non_teaching_calendar_gaps_excluded": True,
            },
            "week_count": len(rows),
            "teaching_week_count": len(rows),
            "calendar_week_count": max(
                [
                    int(item.get("calendar_week") or 0)
                    for item in [*rows, *gaps]
                    if isinstance(item, dict)
                ]
                or [0]
            ),
            "week_calendar": rows,
            "content_week_mappings": content_week_mappings,
            "non_teaching_calendar_gaps": gaps,
            "class_weekdays": list(
                schedule_weeks[0].get("allowed_session_weekdays") or []
            )
            if schedule_weeks
            else [],
        }
    )
    return rebuilt


def build_calendar(report_path: Path, *, course_code: str, source_docs: list[str]) -> dict[str, Any]:
    report = load_json(report_path) or {}
    target = report.get("target_calendar") if isinstance(report.get("target_calendar"), dict) else {}
    return {
        "course_code": course_code,
        "academic_year": str(target.get("label") or ""),
        "target_semester": "first",
        "status": "source-confirmed",
        "source_basis": [repo_rel(report_path)],
        "source_documents": source_docs,
        "semesters": target.get("semesters") or [],
        "refresh_policy": {
            "unmapped_dates": "needs_review",
        },
    }


def normalize_existing_calendar(path: Path, *, course_code: str, source_docs: list[str]) -> dict[str, Any] | None:
    data = load_json(path)
    if not data or not isinstance(data.get("semesters"), (dict, list)):
        return None
    normalized = dict(data)
    normalized["course_code"] = str(normalized.get("course_code") or course_code)
    normalized.setdefault("status", "source-confirmed")
    normalized["source_basis"] = [repo_rel(path)]
    if source_docs:
        normalized["source_documents"] = source_docs
    normalized.setdefault("refresh_policy", {"unmapped_dates": "needs_review"})
    return normalized


def build_syllabus_rules(
    *,
    course_code: str,
    semester: str,
    timeline_path: Path | None,
    calendar_path: Path | None,
    source_manifest_path: Path,
    source_basis: list[str],
) -> dict[str, Any]:
    basis = list(source_basis)
    if timeline_path:
        basis.append(repo_rel(timeline_path))
    if calendar_path:
        basis.append(repo_rel(calendar_path))
    return {
        "course_code": course_code,
        "semester": semester,
        "status": "reproducible_reference",
        "source_basis": basis,
        "reproducibility": {
            "source_manifest": repo_rel(source_manifest_path),
        },
        "rules": [
            {
                "id": "source-first",
                "severity": "hard",
                "text": "Refresh editable Office/source files first when available; generated PDFs are derived artifacts.",
            },
            {
                "id": "dates-grounded",
                "severity": "hard",
                "text": "Visible course dates must be grounded in timeline.json or calendar.json; otherwise mark NEEDS_REVIEW.",
            },
        ],
    }


def default_refresh_rules_markdown(*, course_code: str, semester: str) -> str:
    return f"""# {course_code} Refresh Rules

Scope: single-file refresh previews for {course_code}, {semester}.

## Inputs

- Refresh exactly the file attached or explicitly referenced in the current staff prompt.
- Use this folder as the default reference pack:
  - `timeline.json` for week-level course schedule facts.
  - `calendar.json` for university semester periods, holiday/non-teaching days, teaching suspensions, and date checks.
  - `source_calendar.json`, when present, for resolving old visible dates to the source academic year, semester, calendar week, teaching week, and weekday.
  - `syllabus_rules.json` for hard policy guardrails.
  - `semester_schedule.json` for week-level class meeting dates generated during rebuild from the current staff-provided class weekdays and target calendar.
- Ignore unrelated files outside this reference pack unless the staff prompt explicitly names them.

## Source Priority

When sources disagree, apply this order:

1. Current staff prompt instructions.
2. Current attached file content and structure.
3. `timeline.json` for lecture, workshop, submission, and week-date facts.
4. `semester_schedule.json` for lab, lecture, tutorial, visit, and session-facing class dates.
5. `source_calendar.json` for old-year date coordinates and source non-teaching checks.
6. `calendar.json` for target semester, holiday/non-teaching days, teaching suspensions, reading week, revision period, and assessment period facts.
7. Existing wording in the file when no reference source gives a replacement fact.

## Edit Boundary

- Preserve the file's original purpose, audience, section order, and assessment meaning.
- Make the smallest edits needed to refresh stale course dates or schedule facts.
- Do not change marks, grading weights, deliverables, staff names, policies, or technical requirements unless the current prompt or reference pack supports the change.
- Do not apply broad year or semester text replacements without checking the surrounding meaning.

## Date Handling

- Treat visible dates, week labels, semester labels, and schedule rows as refresh candidates.
- Ground date changes in `timeline.json`, `semester_schedule.json`, or `calendar.json`.
- Resolve old dates through `source_calendar.json` and the typed source fields in `timeline.json` when available; do not infer the old academic year only by subtracting one from the target year.
- For old-course conflicts, `timeline.json` is authoritative over `source_calendar.json`: preserve explicit Timeline teaching rows and use recorded `source_calendar_overrides` evidence. Timeline rows explicitly marked Holiday or Reading Week remain non-teaching.
- Keep two week coordinates: use `teaching_week` for date/content mapping. For a visible Week column, the Date Planner must select `calendar_week` or `teaching_week` from that source table's semantics; it must not assume one display convention globally.
- Lab, lecture, tutorial, visit, and other session-facing dates must use target-week `teaching_dates` from `semester_schedule.json` and must not land on `calendar.json` entries where classes are suspended.
- Assignment/submission/deadline dates should preserve the old teaching-week number, weekday, and time when that source context is known; do not move deadlines solely because the mapped same-weekday date is a holiday or non-teaching day.
- Preserve dates that are citations, bibliography entries, software/version references, screenshot content, or historical examples.
- If a date appears course-facing but cannot be mapped to a reference fact, mark the result `NEEDS_REVIEW`.

## Output Status

- `READY_FOR_REVIEW`: all changed facts are grounded and no unresolved course-facing date remains.
- `BLOCKED`: the file cannot be safely refreshed from the available attachment and reference pack.

## Block Conditions

- The attached file is missing or cannot be opened.
- The file is PDF-only and the staff prompt did not authorize a PDF-copy refresh.
- The requested change depends on information not present in this reference pack.
- The requested change would rewrite course policy or assessment meaning without explicit support.
"""


def build_source_manifest(
    *,
    course_code: str,
    semester: str,
    status: str,
    staged_sources: list[dict[str, Any]],
    generated_files: list[Path],
    stable_files: list[Path],
    missing_inputs: list[str],
) -> dict[str, Any]:
    return {
        "course_code": course_code,
        "semester": semester,
        "status": status,
        "purpose": "Provenance map for the local reference pack rebuild.",
        "generated_references": [
            {
                "path": repo_rel(path),
                "sha256": sha256(path) if path.exists() else "",
            }
            for path in generated_files
        ],
        "stable_references": [
            {
                "path": repo_rel(path),
                "sha256": sha256(path) if path.exists() else "",
            }
            for path in stable_files
        ],
        "missing_inputs": missing_inputs,
        "source_files": staged_sources,
    }


def replace_payload_paths(value: Any, replacements: dict[str, str]) -> Any:
    if isinstance(value, dict):
        return {
            key: replace_payload_paths(item, replacements)
            for key, item in value.items()
        }
    if isinstance(value, list):
        return [replace_payload_paths(item, replacements) for item in value]
    if isinstance(value, str):
        return replacements.get(value, value)
    return value


def commit_staged_reference_sources(
    staged_sources: list[dict[str, Any]],
    target_root: Path,
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    committed: list[dict[str, Any]] = []
    replacements: dict[str, str] = {}
    for record in staged_sources:
        source = resolve_local_path(str(record.get("path") or ""))
        target = copy_to_sources(source, target_root / "sources")
        committed_record = dict(record)
        committed_record["path"] = repo_rel(target)
        replacements[str(record.get("path") or "")] = repo_rel(target)
        committed.append(committed_record)
    return committed, replacements


def ensure_refresh_rules_markdown(target_root: Path, *, course_code: str, semester: str) -> Path:
    refresh_rules_output = target_root / "refresh_rules.md"
    if not refresh_rules_output.exists():
        refresh_rules_output.write_text(
            default_refresh_rules_markdown(course_code=course_code, semester=semester),
            encoding="utf-8",
        )
    return refresh_rules_output


def execute_rebuild_reference_pack_operation(
    operation: dict[str, Any],
    env: dict[str, str],
    run_id: str,
    raw_task: str,
) -> dict[str, Any]:
    errors = validate_rebuild_reference_pack_shape(operation)
    if errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "blocker": "; ".join(errors),
            "attachments": [],
        }

    params = _params(operation)
    course_code = str(params.get("course_code") or env.get("COURSE_AGENT_DEFAULT_COURSE_CODE") or "DASE7137")
    semester = str(params.get("semester") or env.get("COURSE_AGENT_DEFAULT_SEMESTER") or "2026-2027 first semester")
    target_root = reference_pack_path(params, env)
    report_dir = REPORT_ROOT / f"{now_stamp()}_{safe_name(run_id)}"
    sources_root = report_dir / "staged_sources"
    target_root.mkdir(parents=True, exist_ok=True)

    staged_sources: list[dict[str, Any]] = []
    classified: dict[str, list[Path]] = {}
    input_warnings: list[str] = []
    for record in _reference_file_records(params):
        source = resolve_local_path(record["path"])
        copied = copy_to_sources(source, sources_root)
        role = classify_source(copied, record["role"])
        content_error = source_container_error(copied)
        staged = {
            "path": repo_rel(copied),
            "original_path": record["path"],
            "role": role,
            "detected_container": detect_reference_container(copied),
            "sha256": sha256(copied),
        }
        if content_error:
            staged["content_error"] = content_error
            input_warnings.append(f"{copied.name}: {content_error}")
        staged_sources.append(staged)
        if not content_error:
            classified.setdefault(role, []).append(copied)

    timeline_docs = [
        repo_rel(path)
        for path in (
            classified.get("timeline_reference", [])
            + classified.get("syllabus_reference", [])
        )
    ]
    generated: list[Path] = []
    stable_references: list[Path] = []
    missing: list[str] = []
    class_weekdays = class_weekdays_from_params(params)

    calendar_output = target_root / "calendar.json"
    source_calendar_output = target_root / "source_calendar.json"
    calendar_payload: dict[str, Any] | None = None
    target_calendar_candidate_path: Path | None = None
    calendar_candidates = [
        *[("calendar_grounding_report", path) for path in classified.get("calendar_grounding_report", [])],
        *[("calendar_json", path) for path in classified.get("calendar_json", [])],
        *[("target_calendar_reference", path) for path in classified.get("target_calendar_reference", [])],
        *[("calendar_reference", path) for path in classified.get("calendar_reference", [])],
    ]
    for role, candidate_path in calendar_candidates:
        try:
            if role == "calendar_grounding_report":
                candidate = build_calendar(
                    candidate_path,
                    course_code=course_code,
                    source_docs=[repo_rel(candidate_path)],
                )
            elif role == "calendar_json":
                candidate = normalize_existing_calendar(
                    candidate_path,
                    course_code=course_code,
                    source_docs=[repo_rel(candidate_path)],
                )
            else:
                candidate = build_calendar_from_reference(
                    candidate_path,
                    course_code=course_code,
                    source_docs=[repo_rel(candidate_path)],
                )
        except Exception as exc:
            input_warnings.append(
                f"{candidate_path.name}: calendar facts could not be parsed safely: {exc}"
            )
            continue
        if not candidate:
            input_warnings.append(
                f"{candidate_path.name}: no complete semester calendar facts were found"
            )
            continue
        if not calendar_payload_matches_target(candidate, semester):
            input_warnings.append(
                f"{candidate_path.name}: {calendar_target_mismatch_message(candidate, semester)}"
            )
            continue
        calendar_payload = candidate
        calendar_payload["calendar_role"] = "target"
        target_calendar_candidate_path = candidate_path
        break
    if not calendar_payload:
        missing.append(
            "target academic-year calendar reference with readable semester facts"
        )

    source_calendar_declared = bool(classified.get("source_calendar_reference"))
    source_calendar_candidates: list[tuple[Path, dict[str, Any]]] = []
    source_candidate_paths = [
        *classified.get("source_calendar_reference", []),
        *[
            path
            for path in classified.get("calendar_reference", [])
            if target_calendar_candidate_path is None
            or path.resolve(strict=False) != target_calendar_candidate_path.resolve(strict=False)
        ],
    ]
    for candidate_path in source_candidate_paths:
        try:
            candidate = build_calendar_from_reference(
                candidate_path,
                course_code=course_code,
                source_docs=[repo_rel(candidate_path)],
            )
        except Exception as exc:
            input_warnings.append(
                f"{candidate_path.name}: source calendar facts could not be parsed safely: {exc}"
            )
            continue
        if not candidate:
            input_warnings.append(
                f"{candidate_path.name}: no complete source-semester calendar facts were found"
            )
            continue
        if (
            candidate_path not in classified.get("source_calendar_reference", [])
            and calendar_payload_matches_target(candidate, semester)
        ):
            continue
        candidate["calendar_role"] = "source"
        candidate["purpose"] = (
            "Resolve old visible dates to source academic year, semester, calendar week, "
            "teaching week, and weekday."
        )
        source_calendar_candidates.append((candidate_path, candidate))
    if source_calendar_declared and not source_calendar_candidates:
        missing.append(
            "source academic-year calendar reference with readable semester facts"
        )

    timeline_output = target_root / "timeline.json"
    source_timeline_payload: dict[str, Any] | None = None
    source_calendar_payload: dict[str, Any] | None = None
    timeline_candidates = [
        *[("timeline_build_report", path) for path in classified.get("timeline_build_report", [])],
        *[("timeline_json", path) for path in classified.get("timeline_json", [])],
        *[("timeline_reference", path) for path in classified.get("timeline_reference", [])],
        *[("syllabus_reference", path) for path in classified.get("syllabus_reference", [])],
    ]
    for role, candidate_path in timeline_candidates:
        source_calendar_conflict_found = False
        try:
            if role == "timeline_build_report":
                candidate = build_timeline(
                    candidate_path,
                    course_code=course_code,
                    semester=semester,
                    source_docs=timeline_docs,
                )
            elif role == "timeline_json":
                candidate = normalize_existing_timeline(
                    candidate_path,
                    course_code=course_code,
                    semester=semester,
                    source_docs=timeline_docs,
                )
            elif calendar_payload:
                source_options: list[tuple[Path | None, dict[str, Any] | None]] = (
                    list(source_calendar_candidates)
                    if source_calendar_candidates
                    else [(None, None)]
                )
                candidate = None
                for source_calendar_path, source_calendar_candidate in source_options:
                    built = build_timeline_from_reference(
                        candidate_path,
                        course_code=course_code,
                        semester=semester,
                        calendar_payload=calendar_payload,
                        source_calendar_payload=source_calendar_candidate,
                        source_docs=timeline_docs,
                    )
                    conflicts = (
                        built.get("source_calendar_conflicts")
                        if isinstance(built, dict)
                        and isinstance(built.get("source_calendar_conflicts"), list)
                        else []
                    )
                    if conflicts:
                        source_calendar_conflict_found = True
                        label = (
                            source_calendar_path.name
                            if source_calendar_path
                            else "declared source calendar"
                        )
                        first_conflict = conflicts[0]
                        source_date = str(
                            first_conflict.get("source_date_iso")
                            or first_conflict.get("source_date")
                            or "unknown date"
                        )
                        calendar_reason = str(
                            first_conflict.get("calendar_reason")
                            or first_conflict.get("reason")
                            or "source calendar conflict"
                        )
                        input_warnings.append(
                            f"{candidate_path.name} with {label}: source date {source_date} "
                            f"conflicts with {calendar_reason}"
                        )
                        continue
                    candidate = built
                    source_calendar_payload = source_calendar_candidate
                    break
            else:
                candidate = None
        except Exception as exc:
            input_warnings.append(
                f"{candidate_path.name}: timeline rows could not be parsed safely: {exc}"
            )
            continue
        if candidate and teaching_content_rows(candidate):
            source_timeline_payload = candidate
            break
        if not source_calendar_conflict_found:
            input_warnings.append(
                f"{candidate_path.name}: no timeline Week/Lecture structure was found"
            )
    if not source_timeline_payload:
        missing.append(
            "timeline reference containing readable Week/Date/Lecture rows"
        )
        if source_calendar_candidates:
            missing.append(
                "source timeline dates consistent with the declared source academic calendar"
            )

    rules_basis = [
        item["path"]
        for item in staged_sources
        if item["role"]
        in {
            "timeline_build_report",
            "calendar_grounding_report",
            "timeline_reference",
            "calendar_reference",
            "source_calendar_reference",
            "target_calendar_reference",
            "syllabus_reference",
            "refresh_instructions",
        }
    ]
    schedule_output = target_root / "semester_schedule.json"
    schedule_action = ""
    schedule_generation_blocked = False
    schedule_payload: dict[str, Any] | None = None
    timeline_payload: dict[str, Any] | None = None
    if source_timeline_payload and calendar_payload and class_weekdays:
        schedule_payload, schedule_blocker = build_schedule_from_class_weekdays(
            source_timeline_payload,
            calendar_payload,
            class_weekdays=class_weekdays,
            course_code=course_code,
            semester=semester,
            source_basis=[
                "current staff rebuild class weekdays",
                *[
                    item["path"]
                    for item in staged_sources
                    if item["role"]
                    in {
                        "timeline_reference",
                        "syllabus_reference",
                        "timeline_build_report",
                        "source_calendar_reference",
                        "target_calendar_reference",
                        "calendar_reference",
                    }
                ],
            ],
        )
        if schedule_payload:
            schedule_payload["week_numbering"]["display_field"] = "date_planner_selected"
            schedule_payload["week_numbering"]["available_display_fields"] = [
                "calendar_week",
                "teaching_week",
            ]
            schedule_payload["week_numbering"]["rule"] = (
                "Mappings use teaching_week. File refresh Date Planner selects calendar_week or teaching_week "
                "for each visible Week column from the source table semantics."
            )
            timeline_payload = reflow_timeline_to_schedule(
                source_timeline_payload,
                schedule_payload,
                calendar_payload,
            )
            schedule_action = "generated_from_staff_class_weekdays"
        else:
            try:
                from timeline_draft_review import create_timeline_draft
            except ModuleNotFoundError:  # pragma: no cover
                from course_agent_app.timeline_draft_review import create_timeline_draft
            manifest_output = target_root / "source_manifest.json"
            draft_syllabus = build_syllabus_rules(
                course_code=course_code,
                semester=semester,
                timeline_path=timeline_output,
                calendar_path=calendar_output,
                source_manifest_path=manifest_output,
                source_basis=rules_basis,
            )
            draft_result = create_timeline_draft(
                origin=OPERATION_KIND,
                run_id=run_id,
                raw_task=raw_task,
                course_code=course_code,
                semester=semester,
                class_weekdays=class_weekdays,
                reference_pack_path=target_root,
                source_timeline=source_timeline_payload,
                calendar_payload=calendar_payload,
                candidate_payloads={
                    "syllabus_rules.json": draft_syllabus,
                    "source_manifest.json": {},
                    **(
                        {"source_calendar.json": source_calendar_payload}
                        if source_calendar_payload
                        else {}
                    ),
                },
                staged_sources=staged_sources,
                env=env,
            )
            if draft_result.get("ok"):
                draft_result.update(
                    {
                        "raw_task": raw_task,
                        "course_code": course_code,
                        "semester": semester,
                        "semester_schedule_action": "timeline_draft_review",
                        "staged_sources": staged_sources,
                        "generated_files": [],
                        "stable_reference_files": [],
                        "missing_inputs": [],
                        "input_warnings": input_warnings,
                    }
                )
                result_path = report_dir / "reference_pack_rebuild_result.json"
                write_json(result_path, draft_result)
                draft_result["result_path"] = repo_rel(result_path)
                return draft_result
            missing.append(schedule_blocker)
            schedule_generation_blocked = True
    elif not source_timeline_payload:
        missing.append("timeline source required to derive the target teaching schedule")
    elif not calendar_payload:
        missing.append("calendar.json required to derive the target teaching schedule")
    elif not class_weekdays:
        missing.append("staff-provided class weekdays required to derive the target teaching schedule")

    if timeline_payload and calendar_payload and schedule_payload and not missing:
        staged_sources, path_replacements = commit_staged_reference_sources(
            staged_sources,
            target_root,
        )
        calendar_payload = replace_payload_paths(calendar_payload, path_replacements)
        if source_calendar_payload:
            source_calendar_payload = replace_payload_paths(
                source_calendar_payload,
                path_replacements,
            )
        timeline_payload = replace_payload_paths(timeline_payload, path_replacements)
        schedule_payload = replace_payload_paths(schedule_payload, path_replacements)
        rules_basis = [
            path_replacements.get(str(item), str(item))
            for item in rules_basis
        ]
        syllabus_output = target_root / "syllabus_rules.json"
        manifest_output = target_root / "source_manifest.json"
        backup_dir = REPORT_ROOT / f"{now_stamp()}_{safe_name(run_id)}" / "previous_reference_json"
        for existing in (
            calendar_output,
            source_calendar_output,
            timeline_output,
            schedule_output,
            syllabus_output,
            manifest_output,
        ):
            if existing.is_file():
                backup_dir.mkdir(parents=True, exist_ok=True)
                shutil.copy2(existing, backup_dir / existing.name)
        write_json(calendar_output, calendar_payload)
        generated.append(calendar_output)
        if source_calendar_payload:
            write_json(source_calendar_output, source_calendar_payload)
            generated.append(source_calendar_output)
        write_json(timeline_output, timeline_payload)
        generated.append(timeline_output)
        write_json(schedule_output, schedule_payload)
        generated.append(schedule_output)
        write_json(
            syllabus_output,
            build_syllabus_rules(
                course_code=course_code,
                semester=semester,
                timeline_path=timeline_output if timeline_payload else None,
                calendar_path=calendar_output if calendar_payload else None,
                source_manifest_path=manifest_output,
                source_basis=rules_basis,
            ),
        )
        generated.append(syllabus_output)
        stable_references.append(
            ensure_refresh_rules_markdown(target_root, course_code=course_code, semester=semester)
        )

    if schedule_generation_blocked:
        status = "BLOCKED"
    elif not missing and timeline_payload and calendar_payload and schedule_payload:
        status = "READY"
    else:
        status = "NEEDS_STRUCTURED_EXTRACTION"
    manifest_output = target_root / "source_manifest.json"
    if status == "READY":
        write_json(
            manifest_output,
            build_source_manifest(
                course_code=course_code,
                semester=semester,
                status=status.lower(),
                staged_sources=staged_sources,
                generated_files=generated,
                stable_files=stable_references,
                missing_inputs=missing,
            ),
        )
        generated.append(manifest_output)

    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    result_path = REPORT_ROOT / f"{run_id}_{now_stamp()}_reference_pack_rebuild.json"
    result = {
        "ok": status == "READY",
        "status": status,
        "operation_kind": OPERATION_KIND,
        "run_id": run_id,
        "raw_task": raw_task,
        "course_code": course_code,
        "semester": semester,
        "reference_pack_path": repo_rel(target_root),
        "semester_schedule_action": schedule_action,
        "staged_sources": staged_sources,
        "generated_files": [repo_rel(path) for path in generated],
        "stable_reference_files": [repo_rel(path) for path in stable_references],
        "missing_inputs": missing,
        "input_warnings": input_warnings,
        "attachments": [],
    }
    write_json(result_path, result)
    result["result_path"] = repo_rel(result_path)
    if missing:
        prefix = "Schedule generation blocked: " if status == "BLOCKED" else "Missing structured reference input: "
        result["blocker"] = prefix + "; ".join(missing)
    return result
