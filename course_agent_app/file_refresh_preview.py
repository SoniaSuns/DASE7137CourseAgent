from __future__ import annotations

import json
import hashlib
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

try:
    from date_planning_agent import AGENT_NAME as DATE_PLANNER_AGENT_NAME
    from date_planning_agent import run_date_planning_agent
except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
    from course_agent_app.date_planning_agent import AGENT_NAME as DATE_PLANNER_AGENT_NAME
    from course_agent_app.date_planning_agent import run_date_planning_agent

try:
    from document_change_model import (
        apply_document_change_plan,
        build_document_model,
        planner_safe_document_model,
        sanitized_group_snapshot_hash,
        structure_index,
        validate_document_change_plan,
    )
    from file_change_planning_agent import (
        COMPOSITION_AGENT_NAME,
        GROUP_AGENT_NAME,
        run_composition_planner,
        run_group_document_planner,
    )
except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
    from course_agent_app.document_change_model import (
        apply_document_change_plan,
        build_document_model,
        planner_safe_document_model,
        sanitized_group_snapshot_hash,
        structure_index,
        validate_document_change_plan,
    )
    from course_agent_app.file_change_planning_agent import (
        COMPOSITION_AGENT_NAME,
        GROUP_AGENT_NAME,
        run_composition_planner,
        run_group_document_planner,
    )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ROOT = Path(__file__).resolve().parent
OPERATION_KIND = "refresh_file_preview"
EXECUTION_MODE = "local_file_action"

REFERENCE_ROOT = PROJECT_ROOT / "course_files" / "reference"
CURRENT_REFERENCE_ROOT = PROJECT_ROOT / "course_files" / "pre_refresh" / "current_reference"
POST_REFRESH_ROOT = PROJECT_ROOT / "course_files" / "post_refresh" / "current_course" / "per_file"
UPLOAD_READY_ROOT = PROJECT_ROOT / "course_files" / "upload_ready" / "current_course" / "per_file"
REPORT_ROOT = PROJECT_ROOT / "course_files" / "evidence" / "process" / "course_agent_app" / "file_refresh_preview"

OFFICE_SUFFIXES = {".docx", ".pptx", ".xlsx"}
TEXT_SUFFIXES = {".txt", ".md", ".csv"}
PDF_SUFFIXES = {".pdf"}
SUPPORTED_SUFFIXES = OFFICE_SUFFIXES | TEXT_SUFFIXES | PDF_SUFFIXES
REFERENCE_SUFFIXES = {".json", ".md", ".txt", ".yaml", ".yml"}
REFERENCE_PACK_REQUIRED_FILENAMES = (
    "timeline.json",
    "calendar.json",
    "syllabus_rules.json",
    "source_manifest.json",
    "semester_schedule.json",
    "refresh_rules.md",
)
FORBIDDEN_JOB_KEYS = {
    "command",
    "shell",
    "python",
    "python_code",
    "javascript",
    "javascript_code",
    "raw_code",
    "script_text",
}
STUDENT_VISIBLE_TRACE_MARKERS = (
    "Codex",
    "agent",
    "AI-generated",
    "AI modified",
    "AI-modified",
    "minimal refresh",
    "refresh draft",
    "refreshed",
    "revision",
    "marked comparison",
    "修改痕迹",
    "改稿提醒",
    "翻新",
)
STALE_DATE_PATTERNS = (
    re.compile(r"\b202[0-5]\s*-\s*202[1-6]\b"),
    re.compile(r"\b202[0-5]\b"),
    re.compile(r"\b\d{1,2}/\d{1,2}/(?:20)?2[0-5]\b"),
)
DATEISH_PATTERN = re.compile(
    r"\d|jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec|mon|tue|wed|thu|fri|sat|sun|sem",
    re.IGNORECASE,
)
DATE_MENTION_PATTERNS = (
    re.compile(r"\b20\d{2}\s*[-/]\s*20\d{2}\b"),
    re.compile(r"\b20\d{2}\s*[-/]\s*\d{1,2}\s*[-/]\s*\d{1,2}\b"),
    re.compile(r"\b\d{1,2}\s*[-/]\s*\d{1,2}\s*[-/]\s*(?:20)?\d{2}\b"),
    re.compile(
        r"\b\d{1,2}\s*[-\s]\s*(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z.]*\s*[-\s]\s*20\s?\d\s?\d\b",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b(?:Mon|Tue|Tues|Wed|Thu|Thur|Fri|Sat|Sun)[a-z]*,?\s+"
        r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+20\d{2}\b",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+20\d{2}\b",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.?\s+20\d{2}\b",
        re.IGNORECASE,
    ),
    re.compile(
        r"\b\d{1,2}\s*(?:st|nd|rd|th)?\s*,?\s*"
        r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z.]*\b",
        re.IGNORECASE,
    ),
    re.compile(r"\b202[0-5]\b"),
)
ACCEPTED_DATE_PLAN_CONFIDENCE = {"high", "medium"}
OOXML_TEXT_NODE_NAMES = {"t", "instrText"}
OOXML_BLOCK_NODE_NAMES = {"p", "si"}
WEEKDAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
WEEKDAY_ALIASES = {
    "mon": "Monday",
    "monday": "Monday",
    "tue": "Tuesday",
    "tues": "Tuesday",
    "tuesday": "Tuesday",
    "wed": "Wednesday",
    "wednesday": "Wednesday",
    "thu": "Thursday",
    "thur": "Thursday",
    "thurs": "Thursday",
    "thursday": "Thursday",
    "fri": "Friday",
    "friday": "Friday",
    "sat": "Saturday",
    "saturday": "Saturday",
    "sun": "Sunday",
    "sunday": "Sunday",
    "周一": "Monday",
    "星期一": "Monday",
    "周二": "Tuesday",
    "星期二": "Tuesday",
    "周三": "Wednesday",
    "星期三": "Wednesday",
    "周四": "Thursday",
    "星期四": "Thursday",
    "周五": "Friday",
    "星期五": "Friday",
    "周六": "Saturday",
    "星期六": "Saturday",
    "周日": "Sunday",
    "周天": "Sunday",
    "星期日": "Sunday",
    "星期天": "Sunday",
}
MONTH_ALIASES = {
    "jan": 1,
    "january": 1,
    "feb": 2,
    "february": 2,
    "mar": 3,
    "march": 3,
    "apr": 4,
    "april": 4,
    "may": 5,
    "jun": 6,
    "june": 6,
    "jul": 7,
    "july": 7,
    "aug": 8,
    "august": 8,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "oct": 10,
    "october": 10,
    "nov": 11,
    "november": 11,
    "dec": 12,
    "december": 12,
}
SCHEDULE_MAPPING_POLICIES = {"same_week_teaching_date", "same_weekday_deadline"}
SCHEDULE_TABLE_REBUILD_POLICY = "teaching_week_schedule_rebuild"
NON_TEACHING_TIMELINE_MARKERS = (
    "holiday",
    "reading week",
    "reading/field trip",
    "revision period",
    "assessment period",
    "class suspension",
)


def now_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")


def safe_name(value: str, fallback: str = "file") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "-", str(value or "")).strip(" .-_")
    cleaned = re.sub(r"\s+", " ", cleaned)
    cleaned = cleaned or fallback
    limit = 80 if os.name == "nt" else 120
    if len(cleaned) <= limit:
        return cleaned
    digest = hashlib.sha256(cleaned.encode("utf-8")).hexdigest()[:8]
    return f"{cleaned[: limit - len(digest) - 1]}-{digest}"


def resolve_local_path(path_text: str | Path) -> Path:
    path = Path(str(path_text or "")).expanduser()
    if path.is_absolute():
        return path
    for base in (PROJECT_ROOT, ROOT, Path.cwd()):
        candidate = base / path
        if candidate.exists():
            return candidate
    return PROJECT_ROOT / path


def repo_rel(path: str | Path) -> str:
    resolved = resolve_local_path(path)
    try:
        return str(resolved.resolve(strict=False).relative_to(PROJECT_ROOT))
    except ValueError:
        return str(resolved)


def reference_pack_is_ready(path: Path) -> bool:
    return path.is_dir() and all((path / filename).is_file() for filename in REFERENCE_PACK_REQUIRED_FILENAMES)


def canonical_semester_key(value: str) -> str:
    text = str(value or "").strip()
    lowered = text.lower()
    compact = re.sub(r"\s+", "", lowered)
    match = re.search(r"\b(20\d{2})\s*[-/]\s*(20\d{2})\b", text)
    if match and (
        "sem1" in compact
        or "semester1" in compact
        or "firstsemester" in compact
        or "first semester" in lowered
        or "sem 1" in lowered
    ):
        return f"{match.group(1)}-{match.group(2)}-sem1"
    return safe_name(text, "")


def ready_reference_pack_under(root: Path, *, course_code: str = "", semester: str = "") -> Path | None:
    if reference_pack_is_ready(root):
        return root
    if not root.exists() or not root.is_dir():
        return None
    desired_course = course_code.casefold()
    desired_semester = canonical_semester_key(semester).casefold()
    candidates: list[Path] = []
    for manifest in root.rglob("source_manifest.json"):
        pack = manifest.parent
        if any(part.casefold().startswith("backup") for part in pack.parts):
            continue
        if reference_pack_is_ready(pack):
            candidates.append(pack)
    if not candidates:
        return None

    def score(path: Path) -> tuple[int, float]:
        parts = [part.casefold() for part in path.parts]
        match_score = 0
        if desired_course and desired_course in parts:
            match_score += 4
        if desired_semester and desired_semester in parts:
            match_score += 8
        try:
            mtime = (path / "source_manifest.json").stat().st_mtime
        except OSError:
            mtime = path.stat().st_mtime
        return match_score, mtime

    return max(candidates, key=score)


def is_refresh_file_preview_operation(operation: dict[str, Any]) -> bool:
    return str(operation.get("operation_kind", "")).strip() == OPERATION_KIND


def _path_error(label: str, value: str) -> str:
    if value.startswith(("http://", "https://")):
        return f"{label} must be a local path"
    if ".env" in Path(value).parts:
        return f"{label} must not point at a secret env file"
    return ""


def source_file_path_from_operation(operation: dict[str, Any]) -> str:
    params = operation.get("params") if isinstance(operation.get("params"), dict) else {}
    return str(
        params.get("file_path")
        or params.get("source_file_path")
        or params.get("path")
        or operation.get("source_file_path")
        or operation.get("file_path")
        or ""
    ).strip()


def validate_refresh_file_preview_shape(operation: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    params = operation.get("params") if isinstance(operation.get("params"), dict) else {}
    forbidden = sorted(FORBIDDEN_JOB_KEYS & set(params))
    if forbidden:
        errors.append("refresh_file_preview.params contains forbidden executable fields: " + ", ".join(forbidden))
    file_path = source_file_path_from_operation(operation)
    if not file_path:
        errors.append("refresh_file_preview requires params.file_path")
    else:
        path_error = _path_error("refresh_file_preview.params.file_path", file_path)
        if path_error:
            errors.append(path_error)
        elif not resolve_local_path(file_path).exists():
            errors.append(f"refresh_file_preview.params.file_path does not exist: {file_path}")
    for key in ("reference_pack_path", "instructions_path"):
        value = str(params.get(key) or operation.get(key) or "").strip()
        if not value:
            continue
        path_error = _path_error(f"refresh_file_preview.params.{key}", value)
        if path_error:
            errors.append(path_error)
    return errors


def reference_pack_path(params: dict[str, Any], env: dict[str, str]) -> Path:
    explicit = str(params.get("reference_pack_path") or env.get("COURSE_REFERENCE_PACK_PATH") or "").strip()
    course_code = safe_name(
        str(params.get("course_code") or env.get("COURSE_AGENT_DEFAULT_COURSE_CODE") or ""),
        "",
    )
    semester_text = str(params.get("semester") or env.get("COURSE_AGENT_DEFAULT_SEMESTER") or "")
    if explicit:
        resolved = resolve_local_path(explicit)
        if reference_pack_is_ready(resolved):
            return resolved
        nested = ready_reference_pack_under(resolved, course_code=course_code, semester=semester_text)
        if nested:
            return nested
        return resolved
    semester = canonical_semester_key(semester_text)
    candidates: list[Path] = []
    if course_code and semester:
        candidates.append(REFERENCE_ROOT / course_code / semester)
    if course_code:
        candidates.append(REFERENCE_ROOT / course_code)
    candidates.append(REFERENCE_ROOT)
    if str(env.get("COURSE_AGENT_ALLOW_CURRENT_REFERENCE_FALLBACK", "")).strip().lower() in {"1", "true", "yes", "on"}:
        candidates.append(CURRENT_REFERENCE_ROOT)
    for candidate in candidates:
        if reference_pack_is_ready(candidate):
            return candidate
    nested = ready_reference_pack_under(REFERENCE_ROOT, course_code=course_code, semester=semester_text)
    if nested:
        return nested
    return next((candidate for candidate in candidates if candidate.exists()), candidates[-1])


def read_text_lossy(path: Path, limit: int = 80_000) -> str:
    try:
        return path.read_text(encoding="utf-8")[:limit]
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-8", errors="replace")[:limit]


def has_direct_replacement_instructions(params: dict[str, Any]) -> bool:
    for key in ("replacements", "text_replacements"):
        value = params.get(key)
        if isinstance(value, dict) and value:
            return True
        if isinstance(value, list) and value:
            return True
    return False


def extract_date_mentions(text: str, *, limit: int = 80) -> list[dict[str, Any]]:
    mentions: list[dict[str, Any]] = []
    seen: set[tuple[str, int]] = set()
    for pattern in DATE_MENTION_PATTERNS:
        for match in pattern.finditer(text):
            visible = re.sub(r"\s+", " ", match.group(0)).strip()
            if not visible:
                continue
            start = match.start()
            key = (visible.casefold(), start)
            if key in seen:
                continue
            seen.add(key)
            context_start = max(0, start - 90)
            context_end = min(len(text), match.end() + 90)
            context = re.sub(r"\s+", " ", text[context_start:context_end]).strip()
            mentions.append(
                {
                    "visible_text": visible,
                    "start": start,
                    "end": match.end(),
                    "context": context,
                    "stale": bool(any(str(hit) in visible for hit in audit_visible_text(visible)["stale_date_hits"])),
                }
            )
    mentions.sort(key=lambda item: int(item["start"]))
    deduped: list[dict[str, Any]] = []
    covered_spans: list[tuple[int, int]] = []
    for item in mentions:
        start = int(item["start"])
        end = int(item["end"])
        if any(prev_start <= start and end <= prev_end for prev_start, prev_end in covered_spans):
            continue
        covered_spans.append((start, end))
        deduped.append(item)
        if len(deduped) >= limit:
            break
    return deduped


def normalize_academic_year(value: str) -> str:
    text = str(value or "").strip()
    match = re.search(r"\b(20\d{2})\s*[-/]\s*(20\d{2})\b", text)
    if match:
        return f"{match.group(1)}-{match.group(2)}"
    short = re.search(r"(?<!\d)(\d{2})\s*[-_/]\s*(\d{2})(?!\d)", text)
    if short:
        start = int(short.group(1))
        end = int(short.group(2))
        if 0 <= start <= 99 and 0 <= end <= 99:
            return f"20{start:02d}-20{end:02d}"
    return ""


def academic_year_start(value: str) -> int | None:
    academic_year = normalize_academic_year(value)
    if not academic_year:
        return None
    try:
        return int(academic_year.split("-", 1)[0])
    except ValueError:
        return None


def previous_academic_year(value: str) -> str:
    start = academic_year_start(value)
    if start is None:
        return ""
    return f"{start - 1}-{start}"


def academic_year_start_month(reference_payloads: dict[str, Any]) -> int:
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    semesters = calendar.get("semesters") if isinstance(calendar.get("semesters"), dict) else {}
    first = semesters.get("first") if isinstance(semesters.get("first"), dict) else {}
    first_start = parse_date_value(str(first.get("start") or first.get("first_teaching") or ""))
    return first_start.month if first_start else 9


def academic_year_date_for_month_day(academic_year: str, month: int, day: int, *, start_month: int = 9) -> date | None:
    start = academic_year_start(academic_year)
    if start is None:
        return None
    year = start if month >= start_month else start + 1
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_month_day_value(text: str) -> tuple[int, int] | None:
    value = re.sub(r"\s+", " ", str(text or "")).strip()
    if not value or parse_date_value(value):
        return None
    month_names = "|".join(sorted(MONTH_ALIASES, key=len, reverse=True))
    day_month = re.search(
        rf"\b(\d{{1,2}})(?:st|nd|rd|th)?[,\-\s]+({month_names})\.?\b",
        value,
        re.IGNORECASE,
    )
    if day_month:
        month = MONTH_ALIASES.get(day_month.group(2).casefold().rstrip("."))
        if month:
            return month, int(day_month.group(1))
    month_day = re.search(
        rf"\b({month_names})\.?[,\s]+(\d{{1,2}})(?:st|nd|rd|th)?\b",
        value,
        re.IGNORECASE,
    )
    if month_day:
        month = MONTH_ALIASES.get(month_day.group(1).casefold().rstrip("."))
        if month:
            return month, int(month_day.group(2))
    chinese = re.search(r"\b(\d{1,2})\s*月\s*(\d{1,2})\s*日?\b", value)
    if chinese:
        return int(chinese.group(1)), int(chinese.group(2))
    return None


def collect_academic_year_candidates(params: dict[str, Any], env: dict[str, str], reference_payloads: dict[str, Any]) -> dict[str, Any]:
    explicit: list[dict[str, str]] = []
    for key in ("source_academic_year", "old_academic_year"):
        academic_year = normalize_academic_year(str(params.get(key) or ""))
        if academic_year:
            explicit.append({"academic_year": academic_year, "basis": f"refresh_file_preview.params.{key}"})
    for key in ("COURSE_AGENT_SOURCE_ACADEMIC_YEAR", "COURSE_AGENT_OLD_ACADEMIC_YEAR"):
        academic_year = normalize_academic_year(str(env.get(key) or ""))
        if academic_year:
            explicit.append({"academic_year": academic_year, "basis": key})

    target = ""
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    if isinstance(calendar, dict):
        target = normalize_academic_year(str(calendar.get("academic_year") or ""))
    target = target or normalize_academic_year(str(params.get("semester") or env.get("COURSE_AGENT_DEFAULT_SEMESTER") or ""))

    referenced: list[dict[str, str]] = []
    for payload_name in (
        "source_calendar.json",
        "source_manifest.json",
        "calendar.json",
        "timeline.json",
        "syllabus_rules.json",
    ):
        payload = reference_payloads.get(payload_name)
        if not isinstance(payload, dict):
            continue
        for key in ("source_academic_year", "old_academic_year", "source_year"):
            academic_year = normalize_academic_year(str(payload.get(key) or ""))
            if academic_year:
                referenced.append({"academic_year": academic_year, "basis": f"{payload_name}.{key}"})

    derived: list[dict[str, str]] = []
    if not explicit and not referenced:
        previous = previous_academic_year(target)
        if previous:
            derived.append({"academic_year": previous, "basis": "previous academic year derived from target academic_year"})

    source_candidates: list[dict[str, str]] = []
    seen: set[str] = set()
    for candidate in explicit + referenced + derived:
        academic_year = candidate["academic_year"]
        if not academic_year or academic_year == target or academic_year in seen:
            continue
        seen.add(academic_year)
        source_candidates.append(candidate)
    return {
        "target_academic_year": target,
        "source_academic_year_candidates": source_candidates,
    }


def shift_year(value: date, delta_years: int) -> date | None:
    try:
        return value.replace(year=value.year + delta_years)
    except ValueError:
        if value.month == 2 and value.day == 29:
            return value.replace(year=value.year + delta_years, day=28)
    return None


def target_week_anchors(reference_payloads: dict[str, Any]) -> list[dict[str, Any]]:
    anchors: list[dict[str, Any]] = []
    timeline = reference_payloads.get("timeline.json") if isinstance(reference_payloads.get("timeline.json"), dict) else {}
    teaching_week = 0
    for row in timeline.get("week_calendar") if isinstance(timeline.get("week_calendar"), list) else []:
        if not isinstance(row, dict):
            continue
        marker = " ".join(str(row.get(key) or "").lower() for key in ("lecture", "workshop", "tutorial", "notes"))
        if row.get("counts_as_teaching_week") is False or any(token in marker for token in NON_TEACHING_TIMELINE_MARKERS):
            continue
        teaching_week += 1
        stored_teaching_week = _int_field(row, "teaching_week")
        actual_teaching_week = stored_teaching_week or teaching_week
        parsed = parse_date_value(str(row.get("date") or ""))
        if parsed:
            anchors.append(
                {
                    "week": actual_teaching_week,
                    "target_anchor_date": parsed,
                    "basis": "timeline.json.week_calendar actual teaching_week",
                    "source_week": _int_field(row, "source_week", "week"),
                }
            )
    if anchors:
        return sorted(anchors, key=lambda item: int(item["week"]))
    schedule = reference_payloads.get("semester_schedule.json") if isinstance(reference_payloads.get("semester_schedule.json"), dict) else {}
    for row in schedule.get("weeks") if isinstance(schedule.get("weeks"), list) else []:
        if not isinstance(row, dict):
            continue
        week = _int_field(row, "week")
        raw_dates = row.get("teaching_dates") if isinstance(row.get("teaching_dates"), list) else []
        parsed_dates = sorted(item for item in (parse_date_value(str(raw_date)) for raw_date in raw_dates) if item)
        if week is not None and parsed_dates:
            anchors.append({"week": week, "target_anchor_date": parsed_dates[0], "basis": "semester_schedule.json.weeks"})
    return sorted(anchors, key=lambda item: int(item["week"]))


def explicit_source_week_records(
    reference_payloads: dict[str, Any],
) -> list[dict[str, Any]]:
    timeline = (
        reference_payloads.get("timeline.json")
        if isinstance(reference_payloads.get("timeline.json"), dict)
        else {}
    )
    records: dict[tuple[int, str], dict[str, Any]] = {}
    mappings = (
        timeline.get("content_week_mappings")
        if isinstance(timeline.get("content_week_mappings"), list)
        else []
    )
    for item in mappings:
        if not isinstance(item, dict):
            continue
        week = _int_field(item, "source_teaching_week", "source_week")
        source_date_iso = str(item.get("source_date_iso") or "").strip()
        parsed = parse_date_value(source_date_iso)
        if week is None or not parsed:
            continue
        records[(week, parsed.isoformat())] = {
            "source_teaching_week": week,
            "source_date": parsed,
            "source_visible_week": _int_field(item, "source_visible_week"),
            "source_calendar_week": _int_field(item, "source_calendar_week"),
            "source_weekday": str(item.get("source_weekday") or WEEKDAY_NAMES[parsed.weekday()]),
            "source_semester": str(item.get("source_semester") or ""),
            "source_academic_year": str(item.get("source_academic_year") or ""),
        }
    if records:
        return sorted(
            records.values(),
            key=lambda item: (int(item["source_teaching_week"]), item["source_date"]),
        )
    for item in timeline.get("week_calendar") if isinstance(timeline.get("week_calendar"), list) else []:
        if not isinstance(item, dict):
            continue
        week = _int_field(item, "source_teaching_week", "teaching_week", "week")
        parsed = parse_date_value(str(item.get("source_date_iso") or ""))
        if week is None or not parsed:
            continue
        records[(week, parsed.isoformat())] = {
            "source_teaching_week": week,
            "source_date": parsed,
            "source_visible_week": _int_field(item, "source_visible_week", "source_week"),
            "source_calendar_week": _int_field(item, "source_calendar_week"),
            "source_weekday": str(item.get("source_weekday") or WEEKDAY_NAMES[parsed.weekday()]),
            "source_semester": str(item.get("source_semester") or ""),
            "source_academic_year": str(item.get("source_academic_year") or ""),
        }
    return sorted(
        records.values(),
        key=lambda item: (int(item["source_teaching_week"]), item["source_date"]),
    )


def week_mapping_candidates_for_source_date(
    source_date: date,
    *,
    source_academic_year: str,
    target_academic_year: str,
    reference_payloads: dict[str, Any],
    content_kind: str = "",
) -> list[dict[str, Any]]:
    source_start = academic_year_start(source_academic_year)
    target_start = academic_year_start(target_academic_year)
    if source_start is None or target_start is None:
        return []
    delta_years = source_start - target_start
    weekday = WEEKDAY_NAMES[source_date.weekday()]
    candidates: list[dict[str, Any]] = []
    anchors = target_week_anchors(reference_payloads)
    anchors_by_week = {int(item["week"]): item for item in anchors}
    timeline = reference_payloads.get("timeline.json")
    raw_mappings = (
        timeline.get("content_week_mappings")
        if isinstance(timeline, dict) and isinstance(timeline.get("content_week_mappings"), list)
        else []
    )
    explicit_records = explicit_source_week_records(reference_payloads)
    source_week_records: list[dict[str, Any]] = []
    for record in explicit_records:
        record_date = record["source_date"]
        week_start = record_date - timedelta(days=record_date.weekday())
        if week_start <= source_date <= week_start + timedelta(days=6):
            source_week_records.append(record)

    if explicit_records:
        source_weeks = [
            (
                int(record["source_teaching_week"]),
                record,
            )
            for record in source_week_records
        ]
    else:
        source_weeks = []
        for anchor in anchors:
            target_anchor = anchor["target_anchor_date"]
            source_anchor = shift_year(target_anchor, delta_years)
            if not source_anchor:
                continue
            source_week_start = source_anchor - timedelta(days=source_anchor.weekday())
            source_week_end = source_week_start + timedelta(days=6)
            if source_week_start <= source_date <= source_week_end:
                source_weeks.append((int(anchor["week"]), {}))

    for source_week, source_record in source_weeks:
        anchor = anchors_by_week.get(source_week)
        if not anchor:
            continue
        target_anchor = anchor["target_anchor_date"]
        target_week_start = target_anchor - timedelta(days=target_anchor.weekday())
        target_same_weekday = target_week_start + timedelta(days=source_date.weekday())
        mapping_candidates = [
            {
                "item_id": item.get("item_id"),
                "content_kind": item.get("content_kind"),
                "content_text": item.get("content_text"),
                "source_teaching_week": item.get("source_teaching_week"),
                "target_teaching_week": item.get("target_teaching_week"),
                "status": item.get("status"),
            }
            for item in raw_mappings
            if isinstance(item, dict)
            and _int_field(item, "source_teaching_week", "source_week") == source_week
            and (not content_kind or str(item.get("content_kind") or "") == content_kind)
        ]
        candidate = {
            "source_week": source_week,
            "target_week": source_week,
            "source_weekday": weekday,
            "target_same_weekday_date": target_same_weekday.isoformat(),
            "basis": (
                "timeline.json explicit source teaching-week date anchor"
                if source_record
                else anchor["basis"]
            ),
        }
        if source_record:
            candidate.update(
                {
                    key: value
                    for key, value in source_record.items()
                    if key != "source_date" and value not in ("", None)
                }
            )
            candidate["source_date_iso"] = source_date.isoformat()
        if mapping_candidates:
            candidate["content_week_mapping_candidates"] = mapping_candidates
        if len(mapping_candidates) == 1:
            mapping = mapping_candidates[0]
            target_week = _int_field(mapping, "target_teaching_week", "target_week")
            candidate["mapping_item_id"] = mapping.get("item_id")
            candidate["content_kind"] = mapping.get("content_kind")
            candidate["mapping_status"] = mapping.get("status")
            candidate["target_week"] = target_week
            if target_week is None or str(mapping.get("status") or "").lower() == "deleted":
                candidate.pop("target_same_weekday_date", None)
            elif target_week in anchors_by_week:
                mapped_anchor = anchors_by_week[target_week]["target_anchor_date"]
                mapped_week_start = mapped_anchor - timedelta(days=mapped_anchor.weekday())
                candidate["target_same_weekday_date"] = (
                    mapped_week_start + timedelta(days=source_date.weekday())
                ).isoformat()
                candidate["basis"] = (
                    f"{anchor['basis']} plus timeline.json.content_week_mappings "
                    f"{mapping.get('item_id')} {source_week}->{target_week}"
                )
        candidates.append(candidate)
    return candidates


def enrich_date_mentions_with_source_context(
    mentions: list[dict[str, Any]],
    *,
    params: dict[str, Any],
    env: dict[str, str],
    reference_pack: dict[str, Any],
) -> list[dict[str, Any]]:
    reference_payloads = reference_pack.get("reference_payloads") if isinstance(reference_pack.get("reference_payloads"), dict) else {}
    year_context = collect_academic_year_candidates(params, env, reference_payloads)
    source_years = year_context["source_academic_year_candidates"]
    if not source_years:
        return mentions
    enriched: list[dict[str, Any]] = []
    target_academic_year = str(year_context.get("target_academic_year") or "")
    start_month = academic_year_start_month(reference_payloads)
    for mention in mentions:
        item = dict(mention)
        mention_kind = str(item.get("date_kind") or "").strip().lower()
        content_kind = (
            "submission"
            if any(token in mention_kind for token in ("deadline", "submission", "quiz", "deliverable"))
            else ""
        )
        month_day = parse_month_day_value(str(item.get("visible_text") or ""))
        if not month_day:
            enriched.append(item)
            continue
        month, day = month_day
        candidates: list[dict[str, Any]] = []
        for source_year in source_years:
            academic_year = source_year["academic_year"]
            source_date = academic_year_date_for_month_day(academic_year, month, day, start_month=start_month)
            if not source_date:
                continue
            candidate: dict[str, Any] = {
                "source_academic_year": academic_year,
                "source_date_iso": source_date.isoformat(),
                "source_weekday": WEEKDAY_NAMES[source_date.weekday()],
                "basis": source_year["basis"],
            }
            week_candidates = week_mapping_candidates_for_source_date(
                source_date,
                source_academic_year=academic_year,
                target_academic_year=target_academic_year,
                reference_payloads=reference_payloads,
                content_kind=content_kind,
            )
            if week_candidates:
                candidate["same_weekday_deadline_candidates"] = week_candidates
            candidates.append(candidate)
        if candidates:
            item["source_date_candidates"] = candidates
        enriched.append(item)
    return enriched


def deleted_content_mapping_reasons(mentions: list[dict[str, Any]]) -> list[str]:
    reasons: list[str] = []
    for mention in mentions:
        if not isinstance(mention, dict):
            continue
        for source_candidate in mention.get("source_date_candidates") or []:
            if not isinstance(source_candidate, dict):
                continue
            for week_candidate in source_candidate.get("same_weekday_deadline_candidates") or []:
                if not isinstance(week_candidate, dict):
                    continue
                mapping_candidates = week_candidate.get("content_week_mapping_candidates")
                if not isinstance(mapping_candidates, list) or len(mapping_candidates) != 1:
                    continue
                mapping = mapping_candidates[0]
                if not isinstance(mapping, dict):
                    continue
                if (
                    str(mapping.get("status") or "").strip().lower() == "deleted"
                    or mapping.get("target_teaching_week") is None
                ):
                    reasons.append(
                        f"{mention.get('visible_text')}: content item "
                        f"{mapping.get('item_id')} was deleted from the confirmed timeline and has no target week"
                    )
    return list(dict.fromkeys(reasons))


def load_reference_pack(params: dict[str, Any], env: dict[str, str]) -> dict[str, Any]:
    root = reference_pack_path(params, env)
    files: list[dict[str, str]] = []
    reference_payloads: dict[str, Any] = {}
    reference_text_parts: list[str] = []
    paths: list[Path] = []
    if root.exists():
        paths.extend([root] if root.is_file() else sorted(path for path in root.iterdir() if path.is_file()))
    instructions_path = str(params.get("instructions_path") or "").strip()
    if instructions_path:
        resolved_instructions = resolve_local_path(instructions_path)
        if resolved_instructions.exists() and resolved_instructions not in paths:
            paths.append(resolved_instructions)
    for path in paths[:80]:
        if path.suffix.lower() not in REFERENCE_SUFFIXES:
            continue
        rel = repo_rel(path)
        files.append({"path": rel, "name": path.name, "suffix": path.suffix.lower()})
        try:
            path.relative_to(root / "sources")
            continue
        except ValueError:
            pass
        text = read_text_lossy(path)
        reference_text_parts.append(f"\n\n# {rel}\n{text}")
        if path.suffix.lower() == ".json":
            try:
                payload = json.loads(text)
                reference_payloads[path.name] = payload
            except json.JSONDecodeError:
                pass
    reference_text = "".join(reference_text_parts)
    return {
        "root": repo_rel(root),
        "exists": root.exists(),
        "files": files,
        "replacements": [],
        "reference_payloads": reference_payloads,
        "reference_text_excerpt": reference_text[:12_000],
    }


def date_planning_task(
    *,
    report_dir: Path,
    source_path: Path,
    source_text: str,
    date_mentions: list[dict[str, Any]],
    reference_pack: dict[str, Any],
    raw_task: str,
    document_model: dict[str, Any] | None = None,
    composition_plan: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "agent_task_path": str(report_dir / "date_planning_agent_task.md"),
        "agent_output_path": str(report_dir / "date_planning_agent_output.json"),
        "raw_task": raw_task,
        "source_file_name": source_path.name,
        "source_file_path": repo_rel(source_path),
        "source_text_excerpt": source_text[:80_000],
        "date_mentions": date_mentions,
        "reference_payloads": reference_pack.get("reference_payloads") or {},
        "reference_text_excerpt": str(reference_pack.get("reference_text_excerpt") or "")[:20_000],
        "document_model": planner_safe_document_model(document_model or {}) if document_model else {},
        "composition_plan": composition_plan or {},
    }


def date_plan_blockers(date_plan: dict[str, Any]) -> list[str]:
    blockers = date_plan.get("blockers")
    if isinstance(blockers, list):
        readable: list[str] = []
        for item in blockers:
            blocker = human_readable_date_plan_blocker(item)
            if blocker:
                readable.append(blocker)
        return readable
    blocker = str(date_plan.get("blocker") or date_plan.get("error") or "").strip()
    return [blocker] if blocker else []


def human_readable_date_plan_blocker(item: Any) -> str:
    if not isinstance(item, dict):
        return str(item or "").strip()
    visible = str(item.get("visible_text") or "").strip()
    reason = str(item.get("reason") or item.get("summary") or item.get("message") or "").strip()
    needed = str(item.get("needed_to_unblock") or item.get("next_step") or "").strip()
    parts: list[str] = []
    if visible:
        parts.append(f"{visible}:")
    if reason:
        parts.append(reason)
    if needed:
        parts.append(f"Needed: {needed}")
    if parts:
        return " ".join(parts)
    compact = {
        key: value
        for key, value in item.items()
        if key in {"date_kind", "source_time", "confidence"} and str(value or "").strip()
    }
    return json.dumps(compact, ensure_ascii=False) if compact else ""


def _int_field(item: dict[str, Any], *keys: str) -> int | None:
    for key in keys:
        raw = item.get(key)
        if raw in (None, ""):
            continue
        try:
            return int(str(raw).strip())
        except ValueError:
            continue
    return None


def _date_from_parts(year: str, month: str, day: str) -> date | None:
    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def parse_date_value(text: str) -> date | None:
    value = re.sub(r"\s+", " ", str(text or "")).strip()
    if not value:
        return None
    compact_year = re.sub(r"\b(20)\s+(\d)\s+(\d)\b", r"\1\2\3", value)
    for match in re.finditer(r"\b(20\d{2})[-/](\d{1,2})[-/](\d{1,2})\b", compact_year):
        parsed = _date_from_parts(match.group(1), match.group(2), match.group(3))
        if parsed:
            return parsed
    for match in re.finditer(r"\b(\d{1,2})[-/](\d{1,2})[-/](20\d{2})\b", compact_year):
        first = int(match.group(1))
        second = int(match.group(2))
        if first > 12:
            parsed = _date_from_parts(match.group(3), match.group(2), match.group(1))
        else:
            parsed = _date_from_parts(match.group(3), match.group(1), match.group(2))
        if parsed:
            return parsed
    for match in re.finditer(r"\b(20\d{2})年(\d{1,2})月(\d{1,2})日?\b", value):
        parsed = _date_from_parts(match.group(1), match.group(2), match.group(3))
        if parsed:
            return parsed
    month_names = "|".join(sorted(MONTH_ALIASES, key=len, reverse=True))
    day_month = re.compile(
        rf"\b(\d{{1,2}})(?:st|nd|rd|th)?[,\-\s]+({month_names})\.?[,\-\s]+(20\d{{2}})\b",
        re.IGNORECASE,
    )
    for match in day_month.finditer(value):
        month = MONTH_ALIASES.get(match.group(2).casefold().rstrip("."))
        if month:
            parsed = _date_from_parts(match.group(3), str(month), match.group(1))
            if parsed:
                return parsed
    month_day = re.compile(
        rf"\b({month_names})\.?[,\s]+(\d{{1,2}})(?:st|nd|rd|th)?[,]?\s+(20\d{{2}})\b",
        re.IGNORECASE,
    )
    for match in month_day.finditer(value):
        month = MONTH_ALIASES.get(match.group(1).casefold().rstrip("."))
        if month:
            parsed = _date_from_parts(match.group(3), str(month), match.group(2))
            if parsed:
                return parsed
    return None


def normalize_time_value(text: str) -> str:
    match = re.search(r"\b([01]?\d|2[0-3]):([0-5]\d)\b", str(text or ""))
    if not match:
        return ""
    return f"{int(match.group(1)):02d}:{match.group(2)}"


def canonical_weekday_value(value: str) -> str:
    text = re.sub(r"[^A-Za-z\u4e00-\u9fff]+", "", str(value or "")).casefold()
    return WEEKDAY_ALIASES.get(text, "")


def weekday_from_date_or_text(*values: str, parsed_date: date | None = None) -> str:
    for value in values:
        text = str(value or "")
        for alias in sorted(WEEKDAY_ALIASES, key=len, reverse=True):
            if re.search(rf"(?<![A-Za-z]){re.escape(alias)}(?![A-Za-z])", text, re.IGNORECASE) or alias in text:
                weekday = WEEKDAY_ALIASES[alias]
                if weekday:
                    return weekday
    if parsed_date:
        return WEEKDAY_NAMES[parsed_date.weekday()]
    return ""


def reference_semester_schedule(reference_pack: dict[str, Any] | None) -> dict[str, Any]:
    payloads = (reference_pack or {}).get("reference_payloads")
    if not isinstance(payloads, dict):
        return {}
    schedule = payloads.get("semester_schedule.json")
    return schedule if isinstance(schedule, dict) else {}


def schedule_weeks_by_number(schedule: dict[str, Any]) -> dict[int, dict[str, Any]]:
    weeks = schedule.get("weeks")
    if not isinstance(weeks, list):
        return {}
    indexed: dict[int, dict[str, Any]] = {}
    for item in weeks:
        if not isinstance(item, dict):
            continue
        try:
            week = int(item.get("week"))
        except (TypeError, ValueError):
            continue
        dates = []
        raw_dates = item.get("teaching_dates") if isinstance(item.get("teaching_dates"), list) else []
        for raw_date in raw_dates:
            parsed = parse_date_value(str(raw_date))
            if parsed:
                dates.append(parsed)
        if dates:
            indexed[week] = {**item, "_parsed_teaching_dates": sorted(set(dates))}
    return indexed


def dates_between(start: date, end: date) -> list[date]:
    if end < start:
        return []
    return [start + timedelta(days=offset) for offset in range((end - start).days + 1)]


def calendar_first_semester_facts(calendar: dict[str, Any]) -> dict[str, Any]:
    semesters = calendar.get("semesters")
    if isinstance(semesters, dict) and isinstance(semesters.get("first"), dict):
        return semesters["first"]
    if isinstance(semesters, list):
        for item in semesters:
            if isinstance(item, dict) and str(item.get("name") or "").lower() == "first":
                return item
    return {}


def calendar_blocked_teaching_dates(reference_pack: dict[str, Any] | None) -> dict[date, str]:
    payloads = (reference_pack or {}).get("reference_payloads")
    if not isinstance(payloads, dict):
        return {}
    calendar = payloads.get("calendar.json") if isinstance(payloads.get("calendar.json"), dict) else {}
    blocked: dict[date, str] = {}
    for item in calendar.get("non_teaching_days") if isinstance(calendar.get("non_teaching_days"), list) else []:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        parsed = parse_date_value(str(item.get("date") or ""))
        if parsed:
            blocked[parsed] = str(item.get("name") or item.get("type") or "calendar non-teaching day")
    for item in calendar.get("teaching_suspensions") if isinstance(calendar.get("teaching_suspensions"), list) else []:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        start = parse_date_value(str(item.get("start") or ""))
        end = parse_date_value(str(item.get("end") or ""))
        if not start or not end:
            continue
        label = str(item.get("name") or item.get("type") or "calendar teaching suspension")
        for day in dates_between(start, end):
            blocked.setdefault(day, label)
    facts = calendar_first_semester_facts(calendar)
    legacy_periods = (
        ("Reading/Field Trip Week", facts.get("reading_start"), facts.get("reading_end")),
        ("Revision Period", facts.get("revision_start"), facts.get("revision_end")),
        ("Assessment Period", facts.get("assessment_start"), facts.get("assessment_end")),
        ("Class Suspension Period", facts.get("class_suspension_start"), facts.get("class_suspension_end")),
    )
    for label, raw_start, raw_end in legacy_periods:
        if not raw_start or not raw_end:
            continue
        start = parse_date_value(str(raw_start))
        end = parse_date_value(str(raw_end))
        if not start or not end:
            continue
        for day in dates_between(start, end):
            blocked.setdefault(day, label)
    return blocked


def calendar_teaching_suspension_ranges(
    reference_pack: dict[str, Any] | None,
) -> list[tuple[date, date]]:
    """Return grounded class-suspension ranges without interpreting their labels."""
    payloads = (reference_pack or {}).get("reference_payloads")
    if not isinstance(payloads, dict):
        return []
    calendar = payloads.get("calendar.json") if isinstance(payloads.get("calendar.json"), dict) else {}
    ranges: set[tuple[date, date]] = set()
    for item in calendar.get("teaching_suspensions") if isinstance(calendar.get("teaching_suspensions"), list) else []:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        start = parse_date_value(str(item.get("start") or ""))
        end = parse_date_value(str(item.get("end") or ""))
        if start and end and start <= end:
            ranges.add((start, end))
    facts = calendar_first_semester_facts(calendar)
    for raw_start, raw_end in (
        (facts.get("reading_start"), facts.get("reading_end")),
        (facts.get("revision_start"), facts.get("revision_end")),
        (facts.get("assessment_start"), facts.get("assessment_end")),
        (facts.get("class_suspension_start"), facts.get("class_suspension_end")),
    ):
        start = parse_date_value(str(raw_start or ""))
        end = parse_date_value(str(raw_end or ""))
        if start and end and start <= end:
            ranges.add((start, end))
    return sorted(ranges)


def canonical_non_teaching_gap_date(
    proposed: date,
    target_gaps: dict[str, dict[str, Any]],
    reference_pack: dict[str, Any] | None,
) -> tuple[str | None, str]:
    """Resolve a suspension-period date to the canonical course-session gap date.

    A planner may express a reading/suspension week by its range start while the
    rebuilt timeline records the recurring class date that was cancelled. Exact
    timeline dates remain preferred. Range normalization is accepted only when
    the calendar evidence identifies exactly one canonical gap, so multi-session
    or ambiguous weeks still fail closed.
    """
    proposed_iso = proposed.isoformat()
    if proposed_iso in target_gaps:
        return proposed_iso, "exact_timeline_gap_date"
    candidates: set[str] = set()
    for start, end in calendar_teaching_suspension_ranges(reference_pack):
        if not start <= proposed <= end:
            continue
        for target_iso in target_gaps:
            target_date = parse_date_value(target_iso)
            if target_date and start <= target_date <= end:
                candidates.add(target_iso)
    if len(candidates) == 1:
        return next(iter(candidates)), "calendar_suspension_to_course_session_date"
    return None, ""


def expected_weekday_date_for_schedule_week(week_payload: dict[str, Any], weekday: str) -> tuple[date | None, str]:
    weekday_index = WEEKDAY_NAMES.index(weekday) if weekday in WEEKDAY_NAMES else -1
    if weekday_index < 0:
        return None, f"unknown weekday {weekday!r}"
    candidates: set[date] = set()
    for teaching_date in week_payload.get("_parsed_teaching_dates") or []:
        week_start = teaching_date - timedelta(days=teaching_date.weekday())
        candidates.add(week_start + timedelta(days=weekday_index))
    if not candidates:
        return None, "schedule week has no teaching_dates anchor"
    if len(candidates) > 1:
        return None, "schedule week teaching_dates do not share one calendar week"
    return next(iter(candidates)), ""


def timeline_content_week_mappings(
    reference_pack: dict[str, Any] | None,
) -> list[dict[str, Any]]:
    timeline = reference_payload(reference_pack, "timeline.json")
    raw = timeline.get("content_week_mappings")
    if not isinstance(raw, list):
        schedule = reference_semester_schedule(reference_pack)
        raw = schedule.get("content_week_mappings")
    return [dict(item) for item in raw or [] if isinstance(item, dict)]


def content_kind_for_date_edit(edit: dict[str, Any]) -> str:
    explicit = str(edit.get("content_kind") or "").strip().lower()
    if explicit in {"lecture", "workshop", "submission"}:
        return explicit
    date_kind = str(edit.get("date_kind") or "").strip().lower()
    if any(token in date_kind for token in ("deadline", "submission", "quiz", "deliverable")):
        return "submission"
    return ""


def validate_content_week_mapping(
    edit: dict[str, Any],
    *,
    edit_index: int,
    reference_pack: dict[str, Any] | None,
) -> list[str]:
    source_week = _int_field(edit, "source_week")
    target_week = _int_field(edit, "target_week")
    if source_week is None or target_week is None:
        return []
    mappings = timeline_content_week_mappings(reference_pack)
    mapping_item_id = str(edit.get("mapping_item_id") or "").strip()
    content_kind = content_kind_for_date_edit(edit)
    if mapping_item_id:
        candidates = [
            item for item in mappings
            if str(item.get("item_id") or "") == mapping_item_id
        ]
        if len(candidates) != 1:
            return [
                f"date edit {edit_index} references unknown or duplicate content mapping item {mapping_item_id!r}"
            ]
    else:
        candidates = [
            item
            for item in mappings
            if _int_field(item, "source_teaching_week", "source_week") == source_week
            and (not content_kind or str(item.get("content_kind") or "") == content_kind)
        ]
    if source_week == target_week and not mapping_item_id:
        consequential = [
            item
            for item in candidates
            if str(item.get("status") or "").strip().lower() in {"moved", "deleted"}
            or _int_field(item, "target_teaching_week", "target_week") not in {source_week}
        ]
        if not consequential:
            return []
        candidates = consequential
    if not candidates:
        return [
            f"date edit {edit_index} changes Teaching Week {source_week} to {target_week} "
            "without an explicit content_week_mappings record"
        ]
    if len(candidates) != 1:
        return [
            f"date edit {edit_index} has multiple content mappings for Teaching Week {source_week}; "
            "mapping_item_id is required"
        ]
    mapping = candidates[0]
    status = str(mapping.get("status") or "").strip().lower()
    mapped_target = _int_field(mapping, "target_teaching_week", "target_week")
    if status == "deleted" or mapped_target is None:
        return [
            f"date edit {edit_index} refers to deleted content item {str(mapping.get('item_id') or '')!r}; "
            "the refresh must block instead of guessing a date"
        ]
    if (
        _int_field(mapping, "source_teaching_week", "source_week") != source_week
        or mapped_target != target_week
    ):
        return [
            f"date edit {edit_index} Teaching Week {source_week}->{target_week} does not match "
            f"content mapping {mapping.get('item_id')} "
            f"{mapping.get('source_teaching_week')}->{mapping.get('target_teaching_week')}"
        ]
    if source_week != target_week and status != "moved":
        return [
            f"date edit {edit_index} changes teaching weeks but content mapping "
            f"{mapping.get('item_id')} is not marked moved"
        ]
    if content_kind and str(mapping.get("content_kind") or "") != content_kind:
        return [
            f"date edit {edit_index} content_kind={content_kind} does not match "
            f"content mapping {mapping.get('item_id')}"
        ]
    return []


def validate_date_edits_against_schedule(
    edits: list[dict[str, Any]],
    *,
    reference_pack: dict[str, Any] | None,
) -> list[str]:
    errors: list[str] = []
    schedule_edits = [
        edit
        for edit in edits
        if str(edit.get("mapping_policy") or "").strip() in SCHEDULE_MAPPING_POLICIES
    ]
    if not schedule_edits:
        return []
    schedule = reference_semester_schedule(reference_pack)
    weeks_by_number = schedule_weeks_by_number(schedule)
    if not weeks_by_number:
        return ["semester_schedule.json is missing or incomplete for schedule-aware date edits"]
    blocked_teaching_dates = calendar_blocked_teaching_dates(reference_pack)
    all_teaching_dates = {
        teaching_date
        for week_payload in weeks_by_number.values()
        for teaching_date in week_payload.get("_parsed_teaching_dates", [])
    }
    for index, edit in enumerate(schedule_edits, start=1):
        policy = str(edit.get("mapping_policy") or "").strip()
        errors.extend(
            validate_content_week_mapping(
                edit,
                edit_index=index,
                reference_pack=reference_pack,
            )
        )
        target_date = parse_date_value(str(edit.get("target_date_iso") or edit.get("replacement_text") or ""))
        if not target_date:
            errors.append(f"date edit {index} with mapping_policy={policy} is missing a parseable target date")
            continue
        week = _int_field(edit, "target_week", "source_week", "week")
        if week is None:
            errors.append(f"date edit {index} with mapping_policy={policy} is missing structured week metadata")
            continue
        week_payload = weeks_by_number.get(week)
        if not week_payload:
            errors.append(f"date edit {index} targets Week {week}, but semester_schedule.json has no such week")
            continue
        if policy == "same_week_teaching_date":
            if target_date not in week_payload.get("_parsed_teaching_dates", []):
                allowed = ", ".join(item.isoformat() for item in week_payload.get("_parsed_teaching_dates", []))
                errors.append(
                    f"date edit {index} maps a class/session date to {target_date.isoformat()}, "
                    f"but Week {week} teaching_dates are {allowed or '(none)'}"
                )
            if target_date in blocked_teaching_dates:
                errors.append(
                    f"date edit {index} maps a class/session date to {target_date.isoformat()}, "
                    f"but calendar.json marks it as {blocked_teaching_dates[target_date]}"
                )
            continue
        if policy == "same_weekday_deadline":
            source_date = parse_date_value(str(edit.get("source_date_iso") or edit.get("visible_text") or ""))
            source_weekday = canonical_weekday_value(str(edit.get("source_weekday") or ""))
            if not source_weekday:
                source_weekday = weekday_from_date_or_text(str(edit.get("visible_text") or ""), parsed_date=source_date)
            if not source_weekday:
                errors.append(f"date edit {index} with mapping_policy={policy} is missing source_weekday/source_date metadata")
                continue
            expected_date, error = expected_weekday_date_for_schedule_week(week_payload, source_weekday)
            if error or not expected_date:
                errors.append(f"date edit {index} cannot derive Week {week} {source_weekday}: {error}")
                continue
            if target_date != expected_date:
                errors.append(
                    f"date edit {index} maps a deadline to {target_date.isoformat()}, "
                    f"but Week {week} {source_weekday} is {expected_date.isoformat()}"
                )
            source_time = normalize_time_value(str(edit.get("source_time") or edit.get("visible_text") or ""))
            replacement_time = normalize_time_value(str(edit.get("replacement_time") or edit.get("replacement_text") or ""))
            if source_time and replacement_time and source_time != replacement_time:
                errors.append(
                    f"date edit {index} changes deadline time from {source_time} to {replacement_time}"
                )
            if source_time and source_time in str(edit.get("visible_text") or "") and not replacement_time:
                errors.append(f"date edit {index} removes deadline time {source_time}")
            continue
        if target_date not in all_teaching_dates:
            errors.append(f"date edit {index} has unsupported schedule mapping_policy={policy}")
    return errors


def valid_date_edits_from_plan(
    date_plan: dict[str, Any],
    source_text: str,
    reference_pack: dict[str, Any] | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    reasons: list[str] = []
    if str(date_plan.get("agent") or "").strip() != DATE_PLANNER_AGENT_NAME:
        return [], ["date plan was not produced by the Date Planner Agent"]
    status = str(date_plan.get("status") or "").strip().lower()
    if status == "no_changes_needed":
        return [], []
    if status != "ready":
        return [], date_plan_blockers(date_plan) or ["date planner did not return a ready plan"]
    raw_edits = date_plan.get("date_edits")
    if not isinstance(raw_edits, list):
        return [], ["date planner output missing date_edits"]
    edits: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for index, item in enumerate(raw_edits, start=1):
        if not isinstance(item, dict):
            reasons.append(f"date edit {index} is not an object")
            continue
        old = str(item.get("visible_text") or "").strip()
        new = str(item.get("replacement_text") or "").strip()
        confidence = str(item.get("confidence") or "").strip().lower()
        basis = str(item.get("basis") or "").strip()
        if not old or not new or old == new:
            reasons.append(f"date edit {index} is missing distinct visible/replacement text")
            continue
        if old not in source_text:
            reasons.append(f"date edit {index} visible text is not present in the source")
            continue
        if not DATEISH_PATTERN.search(old):
            reasons.append(f"date edit {index} visible text does not look like a date")
            continue
        if confidence not in ACCEPTED_DATE_PLAN_CONFIDENCE:
            reasons.append(f"date edit {index} confidence is not high or medium")
            continue
        if not basis:
            reasons.append(f"date edit {index} is missing grounding basis")
            continue
        key = (old, new)
        if key in seen:
            continue
        seen.add(key)
        edits.append(
            {
                "id": str(item.get("id") or f"date-{index}"),
                "target_id": str(item.get("target_id") or ""),
                "target_fingerprint": str(item.get("target_fingerprint") or ""),
                "visible_text": old,
                "replacement_text": new,
                "date_kind": str(item.get("date_kind") or "other"),
                "content_kind": str(item.get("content_kind") or "").strip(),
                "mapping_item_id": str(item.get("mapping_item_id") or "").strip(),
                "mapping_policy": str(item.get("mapping_policy") or "").strip(),
                "source_week": item.get("source_week"),
                "target_week": item.get("target_week"),
                "source_weekday": item.get("source_weekday"),
                "source_date_iso": item.get("source_date_iso"),
                "target_date_iso": item.get("target_date_iso"),
                "source_time": item.get("source_time"),
                "replacement_time": item.get("replacement_time"),
                "basis": basis,
                "confidence": confidence,
                "requires_human_review": bool(item.get("requires_human_review", True)),
            }
        )
    schedule_errors = validate_date_edits_against_schedule(edits, reference_pack=reference_pack)
    if schedule_errors:
        reasons.extend(schedule_errors)
        return [], reasons
    return edits, reasons


def composition_date_plan_errors(
    date_edits: list[dict[str, Any]],
    *,
    model: dict[str, Any],
    composition_plan: dict[str, Any],
    schedule_table_rebuilds: list[dict[str, Any]],
) -> list[str]:
    """Validate anchored date edits for the composition-enabled date lane.

    Schedule-table rebuilds use the separately validated legacy DOCX lane. They
    remain bounded by an exact header match, the expected source row count, and
    the grounded target rows produced from the reference pack.
    """
    errors: list[str] = []
    assigned = _assigned_target_ids(composition_plan, "date")
    targets = structure_index(model)
    seen: set[tuple[str, str]] = set()
    for index, edit in enumerate(date_edits, start=1):
        target_id = str(edit.get("target_id") or "")
        target = targets.get(target_id)
        if target_id not in assigned:
            errors.append(f"date edit {index} targets a structure not assigned to Date Planner")
            continue
        if target is None:
            errors.append(f"date edit {index} target does not resolve exactly")
            continue
        if target.get("kind") not in {"paragraph", "text_box", "table_cell", "worksheet_cell"}:
            errors.append(f"date edit {index} target is not an indivisible text-bearing structure")
        if str(edit.get("target_fingerprint") or "") != str(target.get("fingerprint") or ""):
            errors.append(f"date edit {index} target fingerprint does not match")
        current = str(target.get("text") if "text" in target else target.get("value") or "")
        visible = str(edit.get("visible_text") or "")
        if visible not in current:
            errors.append(f"date edit {index} expected visible text is not in its assigned structure")
        key = (target_id, visible)
        if key in seen:
            errors.append(f"date edit {index} duplicates an anchored target/value pair")
        seen.add(key)
    return errors


def composed_document_change_plan(
    *,
    model: dict[str, Any],
    group_plan: dict[str, Any],
    date_edits: list[dict[str, Any]],
    group_snapshot_hash: str,
) -> dict[str, Any]:
    operations = [
        dict(item)
        for item in group_plan.get("operations") or []
        if isinstance(item, dict)
    ]
    operations.extend(
        {
            "id": str(edit.get("id") or "date-edit"),
            "type": "anchored_text_edit",
            "target_id": str(edit.get("target_id") or ""),
            "target_fingerprint": str(edit.get("target_fingerprint") or ""),
            "expected_before": str(edit.get("visible_text") or ""),
            "replacement_text": str(edit.get("replacement_text") or ""),
            "evidence": [
                {
                    "source_ref": "date_planner",
                    "fact": str(edit.get("basis") or ""),
                }
            ],
            "confidence": str(edit.get("confidence") or "").lower(),
            "readback": {
                "kind": "text",
                "expected_text": str(edit.get("replacement_text") or ""),
            },
        }
        for edit in date_edits
    )
    return {
        "agent": "file-change-composition-planner",
        "status": "ready",
        "summary": "mechanically combined validated specialist operations",
        "source_document_fingerprint": str(model.get("document_fingerprint") or ""),
        "group_snapshot_hash": group_snapshot_hash,
        "operations": operations,
        "preserve": [],
        "blockers": [],
    }


def reference_payload(reference_pack: dict[str, Any] | None, filename: str) -> dict[str, Any]:
    payloads = (reference_pack or {}).get("reference_payloads")
    if not isinstance(payloads, dict):
        return {}
    payload = payloads.get(filename)
    return payload if isinstance(payload, dict) else {}


def schedule_calendar_week_origin(
    reference_pack: dict[str, Any] | None,
    timeline_rows: list[dict[str, Any]],
    raw_gaps: list[dict[str, Any]],
) -> date | None:
    calendar = reference_payload(reference_pack, "calendar.json")
    semesters = calendar.get("semesters")
    first = semesters.get("first") if isinstance(semesters, dict) else None
    if isinstance(first, dict):
        parsed = parse_date_value(str(first.get("first_teaching") or first.get("start") or ""))
        if parsed:
            return parsed
    candidates: list[date] = []
    for item in [*timeline_rows, *raw_gaps]:
        if not isinstance(item, dict):
            continue
        parsed = parse_date_value(str(item.get("date") or ""))
        if parsed:
            candidates.append(parsed)
    return min(candidates) if candidates else None


def calendar_week_number_for_date(value: date, origin: date | None) -> int | None:
    if origin is None or value < origin:
        return None
    return ((value - origin).days // 7) + 1


def valid_schedule_table_rebuilds_from_plan(
    date_plan: dict[str, Any],
    reference_pack: dict[str, Any] | None,
) -> tuple[list[dict[str, Any]], list[str]]:
    raw_plans = date_plan.get("schedule_table_rebuilds")
    if raw_plans in (None, ""):
        return [], []
    if str(date_plan.get("agent") or "").strip() != DATE_PLANNER_AGENT_NAME:
        return [], ["schedule table plan was not produced by the Date Planner Agent"]
    if str(date_plan.get("status") or "").strip().lower() != "ready":
        return [], date_plan_blockers(date_plan) or ["date planner did not return a ready schedule table plan"]
    if not isinstance(raw_plans, list):
        return [], ["date planner schedule_table_rebuilds must be a list"]

    timeline = reference_payload(reference_pack, "timeline.json")
    timeline_rows = timeline.get("week_calendar") if isinstance(timeline.get("week_calendar"), list) else []
    if not timeline_rows:
        return [], ["timeline.json is missing a complete teaching week_calendar for schedule table rebuild"]
    schedule_weeks = schedule_weeks_by_number(reference_semester_schedule(reference_pack))
    if not schedule_weeks:
        return [], ["semester_schedule.json is missing or incomplete for schedule table rebuild"]
    blocked_dates = calendar_blocked_teaching_dates(reference_pack)
    raw_gaps = timeline.get("non_teaching_calendar_gaps")
    gap_rows = [item for item in raw_gaps if isinstance(item, dict)] if isinstance(raw_gaps, list) else []
    calendar_week_origin = schedule_calendar_week_origin(reference_pack, timeline_rows, gap_rows)

    teaching_rows: list[dict[str, Any]] = []
    teaching_weeks: list[int] = []
    for index, item in enumerate(timeline_rows, start=1):
        if not isinstance(item, dict):
            return [], [f"timeline.json week_calendar row {index} is not an object"]
        week = _int_field(item, "teaching_week", "week")
        calendar_week = _int_field(item, "calendar_week")
        source_week = _int_field(item, "source_week")
        target_date = parse_date_value(str(item.get("date") or ""))
        if week is None or source_week is None or target_date is None:
            return [], [f"timeline.json week_calendar row {index} is missing week, source_week, or ISO date"]
        derived_calendar_week = calendar_week_number_for_date(target_date, calendar_week_origin)
        if calendar_week is None:
            calendar_week = derived_calendar_week
        if calendar_week is None:
            return [], [f"timeline.json Week {week} cannot be assigned a calendar_week"]
        if derived_calendar_week is not None and calendar_week != derived_calendar_week:
            return [], [
                f"timeline.json Teaching Week {week} says Calendar Week {calendar_week}, "
                f"but {target_date.isoformat()} falls in Calendar Week {derived_calendar_week}"
            ]
        schedule_week = schedule_weeks.get(week)
        if not schedule_week:
            return [], [f"semester_schedule.json has no Week {week} required by timeline.json"]
        if target_date not in schedule_week.get("_parsed_teaching_dates", []):
            return [], [
                f"timeline.json Week {week} date {target_date.isoformat()} is not a teaching_date "
                "in semester_schedule.json"
            ]
        if target_date in blocked_dates:
            return [], [
                f"timeline.json Week {week} date {target_date.isoformat()} is blocked by calendar.json "
                f"as {blocked_dates[target_date]}"
            ]
        schedule_calendar_week = _int_field(schedule_week, "calendar_week")
        if schedule_calendar_week is not None and schedule_calendar_week != calendar_week:
            return [], [
                f"semester_schedule.json Teaching Week {week} says Calendar Week {schedule_calendar_week}, "
                f"but timeline.json says Calendar Week {calendar_week}"
            ]
        teaching_weeks.append(week)
        teaching_rows.append(
            {
                "row_type": "teaching",
                "week": week,
                "teaching_week": week,
                "calendar_week": calendar_week,
                "source_week": source_week,
                "target_date_iso": target_date.isoformat(),
                "lecture": str(item.get("lecture") or ""),
                "workshop": str(item.get("workshop") or ""),
            }
        )
    expected_weeks = list(range(1, len(teaching_rows) + 1))
    if teaching_weeks != expected_weeks:
        return [], [
            "timeline.json teaching week_calendar must be ordered and numbered consecutively "
            f"from 1 to {len(teaching_rows)}"
        ]
    try:
        declared_week_count = int(timeline.get("week_count"))
    except (TypeError, ValueError):
        declared_week_count = len(teaching_rows)
    if declared_week_count != len(teaching_rows):
        return [], ["timeline.json week_count does not match week_calendar"]

    excluded_rows = timeline.get("excluded_source_rows")
    excluded_by_week: dict[int, dict[str, Any]] = {}
    if isinstance(excluded_rows, list):
        for item in excluded_rows:
            if not isinstance(item, dict):
                continue
            source_week = _int_field(item, "source_week")
            if source_week is not None:
                excluded_by_week[source_week] = item
    target_gaps: dict[str, dict[str, Any]] = {}
    if isinstance(raw_gaps, list):
        for item in raw_gaps:
            if not isinstance(item, dict):
                continue
            target_date = parse_date_value(str(item.get("date") or ""))
            label = str(item.get("reason") or "").strip()
            if not target_date or not label:
                return [], ["timeline.json has an incomplete non_teaching_calendar_gaps row"]
            if target_date not in blocked_dates:
                return [], [
                    f"timeline.json non-teaching gap {target_date.isoformat()} is not grounded "
                    "in a classes_suspended calendar period"
                ]
            calendar_week = _int_field(item, "calendar_week")
            derived_calendar_week = calendar_week_number_for_date(target_date, calendar_week_origin)
            if calendar_week is None:
                calendar_week = derived_calendar_week
            if calendar_week is None:
                return [], [
                    f"timeline.json non-teaching gap {target_date.isoformat()} cannot be assigned a calendar_week"
                ]
            if derived_calendar_week is not None and calendar_week != derived_calendar_week:
                return [], [
                    f"timeline.json non-teaching gap {target_date.isoformat()} says Calendar Week "
                    f"{calendar_week}, expected {derived_calendar_week}"
                ]
            target_gaps[target_date.isoformat()] = {
                "target_date_iso": target_date.isoformat(),
                "calendar_week": calendar_week,
                "label": label,
            }

    plans: list[dict[str, Any]] = []
    reasons: list[str] = []
    for index, item in enumerate(raw_plans, start=1):
        if not isinstance(item, dict):
            reasons.append(f"schedule table rebuild {index} is not an object")
            continue
        policy = str(item.get("mapping_policy") or "").strip()
        headers = item.get("table_headers")
        basis = str(item.get("basis") or "").strip()
        confidence = str(item.get("confidence") or "").strip().lower()
        display_week_field = str(item.get("display_week_field") or "").strip()
        source_data_row_count = _int_field(item, "source_data_row_count")
        if policy != SCHEDULE_TABLE_REBUILD_POLICY:
            reasons.append(f"schedule table rebuild {index} has unsupported mapping_policy={policy or '(missing)'}")
            continue
        if not isinstance(headers, list) or len(headers) < 2 or any(not str(value or "").strip() for value in headers):
            reasons.append(f"schedule table rebuild {index} requires exact table_headers")
            continue
        if source_data_row_count is None or source_data_row_count < 1:
            reasons.append(f"schedule table rebuild {index} requires source_data_row_count")
            continue
        if confidence not in ACCEPTED_DATE_PLAN_CONFIDENCE:
            reasons.append(f"schedule table rebuild {index} confidence is not high or medium")
            continue
        if not basis:
            reasons.append(f"schedule table rebuild {index} is missing grounding basis")
            continue
        if display_week_field not in {"calendar_week", "teaching_week"}:
            reasons.append(
                f"schedule table rebuild {index} requires display_week_field=calendar_week or teaching_week"
            )
            continue

        raw_non_teaching = item.get("non_teaching_rows")
        if raw_non_teaching in (None, ""):
            raw_non_teaching = []
        if not isinstance(raw_non_teaching, list):
            reasons.append(f"schedule table rebuild {index} non_teaching_rows must be a list")
            continue
        mapped_gaps: dict[str, dict[str, Any]] = {}
        invalid_gap = False
        for gap_index, mapping in enumerate(raw_non_teaching, start=1):
            if not isinstance(mapping, dict):
                reasons.append(f"schedule table rebuild {index} non-teaching row {gap_index} is not an object")
                invalid_gap = True
                continue
            source_week = _int_field(mapping, "source_week")
            target_date = parse_date_value(str(mapping.get("target_date_iso") or ""))
            if source_week is None or source_week not in excluded_by_week:
                reasons.append(
                    f"schedule table rebuild {index} non-teaching row {gap_index} "
                    "must select a source_week from timeline.json excluded_source_rows"
                )
                invalid_gap = True
                continue
            canonical_target_iso = None
            date_resolution = ""
            if target_date:
                canonical_target_iso, date_resolution = canonical_non_teaching_gap_date(
                    target_date,
                    target_gaps,
                    reference_pack,
                )
            if not target_date or not canonical_target_iso:
                reasons.append(
                    f"schedule table rebuild {index} non-teaching row {gap_index} "
                    "must target a timeline.json non_teaching_calendar_gaps date or an "
                    "unambiguous date inside the same grounded class-suspension period"
                )
                invalid_gap = True
                continue
            if canonical_target_iso in mapped_gaps:
                reasons.append(
                    f"schedule table rebuild {index} maps target non-teaching date "
                    f"{canonical_target_iso} more than once"
                )
                invalid_gap = True
                continue
            mapped_gaps[canonical_target_iso] = {
                "source_week": source_week,
                "planner_target_date_iso": target_date.isoformat(),
                "date_resolution": date_resolution,
            }
        if invalid_gap:
            continue
        if set(mapped_gaps) != set(target_gaps):
            missing = sorted(set(target_gaps) - set(mapped_gaps))
            extra = sorted(set(mapped_gaps) - set(target_gaps))
            reasons.append(
                f"schedule table rebuild {index} must cover every target non-teaching gap exactly once"
                f"; missing={missing}; extra={extra}"
            )
            continue

        target_rows = list(teaching_rows)
        for target_date_iso, gap in target_gaps.items():
            target_rows.append(
                {
                    "row_type": "non_teaching",
                    "week": None,
                    "teaching_week": None,
                    "calendar_week": gap["calendar_week"],
                    "source_week": mapped_gaps[target_date_iso]["source_week"],
                    "target_date_iso": target_date_iso,
                    "planner_target_date_iso": mapped_gaps[target_date_iso]["planner_target_date_iso"],
                    "date_resolution": mapped_gaps[target_date_iso]["date_resolution"],
                    "lecture": gap["label"],
                    "workshop": gap["label"],
                }
            )
        target_rows.sort(key=lambda row: str(row["target_date_iso"]))
        if display_week_field == "calendar_week":
            visible_weeks = [int(row["calendar_week"]) for row in target_rows]
            expected_visible_weeks = list(
                range(min(visible_weeks), max(visible_weeks) + 1)
            )
            if visible_weeks != expected_visible_weeks:
                reasons.append(
                    f"schedule table rebuild {index} visible calendar weeks must be continuous; "
                    f"got={visible_weeks}"
                )
                continue
        for row in target_rows:
            row["display_week"] = (
                row.get("calendar_week")
                if display_week_field == "calendar_week"
                else row.get("teaching_week")
            )
        plans.append(
            {
                "id": str(item.get("id") or f"schedule-table-{index}"),
                "table_headers": [str(value).strip() for value in headers],
                "source_data_row_count": source_data_row_count,
                "mapping_policy": policy,
                "display_week_field": display_week_field,
                "basis": basis,
                "confidence": confidence,
                "requires_human_review": bool(item.get("requires_human_review", True)),
                "target_rows": target_rows,
            }
        )
    if reasons:
        return [], reasons
    return plans, []


def valid_unchanged_dates_from_plan(date_plan: dict[str, Any], source_text: str) -> tuple[list[dict[str, Any]], list[str]]:
    reasons: list[str] = []
    if str(date_plan.get("agent") or "").strip() != DATE_PLANNER_AGENT_NAME:
        return [], ["date plan was not produced by the Date Planner Agent"]
    if str(date_plan.get("status") or "").strip().lower() not in {"ready", "no_changes_needed"}:
        return [], date_plan_blockers(date_plan) or ["date planner did not return a ready/no_changes_needed plan"]
    raw_items = date_plan.get("unchanged_dates")
    if raw_items in (None, ""):
        return [], []
    if not isinstance(raw_items, list):
        return [], ["date planner unchanged_dates must be a list"]
    unchanged: list[dict[str, Any]] = []
    seen: set[str] = set()
    for index, item in enumerate(raw_items, start=1):
        if not isinstance(item, dict):
            reasons.append(f"unchanged date {index} is not an object")
            continue
        visible = str(item.get("visible_text") or "").strip()
        reason = str(item.get("reason") or item.get("basis") or "").strip()
        confidence = str(item.get("confidence") or "").strip().lower()
        if not visible or visible not in source_text:
            reasons.append(f"unchanged date {index} visible text is not present in the source")
            continue
        if not DATEISH_PATTERN.search(visible):
            reasons.append(f"unchanged date {index} visible text does not look like a date")
            continue
        if confidence not in ACCEPTED_DATE_PLAN_CONFIDENCE:
            reasons.append(f"unchanged date {index} confidence is not high or medium")
            continue
        if not reason:
            reasons.append(f"unchanged date {index} is missing reason")
            continue
        key = visible.casefold()
        if key in seen:
            continue
        seen.add(key)
        unchanged.append(
            {
                "visible_text": visible,
                "reason": reason,
                "confidence": confidence,
                "date_kind": str(item.get("date_kind") or "other"),
                "requires_human_review": bool(item.get("requires_human_review", False)),
            }
        )
    return unchanged, reasons


def stale_mentions_uncovered_by_unchanged(text: str, unchanged_dates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    accepted = [
        str(item.get("visible_text") or "").strip()
        for item in unchanged_dates
        if str(item.get("visible_text") or "").strip() and not bool(item.get("requires_human_review"))
    ]
    uncovered: list[dict[str, Any]] = []
    for mention in extract_date_mentions(text):
        visible = str(mention.get("visible_text") or "")
        if not audit_visible_text(visible)["stale_date_hits"]:
            continue
        if any(visible == kept or visible in kept or kept in visible for kept in accepted):
            continue
        uncovered.append(mention)
    return uncovered


def audit_with_date_plan_coverage(raw_audit: dict[str, Any], output_text: str, unchanged_dates: list[dict[str, Any]]) -> dict[str, Any]:
    audit = dict(raw_audit)
    uncovered = stale_mentions_uncovered_by_unchanged(output_text, unchanged_dates)
    if raw_audit.get("stale_date_hits") and not uncovered:
        audit["stale_date_hits"] = []
        audit["date_plan_covered_stale_dates"] = True
    elif uncovered:
        audit["stale_date_hits"] = sorted({str(item.get("visible_text") or "") for item in uncovered if str(item.get("visible_text") or "")})[:30]
        audit["date_plan_covered_stale_dates"] = False
    trace_hits = audit.get("student_visible_trace_hits") if isinstance(audit.get("student_visible_trace_hits"), list) else []
    stale_hits = audit.get("stale_date_hits") if isinstance(audit.get("stale_date_hits"), list) else []
    audit["status"] = "PASS" if not stale_hits and not trace_hits else "NEEDS_REVIEW"
    return audit


def apply_date_edits_to_text(text: str, edits: list[dict[str, Any]]) -> tuple[str, dict[str, int], list[str]]:
    updated = text
    counts: dict[str, int] = {}
    applied_ids: list[str] = []
    for edit in edits:
        old = str(edit.get("visible_text") or "")
        new = str(edit.get("replacement_text") or "")
        count = updated.count(old)
        if not count:
            continue
        updated = updated.replace(old, new)
        counts[old] = counts.get(old, 0) + count
        applied_ids.append(str(edit.get("id") or old))
    return updated, counts, sorted(set(applied_ids))


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _ooxml_text_nodes(element: ET.Element) -> list[ET.Element]:
    return [
        node
        for node in element.iter()
        if _local_name(str(node.tag)) in OOXML_TEXT_NODE_NAMES
    ]


def _ooxml_blocks(root: ET.Element) -> list[ET.Element]:
    blocks = [
        node
        for node in root.iter()
        if _local_name(str(node.tag)) in OOXML_BLOCK_NODE_NAMES
    ]
    return blocks or [root]


def _replace_once_in_text_nodes(nodes: list[ET.Element], old: str, new: str) -> bool:
    parts = [node.text or "" for node in nodes]
    joined = "".join(parts)
    start = joined.find(old)
    if start < 0:
        return False
    end = start + len(old)
    ranges: list[tuple[int, int]] = []
    cursor = 0
    for part in parts:
        next_cursor = cursor + len(part)
        ranges.append((cursor, next_cursor))
        cursor = next_cursor
    start_index = next((index for index, (left, right) in enumerate(ranges) if left <= start < right), -1)
    end_index = next((index for index, (left, right) in enumerate(ranges) if left < end <= right), -1)
    if start_index < 0 or end_index < 0:
        return False
    start_left, _start_right = ranges[start_index]
    _end_left, end_right = ranges[end_index]
    prefix = parts[start_index][: start - start_left]
    suffix = parts[end_index][len(parts[end_index]) - (end_right - end) :]
    if start_index == end_index:
        nodes[start_index].text = prefix + new + suffix
        return True
    nodes[start_index].text = prefix + new
    for index in range(start_index + 1, end_index):
        nodes[index].text = ""
    nodes[end_index].text = suffix
    return True


def _replace_in_text_nodes(nodes: list[ET.Element], old: str, new: str) -> int:
    count = 0
    while _replace_once_in_text_nodes(nodes, old, new):
        count += 1
    return count


def apply_date_edits_to_ooxml(path: Path, edits: list[dict[str, Any]]) -> dict[str, Any]:
    if path.suffix.lower() not in OFFICE_SUFFIXES:
        return {"status": "not_ooxml", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
    if not edits:
        return {"status": "unchanged_no_trusted_date_plan", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
    counts: dict[str, int] = {}
    applied_ids: list[str] = []
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    try:
        with zipfile.ZipFile(path, "r") as src, zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                payload = src.read(info.filename)
                if info.filename.endswith(".xml"):
                    try:
                        root = ET.fromstring(payload)
                    except ET.ParseError:
                        root = None
                    if root is not None:
                        changed = False
                        for block in _ooxml_blocks(root):
                            nodes = _ooxml_text_nodes(block)
                            if not nodes:
                                continue
                            for edit in edits:
                                old = str(edit.get("visible_text") or "")
                                new = str(edit.get("replacement_text") or "")
                                count = _replace_in_text_nodes(nodes, old, new)
                                if not count:
                                    continue
                                changed = True
                                counts[old] = counts.get(old, 0) + count
                                applied_ids.append(str(edit.get("id") or old))
                        if changed:
                            payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                dst.writestr(info, payload)
        tmp_path.replace(path)
    except zipfile.BadZipFile as exc:
        if tmp_path.exists():
            tmp_path.unlink()
        return {
            "status": "blocked_bad_ooxml_zip",
            "replacement_count": 0,
            "date_edits": {},
            "applied_date_edit_ids": [],
            "blocker": str(exc),
        }
    replacement_count = sum(counts.values())
    return {
        "status": "modified_by_date_plan" if replacement_count else "unchanged_no_matching_date_edits",
        "replacement_count": replacement_count,
        "date_edits": counts,
        "applied_date_edit_ids": sorted(set(applied_ids)),
    }


def _normalized_table_cell_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _display_month_day(value: str) -> str:
    parsed = parse_date_value(value)
    if not parsed:
        return value
    return f"{parsed.strftime('%b')} {parsed.day}"


def _set_docx_cell_text(cell: Any, value: str) -> None:
    paragraphs = list(cell.paragraphs)
    paragraph = paragraphs[0] if paragraphs else cell.add_paragraph()
    paragraph.text = str(value or "")
    for extra in paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def apply_schedule_table_rebuilds_to_docx(path: Path, plans: list[dict[str, Any]]) -> dict[str, Any]:
    if not plans:
        return {
            "status": "unchanged_no_schedule_table_plan",
            "replacement_count": 0,
            "schedule_table_rebuilds": [],
        }
    if path.suffix.lower() != ".docx":
        return {
            "status": "blocked_schedule_table_requires_docx",
            "replacement_count": 0,
            "schedule_table_rebuilds": [],
            "blocker": "teaching_week_schedule_rebuild currently requires an editable DOCX source",
        }
    try:
        from docx import Document

        document = Document(path)
    except Exception as exc:
        return {
            "status": "blocked_invalid_docx",
            "replacement_count": 0,
            "schedule_table_rebuilds": [],
            "blocker": f"could not open DOCX schedule table: {exc}",
        }

    applied: list[dict[str, Any]] = []
    for plan in plans:
        expected_headers = [
            _normalized_table_cell_text(value)
            for value in plan.get("table_headers") or []
        ]
        matches = []
        for table in document.tables:
            if not table.rows:
                continue
            actual_headers = [
                _normalized_table_cell_text(cell.text)
                for cell in table.rows[0].cells[: len(expected_headers)]
            ]
            if actual_headers == expected_headers:
                matches.append(table)
        if len(matches) != 1:
            return {
                "status": "blocked_schedule_table_match",
                "replacement_count": 0,
                "schedule_table_rebuilds": applied,
                "blocker": (
                    f"schedule table plan {plan.get('id') or '(unnamed)'} matched "
                    f"{len(matches)} DOCX tables; expected exactly one"
                ),
            }
        table = matches[0]
        data_rows = list(table.rows[1:])
        expected_row_count = int(plan.get("source_data_row_count") or 0)
        if len(data_rows) != expected_row_count:
            return {
                "status": "blocked_schedule_table_row_count",
                "replacement_count": 0,
                "schedule_table_rebuilds": applied,
                "blocker": (
                    f"schedule table plan {plan.get('id') or '(unnamed)'} expected "
                    f"{expected_row_count} source rows but found {len(data_rows)}"
                ),
            }
        source_rows_by_week: dict[int, Any] = {}
        for row in data_rows:
            week_text = re.sub(r"\s+", "", row.cells[0].text)
            if not re.fullmatch(r"\d+", week_text):
                continue
            week = int(week_text)
            if week in source_rows_by_week:
                return {
                    "status": "blocked_duplicate_source_week",
                    "replacement_count": 0,
                    "schedule_table_rebuilds": applied,
                    "blocker": f"schedule table contains duplicate source Week {week}",
                }
            source_rows_by_week[week] = deepcopy(row._tr)

        target_rows = plan.get("target_rows") if isinstance(plan.get("target_rows"), list) else []
        missing_source_weeks = sorted(
            {
                int(row["source_week"])
                for row in target_rows
                if isinstance(row, dict)
                and _int_field(row, "source_week") is not None
                and int(row["source_week"]) not in source_rows_by_week
            }
        )
        if missing_source_weeks:
            return {
                "status": "blocked_missing_source_week_rows",
                "replacement_count": 0,
                "schedule_table_rebuilds": applied,
                "blocker": f"schedule table is missing source rows for Weeks {missing_source_weeks}",
            }

        for row in data_rows:
            table._tbl.remove(row._tr)
        for target in target_rows:
            source_week = int(target["source_week"])
            table._tbl.append(deepcopy(source_rows_by_week[source_week]))
            output_row = table.rows[-1]
            if len(output_row.cells) < 4:
                return {
                    "status": "blocked_schedule_table_columns",
                    "replacement_count": 0,
                    "schedule_table_rebuilds": applied,
                    "blocker": "schedule table has fewer than four columns",
                }
            week_value = str(target.get("display_week") or "")
            _set_docx_cell_text(output_row.cells[0], week_value)
            _set_docx_cell_text(output_row.cells[1], _display_month_day(str(target.get("target_date_iso") or "")))
            if target.get("row_type") == "non_teaching":
                _set_docx_cell_text(output_row.cells[2], str(target.get("lecture") or ""))
                _set_docx_cell_text(output_row.cells[3], str(target.get("workshop") or ""))
        applied.append(
            {
                "id": str(plan.get("id") or ""),
                "source_data_row_count": len(data_rows),
                "target_data_row_count": len(target_rows),
            }
        )
    try:
        document.save(path)
    except Exception as exc:
        return {
            "status": "blocked_schedule_table_save",
            "replacement_count": 0,
            "schedule_table_rebuilds": [],
            "blocker": f"could not save rebuilt DOCX schedule table: {exc}",
        }
    return {
        "status": "modified_by_schedule_table_plan",
        "replacement_count": sum(int(item["target_data_row_count"]) for item in applied),
        "schedule_table_rebuilds": applied,
    }


def is_visible_ooxml_member(filename: str, suffix: str) -> bool:
    if suffix == ".docx":
        return (
            filename == "word/document.xml"
            or filename.startswith("word/header")
            or filename.startswith("word/footer")
            or filename in {"word/footnotes.xml", "word/endnotes.xml"}
        )
    if suffix == ".pptx":
        return filename.startswith("ppt/slides/slide") and filename.endswith(".xml")
    if suffix == ".xlsx":
        return filename == "xl/sharedStrings.xml" or (filename.startswith("xl/worksheets/sheet") and filename.endswith(".xml"))
    return filename.endswith(".xml")


def office_text(path: Path) -> str:
    try:
        with zipfile.ZipFile(path, "r") as archive:
            parts: list[str] = []
            for info in archive.infolist():
                if not is_visible_ooxml_member(info.filename, path.suffix.lower()):
                    continue
                payload = archive.read(info.filename)
                try:
                    root = ET.fromstring(payload)
                except ET.ParseError:
                    continue
                for block in _ooxml_blocks(root):
                    text = "".join(node.text or "" for node in _ooxml_text_nodes(block)).strip()
                    if text:
                        parts.append(text)
            return "\n".join(parts)
    except zipfile.BadZipFile:
        return ""


def visible_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return read_text_lossy(path, 200_000)
    if suffix in OFFICE_SUFFIXES:
        return office_text(path)[:200_000]
    if suffix == ".pdf":
        try:
            import pdfplumber  # type: ignore

            with pdfplumber.open(path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages[:20])[:200_000]
        except Exception:
            return ""
    return ""


def audit_visible_text(text: str) -> dict[str, Any]:
    stale_hits: list[str] = []
    for pattern in STALE_DATE_PATTERNS:
        stale_hits.extend(pattern.findall(text))
    trace_hits = [marker for marker in STUDENT_VISIBLE_TRACE_MARKERS if marker.casefold() in text.casefold()]
    stale_unique = sorted(set(str(hit) for hit in stale_hits if str(hit).strip()))
    trace_unique = sorted(set(trace_hits))
    status = "PASS" if not stale_unique and not trace_unique else "NEEDS_REVIEW"
    return {
        "status": status,
        "stale_date_hits": stale_unique[:30],
        "student_visible_trace_hits": trace_unique[:30],
    }


def soffice_candidates(env: dict[str, str]) -> list[str]:
    candidates = []
    configured = env.get("SOFFICE_BIN", "").strip() or env.get("LIBREOFFICE_BIN", "").strip()
    if configured:
        candidates.append(configured)
    bundled = [
        Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/bin/soffice",
        Path.home()
        / ".cache/codex-runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice",
        Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        Path(os.environ.get("PROGRAMFILES", "")) / "LibreOffice" / "program" / "soffice.exe",
        Path(os.environ.get("ProgramFiles(x86)", "")) / "LibreOffice" / "program" / "soffice.exe",
    ]
    found = shutil.which("soffice")
    if found:
        candidates.append(found)
    candidates.extend(str(path) for path in bundled)
    unique: list[str] = []
    for candidate in candidates:
        if candidate and candidate not in unique:
            unique.append(candidate)
    return unique


def convert_to_pdf(source_path: Path, out_dir: Path, env: dict[str, str]) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    expected_pdf = out_dir / f"{source_path.stem}.pdf"
    expected_pdf.unlink(missing_ok=True)
    logs: list[str] = []
    for soffice in soffice_candidates(env):
        if not Path(soffice).exists() and shutil.which(soffice) is None:
            continue
        with tempfile.TemporaryDirectory(prefix="codex-refresh-preview-soffice-") as profile_dir:
            proc_env = os.environ.copy()
            proc_env.setdefault("XDG_CONFIG_HOME", str(Path(profile_dir) / "xdg_config"))
            proc_env.setdefault("XDG_CACHE_HOME", str(Path(profile_dir) / "xdg_cache"))
            Path(proc_env["XDG_CONFIG_HOME"]).mkdir(parents=True, exist_ok=True)
            Path(proc_env["XDG_CACHE_HOME"]).mkdir(parents=True, exist_ok=True)
            command = [
                soffice,
                f"-env:UserInstallation={Path(profile_dir).resolve().as_uri()}",
                "--invisible",
                "--headless",
                "--norestore",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                str(source_path),
            ]
            proc = subprocess.run(
                command,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=proc_env,
                check=False,
            )
            logs.append(f"EXIT {proc.returncode}: {proc.stdout.strip()} {proc.stderr.strip()}".strip())
            if proc.returncode == 0 and expected_pdf.exists() and expected_pdf.stat().st_size > 0:
                return {"ok": True, "pdf_path": repo_rel(expected_pdf), "backend": "libreoffice", "log": logs[-1]}

    if os.name == "nt" and source_path.suffix.lower() in OFFICE_SUFFIXES:
        worker = ROOT / "windows_office_pdf_worker.py"
        timeout = max(30, int(env.get("WINDOWS_OFFICE_PDF_TIMEOUT_SECONDS", "180")))
        with tempfile.TemporaryDirectory(prefix="course-agent-office-pdf-") as stage_text:
            stage = Path(stage_text)
            staged_source = stage / f"source{source_path.suffix.lower()}"
            staged_pdf = stage / "output.pdf"
            shutil.copy2(source_path, staged_source)
            proc_env = os.environ.copy()
            proc_env["PYTHONUTF8"] = "1"
            proc_env["PYTHONIOENCODING"] = "utf-8"
            try:
                proc = subprocess.run(
                    [
                        sys.executable,
                        str(worker),
                        "--source",
                        str(staged_source),
                        "--output",
                        str(staged_pdf),
                    ],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    env=proc_env,
                    timeout=timeout,
                    check=False,
                )
            except subprocess.TimeoutExpired:
                logs.append(f"Microsoft Office conversion timed out after {timeout}s.")
            else:
                logs.append(f"OFFICE EXIT {proc.returncode}: {proc.stdout.strip()} {proc.stderr.strip()}".strip())
                if proc.returncode == 0 and staged_pdf.is_file() and staged_pdf.stat().st_size > 0:
                    shutil.copy2(staged_pdf, expected_pdf)
                    return {
                        "ok": True,
                        "pdf_path": repo_rel(expected_pdf),
                        "backend": "microsoft_office_com",
                        "log": logs[-1],
                    }
    return {
        "ok": False,
        "pdf_path": repo_rel(expected_pdf),
        "error": "No working Office-to-PDF backend was available.",
        "logs": logs,
    }


def output_filename(source_path: Path) -> str:
    stem = source_path.stem
    for marker in STUDENT_VISIBLE_TRACE_MARKERS:
        stem = re.sub(re.escape(marker), "", stem, flags=re.IGNORECASE)
    stem = safe_name(stem, source_path.stem)
    return f"{stem}{source_path.suffix.lower()}"


def write_review_report(path: Path, result: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# File Refresh Preview",
        "",
        f"- Status: `{result.get('status')}`",
        f"- Source: `{result.get('source_file_name')}`",
        f"- Refreshed file: `{result.get('refreshed_file_path', '')}`",
        f"- Upload-ready file: `{result.get('upload_ready_path', '')}`",
        f"- Reference pack: `{(result.get('reference_pack') or {}).get('root', '')}`",
        f"- Edits applied: `{result.get('replacement_count', 0)}`",
        "",
        "## Audit",
        "",
        f"- Audit status: `{(result.get('audit') or {}).get('status', '')}`",
        f"- Stale date hits: `{', '.join((result.get('audit') or {}).get('stale_date_hits', []))}`",
        f"- Trace hits: `{', '.join((result.get('audit') or {}).get('student_visible_trace_hits', []))}`",
        "",
        "## Notes",
        "",
        result.get("note", ""),
    ]
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def _planner_context(operation: dict[str, Any]) -> dict[str, Any]:
    params = operation.get("params") if isinstance(operation.get("params"), dict) else {}
    for value in (operation.get("planner_context"), params.get("planner_context")):
        if isinstance(value, dict):
            return dict(value)
    return {}


def _sanitized_group_snapshot(context: dict[str, Any]) -> dict[str, Any]:
    course_readback = context.get("course_readback") if isinstance(context.get("course_readback"), dict) else {}
    source = context.get("group_structure")
    if not isinstance(source, dict):
        source = course_readback.get("group_structure") if isinstance(course_readback.get("group_structure"), dict) else {}
    groups = []
    for item in source.get("groups") or []:
        if not isinstance(item, dict) or not str(item.get("name") or "").strip():
            continue
        groups.append(
            {
                "id": str(item.get("id") or ""),
                "name": str(item.get("name") or "").strip(),
                "idnumber": str(item.get("idnumber") or ""),
                "member_count": item.get("member_count"),
                "grouping_names": [str(value) for value in item.get("grouping_names") or []],
            }
        )
    groupings = []
    for item in source.get("groupings") or []:
        if not isinstance(item, dict) or not str(item.get("name") or "").strip():
            continue
        groupings.append(
            {
                "id": str(item.get("id") or ""),
                "name": str(item.get("name") or "").strip(),
                "group_names": [str(value) for value in item.get("group_names") or []],
            }
        )
    snapshot = {
        "status": str(source.get("status") or "unavailable"),
        "course_id": str(source.get("course_id") or course_readback.get("course_id") or ""),
        "groups": groups,
        "groupings": groupings,
        "choicegroup_activities": [
            {
                "title": str(item.get("title") or ""),
                "section": str(item.get("section") or ""),
                "selected_group_names": [str(value) for value in item.get("selected_group_names") or []],
                "maxanswers": item.get("maxanswers"),
            }
            for item in source.get("choicegroup_activities") or []
            if isinstance(item, dict)
        ],
    }
    snapshot["snapshot_hash"] = sanitized_group_snapshot_hash(snapshot)
    return snapshot


def _composition_plan_errors(plan: dict[str, Any], model: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if str(plan.get("agent") or "") != COMPOSITION_AGENT_NAME:
        errors.append("composition plan was not produced by the File Change Composition Planner")
    status = str(plan.get("status") or "").lower()
    if status not in {"ready", "blocked", "no_changes_needed"}:
        errors.append("composition plan has an unsupported status")
    if status == "blocked":
        if not _plan_blockers(plan):
            errors.append("blocked composition plan must include a blocker")
        return errors
    selected = plan.get("selected_specialists")
    if not isinstance(selected, list) or any(value not in {"date", "group_document"} for value in selected):
        errors.append("composition plan selected_specialists is invalid")
        selected = []
    if len(selected) != len(set(selected)):
        errors.append("composition plan selected_specialists contains duplicates")
    targets = structure_index(model)
    assigned: set[str] = set()
    for index, item in enumerate(plan.get("targets") or [], start=1):
        if not isinstance(item, dict):
            errors.append(f"composition target {index} must be an object")
            continue
        target_id = str(item.get("target_id") or "")
        target = targets.get(target_id)
        if target is None:
            errors.append(f"composition target {index} does not resolve: {target_id!r}")
            continue
        if str(item.get("target_fingerprint") or "") != str(target.get("fingerprint") or ""):
            errors.append(f"composition target {index} fingerprint does not match")
        specialist = str(item.get("specialist") or "")
        if specialist not in selected:
            errors.append(f"composition target {index} specialist was not selected")
        if not str(item.get("reason") or "").strip():
            errors.append(f"composition target {index} is missing a semantic reason")
        if target_id in assigned:
            errors.append(f"composition target is assigned more than once: {target_id}")
        assigned.add(target_id)
    preserved: set[str] = set()
    for index, item in enumerate(plan.get("preserve") or [], start=1):
        if not isinstance(item, dict):
            errors.append(f"composition preserve item {index} must be an object")
            continue
        target_id = str(item.get("target_id") or "")
        if target_id not in targets:
            errors.append(f"composition preserve item {index} does not resolve: {target_id!r}")
        if target_id in assigned:
            errors.append(f"composition structure cannot be both assigned and preserved: {target_id}")
        if not str(item.get("reason") or "").strip():
            errors.append(f"composition preserve item {index} is missing a reason")
        preserved.add(target_id)
    for index, item in enumerate(plan.get("resolved_shared_targets") or [], start=1):
        if not isinstance(item, dict):
            errors.append(f"resolved shared target {index} must be an object")
            continue
        target_id = str(item.get("target_id") or "")
        owner = str(item.get("owner") or "")
        assigned_owner = next(
            (
                str(target.get("specialist") or "")
                for target in plan.get("targets") or []
                if isinstance(target, dict) and str(target.get("target_id") or "") == target_id
            ),
            "",
        )
        if target_id not in targets or owner not in selected or assigned_owner != owner:
            errors.append(f"resolved shared target {index} must resolve to its single assigned owner")
        if not str(item.get("reason") or "").strip():
            errors.append(f"resolved shared target {index} is missing an ownership reason")
    if status == "ready" and selected and not assigned:
        errors.append("composition plan selected specialists without assigning any targets")
    if status == "no_changes_needed" and (selected or assigned):
        errors.append("no_changes_needed composition plan cannot select specialists or targets")
    if plan.get("preserve_unselected") is not True:
        errors.append("composition plan must explicitly preserve all unselected structures")
    if not str(plan.get("preserve_unselected_reason") or "").strip():
        errors.append("composition plan must explain why unselected structures are preserved")
    return errors


def _assigned_target_ids(plan: dict[str, Any], specialist: str) -> set[str]:
    return {
        str(item.get("target_id") or "")
        for item in plan.get("targets") or []
        if isinstance(item, dict) and str(item.get("specialist") or "") == specialist
    }


def _group_plan_errors(
    plan: dict[str, Any],
    *,
    model: dict[str, Any],
    composition_plan: dict[str, Any],
    group_snapshot: dict[str, Any],
) -> list[str]:
    errors: list[str] = []
    if str(plan.get("agent") or "") != GROUP_AGENT_NAME:
        errors.append("Group document plan was not produced by the Group-aware Document Planner")
    status = str(plan.get("status") or "").lower()
    if status not in {"ready", "blocked", "no_changes_needed"}:
        errors.append("Group document plan has an unsupported status")
    if status == "blocked":
        if not _plan_blockers(plan):
            errors.append("blocked Group document plan must include a blocker")
        return errors
    if str(plan.get("source_document_fingerprint") or "") != str(model.get("document_fingerprint") or ""):
        errors.append("Group document plan source fingerprint does not match")
    expected_snapshot_hash = str(group_snapshot.get("snapshot_hash") or "")
    if str(plan.get("group_snapshot_hash") or "") != expected_snapshot_hash:
        errors.append("Group document plan snapshot hash does not match current sanitized Groups")
    assigned = _assigned_target_ids(composition_plan, "group_document")
    classified: set[str] = set()
    for index, operation in enumerate(plan.get("operations") or [], start=1):
        if isinstance(operation, dict) and str(operation.get("target_id") or "") not in assigned:
            errors.append(f"Group document operation {index} targets a structure not assigned by Composition Planner")
        if isinstance(operation, dict):
            classified.add(str(operation.get("target_id") or ""))
    for index, item in enumerate(plan.get("preserve") or [], start=1):
        if not isinstance(item, dict):
            errors.append(f"Group document preserve item {index} must be an object")
            continue
        target_id = str(item.get("target_id") or "")
        if target_id not in assigned:
            errors.append(f"Group document preserve item {index} was not assigned by Composition Planner")
        if not str(item.get("reason") or "").strip():
            errors.append(f"Group document preserve item {index} is missing a reason")
        classified.add(target_id)
    if status == "ready":
        if not plan.get("operations"):
            errors.append("ready Group document plan must contain at least one operation")
        missing_classification = sorted(assigned - classified)
        if missing_classification:
            errors.append(
                "Group document plan did not classify every assigned structure: "
                + ", ".join(missing_classification[:5])
            )
        errors.extend(validate_document_change_plan(plan, model))
    return errors


def _plan_blockers(plan: dict[str, Any]) -> list[str]:
    blockers = plan.get("blockers")
    if isinstance(blockers, list):
        return [str(value) for value in blockers if str(value).strip()]
    return []


def run_file_change_planners(
    operation: dict[str, Any],
    env: dict[str, str],
    *,
    report_dir: Path,
    source_path: Path,
    raw_task: str,
    reference_pack: dict[str, Any],
) -> dict[str, Any]:
    context = _planner_context(operation)
    if not context or source_path.suffix.lower() not in OFFICE_SUFFIXES:
        return {"enabled": False}
    try:
        model = build_document_model(source_path)
    except Exception as exc:
        return {
            "enabled": True,
            "blocker": f"Could not build the generic Office document model: {exc}",
            "composition_plan": {},
            "group_plan": {},
        }
    safe_model = planner_safe_document_model(model)
    group_snapshot = _sanitized_group_snapshot(context)
    course_readback = context.get("course_readback") if isinstance(context.get("course_readback"), dict) else {}
    historical_evidence = context.get("historical_evidence") if isinstance(context.get("historical_evidence"), list) else []
    composition_task = {
        "agent_task_path": str(report_dir / "file_change_composition_task.md"),
        "agent_output_path": str(report_dir / "file_change_composition_output.json"),
        "raw_task": raw_task,
        "document_model": safe_model,
        "course_readback": course_readback,
        "reference_payloads": reference_pack.get("reference_payloads") or {},
        "historical_evidence": historical_evidence,
    }
    composition_plan = run_composition_planner(env, composition_task)
    composition_errors = _composition_plan_errors(composition_plan, model)
    if composition_errors:
        return {
            "enabled": True,
            "blocker": "File Change Composition Planner output failed validation: " + "; ".join(composition_errors[:5]),
            "composition_plan": composition_plan,
            "group_plan": {},
            "document_model": model,
            "group_snapshot": group_snapshot,
        }
    if str(composition_plan.get("status") or "").lower() == "blocked":
        return {
            "enabled": True,
            "blocker": "; ".join(_plan_blockers(composition_plan)[:5]) or "File Change Composition Planner blocked the refresh.",
            "composition_plan": composition_plan,
            "group_plan": {},
            "document_model": model,
            "group_snapshot": group_snapshot,
        }
    selected = set(composition_plan.get("selected_specialists") or [])
    group_plan: dict[str, Any] = {}
    if "group_document" in selected:
        group_task = {
            "agent_task_path": str(report_dir / "group_document_planner_task.md"),
            "agent_output_path": str(report_dir / "group_document_planner_output.json"),
            "raw_task": raw_task,
            "document_model": safe_model,
            "composition_plan": composition_plan,
            "group_snapshot": group_snapshot,
            "reference_payloads": reference_pack.get("reference_payloads") or {},
            "historical_evidence": historical_evidence,
        }
        group_plan = run_group_document_planner(env, group_task)
        group_errors = _group_plan_errors(
            group_plan,
            model=model,
            composition_plan=composition_plan,
            group_snapshot=group_snapshot,
        )
        if group_errors:
            return {
                "enabled": True,
                "blocker": "Group-aware Document Planner output failed validation: " + "; ".join(group_errors[:5]),
                "composition_plan": composition_plan,
                "group_plan": group_plan,
                "document_model": model,
                "group_snapshot": group_snapshot,
            }
        if str(group_plan.get("status") or "").lower() == "blocked":
            return {
                "enabled": True,
                "blocker": "; ".join(_plan_blockers(group_plan)[:5]) or "Group-aware Document Planner blocked the refresh.",
                "composition_plan": composition_plan,
                "group_plan": group_plan,
                "document_model": model,
                "group_snapshot": group_snapshot,
            }
    return {
        "enabled": True,
        "composition_plan": composition_plan,
        "group_plan": group_plan,
        "document_model": model,
        "group_snapshot": group_snapshot,
        "selected_specialists": sorted(selected),
        "group_snapshot_hash": group_snapshot.get("snapshot_hash", ""),
        "blocker": "",
    }


def _date_group_overlap_reasons(
    date_edits: list[dict[str, Any]],
    group_plan: dict[str, Any],
    model: dict[str, Any],
) -> list[str]:
    targets = structure_index(model)
    reasons: list[str] = []
    for operation in group_plan.get("operations") or []:
        if not isinstance(operation, dict):
            continue
        target_id = str(operation.get("target_id") or "")
        target = targets.get(target_id) or {}
        serialized = json.dumps(target, ensure_ascii=False)
        for edit in date_edits:
            visible = str(edit.get("visible_text") or "")
            if visible and visible in serialized:
                reasons.append(
                    f"Date Planner and Group-aware Document Planner both own {target_id}; "
                    "Composition Planner must return one non-overlapping plan"
                )
                break
    return list(dict.fromkeys(reasons))


def execute_refresh_file_preview_operation(
    operation: dict[str, Any],
    env: dict[str, str],
    run_id: str,
    raw_task: str,
) -> dict[str, Any]:
    params = operation.get("params") if isinstance(operation.get("params"), dict) else {}
    errors = validate_refresh_file_preview_shape(operation)
    if errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "blocker": "; ".join(errors),
            "run_id": run_id,
        }

    source_path = resolve_local_path(source_file_path_from_operation(operation))
    suffix = source_path.suffix.lower()
    report_dir = REPORT_ROOT / f"{now_stamp()}_{safe_name(run_id)}"
    post_dir = POST_REFRESH_ROOT / safe_name(run_id)
    upload_dir = UPLOAD_READY_ROOT / safe_name(run_id)
    report_dir.mkdir(parents=True, exist_ok=True)
    post_dir.mkdir(parents=True, exist_ok=True)
    upload_dir.mkdir(parents=True, exist_ok=True)

    reference_pack = load_reference_pack(params, env)
    source_text = visible_text(source_path)
    source_audit = audit_visible_text(source_text)
    source_date_mentions = enrich_date_mentions_with_source_context(
        extract_date_mentions(source_text),
        params=params,
        env=env,
        reference_pack=reference_pack,
    )
    direct_replacements_requested = has_direct_replacement_instructions(params)
    target_name = output_filename(source_path)
    output_path = post_dir / target_name
    upload_ready_path = upload_dir / target_name
    edit_result: dict[str, Any] = {}
    preview_pdf: dict[str, Any] = {}
    date_plan: dict[str, Any] = {}
    date_edits: list[dict[str, Any]] = []
    schedule_table_rebuilds: list[dict[str, Any]] = []
    unchanged_dates: list[dict[str, Any]] = []
    date_plan_rejections: list[str] = []
    composition_date_rejections: list[str] = []
    schedule_table_rejections: list[str] = []
    unchanged_date_rejections: list[str] = []
    deleted_mapping_rejections: list[str] = []
    blocker = ""
    user_reason = ""
    generated = False

    file_change_planning = run_file_change_planners(
        operation,
        env,
        report_dir=report_dir,
        source_path=source_path,
        raw_task=raw_task,
        reference_pack=reference_pack,
    )
    composition_plan = (
        file_change_planning.get("composition_plan")
        if isinstance(file_change_planning.get("composition_plan"), dict)
        else {}
    )
    group_plan = (
        file_change_planning.get("group_plan")
        if isinstance(file_change_planning.get("group_plan"), dict)
        else {}
    )
    group_snapshot = (
        file_change_planning.get("group_snapshot")
        if isinstance(file_change_planning.get("group_snapshot"), dict)
        else {}
    )
    selected_specialists = set(file_change_planning.get("selected_specialists") or [])
    composition_enabled = bool(file_change_planning.get("enabled"))
    date_planner_selected = not composition_enabled or "date" in selected_specialists
    group_operations = [item for item in group_plan.get("operations") or [] if isinstance(item, dict)]
    group_changes_requested = bool(
        "group_document" in selected_specialists
        and str(group_plan.get("status") or "").lower() == "ready"
        and group_operations
    )
    group_edit_result: dict[str, Any] = {}

    if direct_replacements_requested:
        blocker = "Direct text replacement instructions are not executable in the file refresh lane."
        user_reason = (
            "I cannot safely refresh this file from direct replacement text alone. "
            "The Date Planner Agent must first map the visible dates to grounded current-semester dates."
        )
    elif str(file_change_planning.get("blocker") or "").strip():
        blocker = str(file_change_planning.get("blocker") or "").strip()
        user_reason = blocker
    elif suffix not in SUPPORTED_SUFFIXES:
        blocker = f"Unsupported file type for refresh preview: {suffix or '(none)'}"
    elif suffix in PDF_SUFFIXES and not bool(params.get("allow_pdf_only_copy")):
        blocker = "PDF-only refresh preview is blocked; provide an editable source or explicitly allow a PDF-only copy."
        user_reason = (
            "This attachment is a PDF only. I need an editable source file to update the dates safely, "
            "or you can explicitly allow a PDF-only copy."
        )
    elif source_date_mentions and date_planner_selected:
        deleted_mapping_rejections = deleted_content_mapping_reasons(source_date_mentions)
        task = date_planning_task(
            report_dir=report_dir,
            source_path=source_path,
            source_text=source_text,
            date_mentions=source_date_mentions,
            reference_pack=reference_pack,
            raw_task=raw_task,
            document_model=(
                file_change_planning.get("document_model")
                if composition_enabled and isinstance(file_change_planning.get("document_model"), dict)
                else None
            ),
            composition_plan=composition_plan if composition_enabled else None,
        )
        date_plan = run_date_planning_agent(env, task)
        date_edits, date_plan_rejections = valid_date_edits_from_plan(date_plan, source_text, reference_pack)
        schedule_table_rebuilds, schedule_table_rejections = valid_schedule_table_rebuilds_from_plan(
            date_plan,
            reference_pack,
        )
        if schedule_table_rebuilds and suffix != ".docx":
            schedule_table_rebuilds = []
            schedule_table_rejections.append(
                "teaching_week_schedule_rebuild currently requires an editable DOCX source"
            )
        unchanged_dates, unchanged_date_rejections = valid_unchanged_dates_from_plan(date_plan, source_text)
        if composition_enabled:
            composition_date_rejections = composition_date_plan_errors(
                date_edits,
                model=(
                    file_change_planning.get("document_model")
                    if isinstance(file_change_planning.get("document_model"), dict)
                    else {}
                ),
                composition_plan=composition_plan,
                schedule_table_rebuilds=schedule_table_rebuilds,
            )
            date_plan_rejections.extend(composition_date_rejections)
        uncovered_source_dates = stale_mentions_uncovered_by_unchanged(source_text, unchanged_dates)
        if composition_date_rejections:
            blocker = "The Date Planner Agent output did not satisfy Composition Planner structure ownership."
            user_reason = (
                "I stopped before editing because the date plan did not use the exact Office structures "
                f"assigned by the Composition Planner. Reason: {'; '.join(composition_date_rejections[:3])}."
            )
        elif deleted_mapping_rejections:
            blocker = "The source refers to course content that was deleted in the confirmed timeline."
            user_reason = (
                "I found a date tied to deleted timeline content with no target Teaching Week. "
                "I blocked the refresh instead of guessing another date. "
                f"Reason: {'; '.join(deleted_mapping_rejections[:3])}."
            )
        elif schedule_table_rejections:
            blocker = "The Date Planner Agent schedule table plan did not pass reference validation."
            user_reason = (
                "I found a teaching schedule table that needs row-level rebuilding, but its plan did not pass "
                f"the timeline, semester schedule, and calendar checks. Reason: {'; '.join(schedule_table_rejections[:3])}. "
                "I did not create a refreshed review copy."
            )
        elif not date_edits and not schedule_table_rebuilds and (
            date_plan_blockers(date_plan) or date_plan_rejections or uncovered_source_dates
        ):
            blocker = "The file still contains old visible dates and no trusted date plan was available to edit them."
            reasons = (
                date_plan_blockers(date_plan)
                or date_plan_rejections
                or unchanged_date_rejections
                or [f"uncovered old date: {item.get('visible_text')}" for item in uncovered_source_dates[:3]]
            )
            detail = f" Reason: {'; '.join(reasons[:3])}." if reasons else ""
            user_reason = (
                "I found old visible dates in the file, but the Date Planner Agent did not produce "
                f"a grounded edit plan yet.{detail} I did not create a refreshed review copy."
            )
        elif not date_edits and not schedule_table_rebuilds:
            if suffix in OFFICE_SUFFIXES:
                shutil.copy2(source_path, output_path)
                edit_result = {"status": "unchanged_date_plan_no_edits_needed", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
                generated = True
                preview_pdf = convert_to_pdf(output_path, upload_dir, env)
                if preview_pdf.get("ok"):
                    upload_ready_path = resolve_local_path(str(preview_pdf["pdf_path"]))
                else:
                    shutil.copy2(output_path, upload_ready_path)
            elif suffix in TEXT_SUFFIXES:
                shutil.copy2(source_path, output_path)
                shutil.copy2(output_path, upload_ready_path)
                edit_result = {"status": "unchanged_date_plan_no_edits_needed", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
                generated = True
            elif suffix in PDF_SUFFIXES:
                shutil.copy2(source_path, output_path)
                shutil.copy2(output_path, upload_ready_path)
                edit_result = {"status": "pdf_only_copy", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
                generated = True
    elif suffix in OFFICE_SUFFIXES:
        shutil.copy2(source_path, output_path)
        edit_result = {"status": "unchanged_no_date_edits_needed", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
        generated = True
        preview_pdf = convert_to_pdf(output_path, upload_dir, env)
        if preview_pdf.get("ok"):
            upload_ready_path = resolve_local_path(str(preview_pdf["pdf_path"]))
        else:
            shutil.copy2(output_path, upload_ready_path)
    elif suffix in TEXT_SUFFIXES:
        shutil.copy2(source_path, output_path)
        shutil.copy2(output_path, upload_ready_path)
        edit_result = {"status": "unchanged_no_date_edits_needed", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
        generated = True
    elif suffix in PDF_SUFFIXES:
        shutil.copy2(source_path, output_path)
        shutil.copy2(output_path, upload_ready_path)
        edit_result = {"status": "pdf_only_copy", "replacement_count": 0, "date_edits": {}, "applied_date_edit_ids": []}
        generated = True

    overlap_reasons = _date_group_overlap_reasons(
        date_edits,
        group_plan,
        file_change_planning.get("document_model")
        if isinstance(file_change_planning.get("document_model"), dict)
        else {},
    )
    if not blocker and overlap_reasons:
        blocker = overlap_reasons[0]
        user_reason = (
            "The specialist plans overlap on one Office structure. I stopped before editing so the "
            "Composition Planner can assign one safe owner."
        )

    if not blocker and (group_changes_requested or date_edits or schedule_table_rebuilds):
        if suffix in OFFICE_SUFFIXES:
            shutil.copy2(source_path, output_path)
            if composition_enabled:
                model = (
                    file_change_planning.get("document_model")
                    if isinstance(file_change_planning.get("document_model"), dict)
                    else {}
                )
                combined_plan = composed_document_change_plan(
                    model=model,
                    group_plan=group_plan,
                    date_edits=date_edits,
                    group_snapshot_hash=str(group_snapshot.get("snapshot_hash") or ""),
                )
                combined_result = apply_document_change_plan(
                    output_path,
                    combined_plan,
                    group_snapshot=group_snapshot,
                    run_id=run_id,
                )
                if not combined_result.get("ok"):
                    blocker = "; ".join(
                        str(value) for value in combined_result.get("blockers") or []
                    ) or "The validated composed document plan could not be applied."
                    user_reason = blocker
                else:
                    schedule_edit_result = apply_schedule_table_rebuilds_to_docx(
                        output_path, schedule_table_rebuilds
                    )
                    if str(schedule_edit_result.get("status") or "").startswith("blocked_"):
                        blocker = str(
                            schedule_edit_result.get("blocker")
                            or "Could not rebuild the DOCX teaching schedule table."
                        )
                        user_reason = blocker
                    else:
                        if group_operations:
                            group_edit_result = {
                                "ok": True,
                                "status": "modified_by_composed_document_change_plan",
                                "operation_count": len(group_operations),
                                "composed_result": combined_result,
                            }
                        edit_result = {
                            "status": "modified_by_composed_plans",
                            "replacement_count": (
                                len(group_operations)
                                + len(date_edits)
                                + int(schedule_edit_result.get("replacement_count") or 0)
                            ),
                            "group_change_result": group_edit_result,
                            "date_edits": {
                                str(item.get("visible_text") or ""): 1
                                for item in date_edits
                            },
                            "applied_date_edit_ids": [
                                str(item.get("id") or "") for item in date_edits
                            ],
                            "schedule_table_rebuilds": schedule_edit_result.get("schedule_table_rebuilds") or [],
                            "composed_change_result": combined_result,
                        }
                        generated = True
                        preview_pdf = convert_to_pdf(output_path, upload_dir, env)
                        if preview_pdf.get("ok"):
                            upload_ready_path = resolve_local_path(str(preview_pdf["pdf_path"]))
                        else:
                            shutil.copy2(output_path, upload_ready_path)
            else:
                if group_changes_requested:
                    group_edit_result = apply_document_change_plan(
                        output_path,
                        group_plan,
                        group_snapshot=group_snapshot,
                        run_id=run_id,
                    )
                    if not group_edit_result.get("ok"):
                        blocker = "; ".join(
                            str(value) for value in group_edit_result.get("blockers") or []
                        ) or "The validated Group-aware document plan could not be applied."
                        user_reason = blocker
                schedule_edit_result = (
                    apply_schedule_table_rebuilds_to_docx(output_path, schedule_table_rebuilds)
                    if not blocker
                    else {"status": "skipped_after_group_blocker", "replacement_count": 0}
                )
                if not blocker and str(schedule_edit_result.get("status") or "").startswith("blocked_"):
                    blocker = str(schedule_edit_result.get("blocker") or "Could not rebuild the DOCX teaching schedule table.")
                elif not blocker:
                    date_edit_result = apply_date_edits_to_ooxml(output_path, date_edits)
                    if date_edit_result.get("status") == "blocked_bad_ooxml_zip":
                        blocker = str(date_edit_result.get("blocker") or "Office file is not a valid OOXML package.")
                    else:
                        edit_result = {
                            "status": "modified_by_date_plan",
                            "replacement_count": int(schedule_edit_result.get("replacement_count") or 0)
                            + int(date_edit_result.get("replacement_count") or 0),
                            "date_edits": date_edit_result.get("date_edits") or {},
                            "applied_date_edit_ids": date_edit_result.get("applied_date_edit_ids") or [],
                            "schedule_table_rebuilds": schedule_edit_result.get("schedule_table_rebuilds") or [],
                        }
                        generated = True
                        preview_pdf = convert_to_pdf(output_path, upload_dir, env)
                        if preview_pdf.get("ok"):
                            upload_ready_path = resolve_local_path(str(preview_pdf["pdf_path"]))
                        else:
                            shutil.copy2(output_path, upload_ready_path)
        elif suffix in TEXT_SUFFIXES:
            updated, counts, applied_ids = apply_date_edits_to_text(read_text_lossy(source_path, 1_000_000), date_edits)
            output_path.write_text(updated, encoding="utf-8")
            shutil.copy2(output_path, upload_ready_path)
            edit_result = {
                "status": "modified_by_date_plan" if counts else "unchanged_no_matching_date_edits",
                "replacement_count": sum(counts.values()),
                "date_edits": counts,
                "applied_date_edit_ids": applied_ids,
            }
            generated = True

    output_text = visible_text(upload_ready_path if upload_ready_path.exists() else output_path) if generated else source_text
    raw_audit = (
        audit_visible_text(output_text)
        if generated
        else source_audit
    )
    audit = audit_with_date_plan_coverage(raw_audit, output_text, unchanged_dates)
    replacement_count = int(edit_result.get("replacement_count", 0) or 0)
    if not blocker and replacement_count == 0 and audit["status"] != "PASS":
        blocker = "The file still contains old visible dates and no trusted date plan was available to edit them."
        user_reason = (
            "I found old visible dates in the file, but the Date Planner Agent did not produce a grounded edit plan yet. "
            "I did not create a refreshed review copy."
        )
        generated = False
    if blocker:
        # A blocked plan must never publish a partially written Office file or
        # preview as a review/upload candidate.  The report still retains the
        # structured blocker and planner artifacts for diagnosis.
        generated = False
        preview_pdf = {}
        status = "BLOCKED"
        ok = False
        note = blocker
    elif audit["status"] != "PASS":
        status = "NEEDS_REVIEW"
        ok = True
        note = "Generated a refreshed file, but audit found visible dates or trace markers that need human review before upload."
    else:
        status = "READY_FOR_REVIEW"
        ok = True
        note = "Generated a refreshed file for staff review. No Moodle change was made."

    result = {
        "ok": ok,
        "status": status,
        "operation_kind": OPERATION_KIND,
        "execution_mode": EXECUTION_MODE,
        "run_id": run_id,
        "raw_task": raw_task,
        "source_file_path": repo_rel(source_path),
        "source_file_name": source_path.name,
        "refreshed_file_path": repo_rel(output_path) if generated and output_path.exists() else "",
        "upload_ready_path": repo_rel(upload_ready_path) if generated and upload_ready_path.exists() else "",
        "preview_pdf_path": preview_pdf.get("pdf_path", "") if generated and preview_pdf.get("ok") else "",
        "reference_pack": {
            "root": reference_pack["root"],
            "exists": reference_pack["exists"],
            "files": reference_pack["files"],
            "replacement_count": 0,
        },
        "composition_plan": composition_plan,
        "selected_specialists": sorted(selected_specialists),
        "group_plan": group_plan,
        "input_document_fingerprint": str(
            (file_change_planning.get("document_model") or {}).get("document_fingerprint")
            if isinstance(file_change_planning.get("document_model"), dict)
            else ""
        ),
        "group_snapshot_hash": str(file_change_planning.get("group_snapshot_hash") or ""),
        "group_edit_result": group_edit_result,
        "date_plan": date_plan,
        "date_plan_rejections": date_plan_rejections,
        "schedule_table_rebuilds": schedule_table_rebuilds,
        "schedule_table_rejections": schedule_table_rejections,
        "unchanged_dates": unchanged_dates,
        "unchanged_date_rejections": unchanged_date_rejections,
        "date_mentions": source_date_mentions,
        "date_edit_count": len(date_edits) + len(schedule_table_rebuilds),
        "group_edit_count": int(group_edit_result.get("operation_count") or 0),
        "edit_result": edit_result,
        "replacement_count": replacement_count,
        "raw_audit": raw_audit,
        "audit": audit,
        "blocker": blocker,
        "user_reason": user_reason,
        "note": note,
        "report_dir": repo_rel(report_dir),
        "attachments": [],
    }
    attachments = []
    for path_text in (result["preview_pdf_path"],):
        if path_text and path_text not in attachments:
            attachments.append(path_text)
    result["attachments"] = attachments
    result_path = report_dir / "result.json"
    report_path = report_dir / "review.md"
    result["result_path"] = repo_rel(result_path)
    result["review_report_path"] = repo_rel(report_path)
    result_path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    write_review_report(report_path, result)
    return result
