from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

try:
    from schedule_planning_agent import AGENT_NAME as SCHEDULE_PLANNER_AGENT_NAME
    from schedule_planning_agent import run_schedule_planning_agent
except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
    from course_agent_app.schedule_planning_agent import AGENT_NAME as SCHEDULE_PLANNER_AGENT_NAME
    from course_agent_app.schedule_planning_agent import run_schedule_planning_agent

try:
    from reference_pack_rebuild import reference_pack_path as rebuild_reference_pack_path
except ModuleNotFoundError:  # pragma: no cover
    from course_agent_app.reference_pack_rebuild import reference_pack_path as rebuild_reference_pack_path


PROJECT_ROOT = Path(__file__).resolve().parents[1]
ROOT = Path(__file__).resolve().parent
REFERENCE_ROOT = PROJECT_ROOT / "course_files" / "reference"
REPORT_ROOT = PROJECT_ROOT / "course_files" / "evidence" / "process" / "course_agent_app" / "semester_schedule_update"
OPERATION_KIND = "update_semester_schedule"
EXECUTION_MODE = "local_file_action"
SCHEDULE_FILENAME = "semester_schedule.json"
DEFAULT_TIMEZONE = "Asia/Hong_Kong"
FORBIDDEN_JOB_KEYS = {
    "command",
    "shell",
    "python",
    "python_code",
    "javascript",
    "javascript_code",
    "raw_code",
    "script_text",
}
WEEKDAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
WEEKDAY_ALIASES = {
    "mon": "Monday",
    "monday": "Monday",
    "tue": "Tuesday",
    "tues": "Tuesday",
    "tuesday": "Tuesday",
    "wed": "Wednesday",
    "wednesday": "Wednesday",
    "thu": "Thursday",
    "thur": "Thursday",
    "thurs": "Thursday",
    "thursday": "Thursday",
    "fri": "Friday",
    "friday": "Friday",
    "sat": "Saturday",
    "saturday": "Saturday",
    "sun": "Sunday",
    "sunday": "Sunday",
}
CHINESE_WEEKDAYS = {
    "一": "Monday",
    "二": "Tuesday",
    "三": "Wednesday",
    "四": "Thursday",
    "五": "Friday",
    "六": "Saturday",
    "日": "Sunday",
    "天": "Sunday",
}
NON_TEACHING_TIMELINE_MARKERS = (
    "holiday",
    "reading week",
    "reading/field trip",
    "revision period",
    "assessment period",
    "class suspension",
)


def now_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")


def safe_name(value: str, fallback: str = "schedule") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "-", str(value or "")).strip(" .-_")
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned[:120] or fallback


def resolve_local_path(path_text: str | Path) -> Path:
    path = Path(str(path_text or "")).expanduser()
    if path.is_absolute():
        return path
    for base in (PROJECT_ROOT, ROOT, Path.cwd()):
        candidate = base / path
        if candidate.exists():
            return candidate
    return PROJECT_ROOT / path


def repo_rel(path: str | Path) -> str:
    resolved = resolve_local_path(path)
    try:
        return str(resolved.resolve(strict=False).relative_to(PROJECT_ROOT))
    except ValueError:
        return str(resolved)


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def load_json(path: Path) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def is_update_semester_schedule_operation(operation: dict[str, Any]) -> bool:
    return str(operation.get("operation_kind") or "").strip() == OPERATION_KIND


def _params(operation: dict[str, Any]) -> dict[str, Any]:
    return operation.get("params") if isinstance(operation.get("params"), dict) else {}


def _path_error(label: str, value: str) -> str:
    if value.startswith(("http://", "https://")):
        return f"{label} must be a local path"
    if ".env" in Path(value).parts:
        return f"{label} must not point at a secret env file"
    return ""


def schedule_file_records(operation: dict[str, Any]) -> list[dict[str, str]]:
    params = _params(operation)
    raw = params.get("schedule_files") or params.get("schedule_file_paths") or operation.get("schedule_files") or []
    records: list[dict[str, str]] = []
    if isinstance(raw, (str, Path)):
        raw = [str(raw)]
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, str):
                path = item
                filename = Path(item).name
            elif isinstance(item, dict):
                path = str(item.get("local_path") or item.get("path") or item.get("file_path") or "").strip()
                filename = str(item.get("filename") or Path(path).name).strip()
            else:
                continue
            if path.strip():
                records.append({"path": path.strip(), "filename": filename})
    for key in ("schedule_file_path", "file_path", "path"):
        path = str(params.get(key) or operation.get(key) or "").strip()
        if path and all(record["path"] != path for record in records):
            records.append({"path": path, "filename": Path(path).name})
    return records


def schedule_text_from_operation(operation: dict[str, Any]) -> str:
    params = _params(operation)
    return str(params.get("schedule_text") or params.get("text") or operation.get("schedule_text") or "").strip()


def validate_update_semester_schedule_shape(operation: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    params = _params(operation)
    forbidden = sorted(FORBIDDEN_JOB_KEYS & set(params))
    if forbidden:
        errors.append("update_semester_schedule.params contains forbidden executable fields: " + ", ".join(forbidden))
    merge_mode = str(params.get("merge_mode") or operation.get("merge_mode") or "replace_all").strip()
    if merge_mode and merge_mode != "replace_all":
        errors.append("update_semester_schedule only supports merge_mode=replace_all")
    file_records = schedule_file_records(operation)
    schedule_text = schedule_text_from_operation(operation)
    if not file_records and not schedule_text:
        errors.append("update_semester_schedule requires params.schedule_files or params.schedule_text")
    for index, record in enumerate(file_records):
        path_text = record["path"]
        path_error = _path_error(f"update_semester_schedule.params.schedule_files[{index}]", path_text)
        if path_error:
            errors.append(path_error)
        elif not resolve_local_path(path_text).exists():
            errors.append(f"update_semester_schedule.params.schedule_files[{index}] does not exist: {path_text}")
    target = str(params.get("reference_pack_path") or operation.get("reference_pack_path") or "").strip()
    if target:
        path_error = _path_error("update_semester_schedule.params.reference_pack_path", target)
        if path_error:
            errors.append(path_error)
    return errors


def reference_pack_path(params: dict[str, Any], env: dict[str, str]) -> Path:
    return rebuild_reference_pack_path(params, env)


def weekday_index(name: str) -> int:
    return WEEKDAY_NAMES.index(canonical_weekday(name))


def canonical_weekday(value: str) -> str:
    text = str(value or "").strip()
    if text in WEEKDAY_NAMES:
        return text
    lowered = re.sub(r"[^a-z]", "", text.lower())
    if lowered in WEEKDAY_ALIASES:
        return WEEKDAY_ALIASES[lowered]
    if text in CHINESE_WEEKDAYS:
        return CHINESE_WEEKDAYS[text]
    raise ValueError(f"unsupported weekday {value!r}")


def parse_iso_date(value: str) -> date:
    return date.fromisoformat(str(value).strip())


def ordered_unique(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value not in seen:
            result.append(value)
            seen.add(value)
    return result


def parse_week_number(value: Any) -> int | None:
    if isinstance(value, int):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    match = re.search(r"\b(?:week|wk)?\s*(\d{1,2})\b", text, flags=re.IGNORECASE)
    if not match:
        match = re.search(r"第?\s*(\d{1,2})\s*周", text)
    return int(match.group(1)) if match else None


def parse_week_range(text: str) -> tuple[int, int] | None:
    patterns = (
        r"\bweek\s*(\d{1,2})\s*(?:-|–|—|to|through|到|至)\s*(\d{1,2})\b",
        r"\bwk\s*(\d{1,2})\s*(?:-|–|—|to|through|到|至)\s*(\d{1,2})\b",
        r"第?\s*(\d{1,2})\s*(?:-|–|—|到|至)\s*(\d{1,2})\s*周",
        r"\bweek\s*(\d{1,2})\b",
        r"第?\s*(\d{1,2})\s*周",
    )
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            start = int(match.group(1))
            end = int(match.group(2)) if len(match.groups()) >= 2 and match.group(2) else start
            return min(start, end), max(start, end)
    return None


def parse_date_list(value: Any) -> list[str]:
    if isinstance(value, list):
        raw_values = [str(item) for item in value]
    else:
        raw_values = re.split(r"[,;/\n]+|\s+and\s+|\s*&\s*", str(value or ""), flags=re.IGNORECASE)
    dates: list[str] = []
    for raw in raw_values:
        for match in re.finditer(r"\b20\d{2}-\d{1,2}-\d{1,2}\b", str(raw)):
            try:
                dates.append(parse_iso_date(match.group(0)).isoformat())
            except ValueError:
                continue
    return ordered_unique(dates)


def parse_weekdays(value: Any) -> list[str]:
    if isinstance(value, list):
        values = [str(item) for item in value]
    else:
        text = str(value or "")
        values = re.split(r"[,;/、，和与&\s]+", text)
        for match in re.finditer(r"(?:星期|礼拜|周)([一二三四五六日天]+)", text):
            values.extend(list(match.group(1)))
        for char in CHINESE_WEEKDAYS:
            if re.search(rf"(?:星期|礼拜|周){char}", text):
                values.append(char)
        compact = re.sub(r"\s+", "", text)
        chinese_pair = re.search(r"[周星期礼拜]?([一二三四五六日天]{2,})", compact)
        if chinese_pair:
            values.extend(list(chinese_pair.group(1)))
    result: list[str] = []
    for raw in values:
        token = str(raw or "").strip()
        if not token:
            continue
        token = token.replace("每周", "").replace("星期", "").replace("礼拜", "").replace("周", "")
        if len(token) > 1 and all(char in CHINESE_WEEKDAYS for char in token):
            for char in token:
                result.append(CHINESE_WEEKDAYS[char])
            continue
        lowered = re.sub(r"[^a-z]", "", token.lower())
        if lowered in WEEKDAY_ALIASES:
            result.append(WEEKDAY_ALIASES[lowered])
        elif token in CHINESE_WEEKDAYS:
            result.append(CHINESE_WEEKDAYS[token])
    ordered = ordered_unique(result)
    return sorted(ordered, key=lambda item: WEEKDAY_NAMES.index(item))


def normalized_header(value: Any) -> str:
    text = re.sub(r"[^A-Za-z0-9一二三四五六日天周星期日期上课教学备注说明_-]+", "_", str(value or "").strip().lower())
    text = text.strip("_")
    aliases = {
        "week": "week",
        "wk": "week",
        "周": "week",
        "周次": "week",
        "week_no": "week",
        "week_number": "week",
        "date": "teaching_dates",
        "dates": "teaching_dates",
        "teaching_date": "teaching_dates",
        "teaching_dates": "teaching_dates",
        "class_date": "teaching_dates",
        "上课日期": "teaching_dates",
        "日期": "teaching_dates",
        "weekday": "allowed_session_weekdays",
        "weekdays": "allowed_session_weekdays",
        "allowed_session_weekdays": "allowed_session_weekdays",
        "class_day": "allowed_session_weekdays",
        "class_days": "allowed_session_weekdays",
        "上课日": "allowed_session_weekdays",
        "星期": "allowed_session_weekdays",
        "notes": "notes",
        "note": "notes",
        "备注": "notes",
        "说明": "notes",
    }
    return aliases.get(text, text)


def row_value(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row and str(row[key] or "").strip():
            return row[key]
    return ""


def rows_from_csv(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig")
    reader = csv.DictReader(text.splitlines())
    rows = []
    for row in reader:
        rows.append({normalized_header(key): value for key, value in row.items()})
    return rows


def rows_from_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read schedule XLSX files") from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    raw_rows = list(sheet.iter_rows(values_only=True))
    if not raw_rows:
        return []
    headers = [normalized_header(item) for item in raw_rows[0]]
    rows: list[dict[str, Any]] = []
    for raw in raw_rows[1:]:
        if not any(item not in (None, "") for item in raw):
            continue
        rows.append({headers[index]: value for index, value in enumerate(raw) if index < len(headers)})
    return rows


def rows_and_text_from_docx(path: Path) -> tuple[list[dict[str, Any]], str]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to read schedule DOCX files") from exc
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    rows: list[dict[str, Any]] = []
    for table in document.tables:
        raw_rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not raw_rows:
            continue
        headers = [normalized_header(item) for item in raw_rows[0]]
        for raw in raw_rows[1:]:
            if not any(raw):
                continue
            rows.append({headers[index]: value for index, value in enumerate(raw) if index < len(headers)})
    return rows, text


def first_semester_facts(calendar_payload: dict[str, Any]) -> dict[str, Any]:
    semesters = calendar_payload.get("semesters")
    if isinstance(semesters, dict) and isinstance(semesters.get("first"), dict):
        return semesters["first"]
    if isinstance(semesters, list):
        for item in semesters:
            name = str(item.get("name") or "").strip().lower() if isinstance(item, dict) else ""
            if isinstance(item, dict) and name in {
                "first",
                "first semester",
                "semester 1",
                "sem 1",
            }:
                return item
    return {}


def target_start_year(semester: str, default: int = 2026) -> int:
    match = re.search(r"\b(20\d{2})\s*[-/]\s*20\d{2}\b", str(semester or ""))
    if match:
        return int(match.group(1))
    match = re.search(r"\b(20\d{2})\b", str(semester or ""))
    return int(match.group(1)) if match else default


def semester_bounds(reference_payloads: dict[str, Any], semester: str) -> tuple[date, date]:
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    facts = first_semester_facts(calendar)
    start_text = str(facts.get("start") or facts.get("first_teaching") or "").strip()
    end_text = str(facts.get("end") or facts.get("assessment_end") or facts.get("last_teaching") or "").strip()
    try:
        if start_text and end_text:
            return parse_iso_date(start_text), parse_iso_date(end_text)
    except ValueError:
        pass
    year = target_start_year(semester)
    lowered = str(semester or "").lower()
    if "sem2" in lowered or "second" in lowered:
        return date(year + 1, 1, 1), date(year + 1, 6, 30)
    if "sem1" in lowered or "first" in lowered:
        return date(year, 9, 1), date(year, 12, 31)
    return date(year, 9, 1), date(year + 1, 8, 31)


def calendar_week_for_date(reference_payloads: dict[str, Any], value: date) -> int | None:
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    facts = first_semester_facts(calendar)
    first_text = str(facts.get("first_teaching") or facts.get("start") or "").strip()
    try:
        first_teaching = parse_iso_date(first_text)
    except ValueError:
        return None
    if value < first_teaching:
        return None
    return ((value - first_teaching).days // 7) + 1


def timeline_anchor_for_week(reference_payloads: dict[str, Any], week: int) -> date | None:
    timeline = reference_payloads.get("timeline.json")
    rows = timeline.get("week_calendar") if isinstance(timeline, dict) else []
    teaching_week = 0
    if isinstance(rows, list):
        for row in rows:
            if not isinstance(row, dict):
                continue
            marker = " ".join(str(row.get(key) or "").lower() for key in ("lecture", "workshop", "tutorial", "notes"))
            if row.get("counts_as_teaching_week") is False or any(token in marker for token in NON_TEACHING_TIMELINE_MARKERS):
                continue
            teaching_week += 1
            if teaching_week != week:
                continue
            try:
                return parse_iso_date(str(row.get("date") or ""))
            except ValueError:
                return None
    return None


def calendar_anchor_exclusion_reason(calendar: dict[str, Any], anchor: date) -> str:
    periods = calendar.get("teaching_suspensions") if isinstance(calendar.get("teaching_suspensions"), list) else []
    for item in periods:
        if not isinstance(item, dict):
            continue
        should_exclude = (
            item.get("counts_as_teaching_week") is False
            or bool(item.get("classes_suspended"))
            or str(item.get("type") or "") == "reading_or_field_trip_week"
        )
        if not should_exclude:
            continue
        try:
            start = parse_iso_date(str(item.get("start") or ""))
            end = parse_iso_date(str(item.get("end") or ""))
        except ValueError:
            continue
        if start <= anchor <= end:
            return str(item.get("name") or item.get("type") or "calendar non-teaching period")
    facts = first_semester_facts(calendar)
    legacy_periods = (
        ("Reading/Field Trip Week", facts.get("reading_start"), facts.get("reading_end")),
        ("Revision Period", facts.get("revision_start"), facts.get("revision_end")),
        ("Assessment Period", facts.get("assessment_start"), facts.get("assessment_end")),
        ("Class Suspension Period", facts.get("class_suspension_start"), facts.get("class_suspension_end")),
    )
    for label, raw_start, raw_end in legacy_periods:
        if not raw_start or not raw_end:
            continue
        try:
            start = parse_iso_date(str(raw_start))
            end = parse_iso_date(str(raw_end))
        except ValueError:
            continue
        if start <= anchor <= end:
            return label
    days = calendar.get("non_teaching_days") if isinstance(calendar.get("non_teaching_days"), list) else []
    for item in days:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        if str(item.get("date") or "") == anchor.isoformat():
            return str(item.get("name") or item.get("type") or "calendar non-teaching day")
    return ""


def calendar_anchor_for_week(reference_payloads: dict[str, Any], week: int) -> date | None:
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    facts = first_semester_facts(calendar)
    first_teaching = str(facts.get("first_teaching") or facts.get("start") or "").strip()
    last_teaching = str(facts.get("last_teaching") or facts.get("revision_start") or facts.get("end") or "").strip()
    if not first_teaching or not last_teaching:
        return None
    try:
        anchor = parse_iso_date(first_teaching)
        end = parse_iso_date(last_teaching)
    except ValueError:
        return None
    teaching_week = 0
    while anchor <= end:
        if not calendar_anchor_exclusion_reason(calendar, anchor):
            teaching_week += 1
            if teaching_week == week:
                return anchor
        anchor += timedelta(days=7)
    return None


def recurring_teaching_week_dates(
    reference_payloads: dict[str, Any],
    weekdays: list[str],
    *,
    through_week: int,
) -> dict[int, list[str]]:
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    facts = first_semester_facts(calendar)
    first_text = str(facts.get("first_teaching") or facts.get("start") or "").strip()
    last_text = str(facts.get("last_teaching") or "").strip()
    if not last_text and facts.get("revision_start"):
        try:
            last_text = (parse_iso_date(str(facts["revision_start"])) - timedelta(days=1)).isoformat()
        except ValueError:
            last_text = ""
    if not last_text:
        last_text = str(facts.get("end") or "").strip()
    try:
        first_teaching = parse_iso_date(first_text)
        last_teaching = parse_iso_date(last_text)
        canonical_weekdays = sorted(ordered_unique([canonical_weekday(item) for item in weekdays]), key=weekday_index)
    except (ValueError, TypeError):
        return {}
    blocked_dates = calendar_blocked_teaching_dates(reference_payloads)
    blocked_calendar_weeks: set[date] = set()
    timeline = reference_payloads.get("timeline.json")
    timeline_rows = timeline.get("week_calendar") if isinstance(timeline, dict) else []
    if isinstance(timeline_rows, list):
        for row in timeline_rows:
            if not isinstance(row, dict):
                continue
            marker = " ".join(
                str(row.get(key) or "").lower()
                for key in ("lecture", "workshop", "tutorial", "notes")
            )
            if row.get("counts_as_teaching_week") is not False and not any(
                token in marker for token in NON_TEACHING_TIMELINE_MARKERS
            ):
                continue
            try:
                row_date = parse_iso_date(str(row.get("date") or ""))
            except ValueError:
                continue
            blocked_calendar_weeks.add(row_date - timedelta(days=row_date.weekday()))
    calendar_week_start = first_teaching - timedelta(days=first_teaching.weekday())
    teaching_week = 0
    result: dict[int, list[str]] = {}
    while calendar_week_start <= last_teaching and teaching_week < through_week:
        if calendar_week_start in blocked_calendar_weeks:
            calendar_week_start += timedelta(days=7)
            continue
        candidates = [
            calendar_week_start + timedelta(days=weekday_index(weekday))
            for weekday in canonical_weekdays
        ]
        eligible = [
            candidate.isoformat()
            for candidate in candidates
            if first_teaching <= candidate <= last_teaching and candidate not in blocked_dates
        ]
        if eligible:
            teaching_week += 1
            result[teaching_week] = eligible
        calendar_week_start += timedelta(days=7)
    return result


def dates_between(start: date, end: date) -> list[date]:
    if end < start:
        return []
    return [start + timedelta(days=offset) for offset in range((end - start).days + 1)]


def calendar_blocked_teaching_dates(reference_payloads: dict[str, Any]) -> dict[date, str]:
    calendar = reference_payloads.get("calendar.json") if isinstance(reference_payloads.get("calendar.json"), dict) else {}
    blocked: dict[date, str] = {}
    non_teaching_days = calendar.get("non_teaching_days") if isinstance(calendar.get("non_teaching_days"), list) else []
    for item in non_teaching_days:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        try:
            day = parse_iso_date(str(item.get("date") or ""))
        except ValueError:
            continue
        label = str(item.get("name") or item.get("type") or "calendar non-teaching day")
        blocked[day] = label
    periods = calendar.get("teaching_suspensions") if isinstance(calendar.get("teaching_suspensions"), list) else []
    for item in periods:
        if not isinstance(item, dict) or not bool(item.get("classes_suspended")):
            continue
        try:
            start = parse_iso_date(str(item.get("start") or ""))
            end = parse_iso_date(str(item.get("end") or ""))
        except ValueError:
            continue
        label = str(item.get("name") or item.get("type") or "calendar teaching suspension")
        for day in dates_between(start, end):
            blocked.setdefault(day, label)
    facts = first_semester_facts(calendar)
    legacy_periods = (
        ("Reading/Field Trip Week", facts.get("reading_start"), facts.get("reading_end")),
        ("Revision Period", facts.get("revision_start"), facts.get("revision_end")),
        ("Assessment Period", facts.get("assessment_start"), facts.get("assessment_end")),
        ("Class Suspension Period", facts.get("class_suspension_start"), facts.get("class_suspension_end")),
    )
    for label, raw_start, raw_end in legacy_periods:
        if not raw_start or not raw_end:
            continue
        try:
            start = parse_iso_date(str(raw_start))
            end = parse_iso_date(str(raw_end))
        except ValueError:
            continue
        for day in dates_between(start, end):
            blocked.setdefault(day, label)
    return blocked


def derive_dates_for_week(reference_payloads: dict[str, Any], week: int, weekdays: list[str]) -> list[str]:
    recurring_dates = recurring_teaching_week_dates(reference_payloads, weekdays, through_week=week)
    if week in recurring_dates:
        return recurring_dates[week]
    anchor = timeline_anchor_for_week(reference_payloads, week) or calendar_anchor_for_week(reference_payloads, week)
    if not anchor:
        return []
    week_start = anchor - timedelta(days=anchor.weekday())
    blocked_dates = calendar_blocked_teaching_dates(reference_payloads)
    return [
        candidate.isoformat()
        for weekday in weekdays
        for candidate in [week_start + timedelta(days=weekday_index(weekday))]
        if candidate not in blocked_dates
    ]


def reference_payloads(target_root: Path) -> dict[str, Any]:
    payloads: dict[str, Any] = {}
    for name in (
        "timeline.json",
        "calendar.json",
        "syllabus_rules.json",
        "source_manifest.json",
        SCHEDULE_FILENAME,
    ):
        path = target_root / name
        if path.is_file():
            payload = load_json(path)
            if isinstance(payload, dict):
                payloads[name] = payload
    return payloads


def source_manifest_after_schedule_update(
    manifest: dict[str, Any],
    *,
    timeline_path: Path,
    timeline_sha256: str,
    schedule_path: Path,
    schedule_sha256: str,
    run_id: str,
) -> dict[str, Any]:
    updated = dict(manifest)
    references = [
        dict(item)
        for item in updated.get("generated_references") or []
        if isinstance(item, dict)
    ]
    replacements = {
        "timeline.json": (timeline_path, timeline_sha256),
        SCHEDULE_FILENAME: (schedule_path, schedule_sha256),
    }
    seen: set[str] = set()
    for item in references:
        name = Path(str(item.get("path") or "")).name
        replacement = replacements.get(name)
        if not replacement:
            continue
        path, digest = replacement
        item["path"] = repo_rel(path)
        item["sha256"] = digest
        seen.add(name)
    for name, (path, digest) in replacements.items():
        if name not in seen:
            references.append({"path": repo_rel(path), "sha256": digest})
    updated["generated_references"] = references
    updated["last_schedule_update"] = {
        "run_id": run_id,
        "operation": "staff_schedule_reflow",
        "updated_references": ["timeline.json", SCHEDULE_FILENAME],
        "week_coordinate": "teaching_week",
    }
    return updated


def schedule_from_rows(
    rows: list[dict[str, Any]],
    *,
    course_code: str,
    semester: str,
    source_basis: list[str],
    reference_payloads_value: dict[str, Any],
) -> dict[str, Any] | None:
    weeks: list[dict[str, Any]] = []
    for row in rows:
        week = parse_week_number(row_value(row, "week", "week_number", "周次"))
        if not week:
            continue
        teaching_dates = parse_date_list(row_value(row, "teaching_dates", "date", "dates", "上课日期"))
        weekdays = parse_weekdays(row_value(row, "allowed_session_weekdays", "weekday", "weekdays", "上课日"))
        if not teaching_dates and weekdays:
            teaching_dates = derive_dates_for_week(reference_payloads_value, week, weekdays)
        if teaching_dates and not weekdays:
            weekdays = sorted(
                {WEEKDAY_NAMES[parse_iso_date(item).weekday()] for item in teaching_dates},
                key=lambda item: WEEKDAY_NAMES.index(item),
            )
        if not teaching_dates and not weekdays:
            continue
        weeks.append(
            {
                "week": week,
                "label": f"Week {week}",
                "teaching_dates": teaching_dates,
                "allowed_session_weekdays": weekdays,
                "notes": str(row_value(row, "notes", "note", "备注") or "").strip(),
            }
        )
    if not weeks:
        return None
    return base_schedule_payload(course_code, semester, source_basis, weeks)


def schedule_from_natural_text(
    text: str,
    *,
    course_code: str,
    semester: str,
    source_basis: list[str],
    reference_payloads_value: dict[str, Any],
) -> dict[str, Any] | None:
    segments = [segment.strip() for segment in re.split(r"[;；,，\n。]+", text) if segment.strip()]
    weeks: list[dict[str, Any]] = []
    for segment in segments:
        week_range = parse_week_range(segment)
        if not week_range:
            continue
        weekdays = parse_weekdays(segment)
        teaching_dates = parse_date_list(segment)
        if not weekdays and teaching_dates:
            weekdays = sorted(
                {WEEKDAY_NAMES[parse_iso_date(item).weekday()] for item in teaching_dates},
                key=lambda item: WEEKDAY_NAMES.index(item),
            )
        if not weekdays and not teaching_dates:
            continue
        start, end = week_range
        for week in range(start, end + 1):
            week_dates = teaching_dates if start == end and teaching_dates else []
            if not week_dates and weekdays:
                week_dates = derive_dates_for_week(reference_payloads_value, week, weekdays)
            weeks.append(
                {
                    "week": week,
                    "label": f"Week {week}",
                    "teaching_dates": week_dates,
                    "allowed_session_weekdays": weekdays,
                    "notes": segment,
                }
            )
    if not weeks:
        return None
    return base_schedule_payload(course_code, semester, source_basis, weeks)


def base_schedule_payload(
    course_code: str,
    semester: str,
    source_basis: list[str],
    weeks: list[dict[str, Any]],
) -> dict[str, Any]:
    return {
        "course_code": course_code,
        "semester": semester,
        "timezone": DEFAULT_TIMEZONE,
        "status": "source-confirmed",
        "merge_mode": "replace_all",
        "source_basis": source_basis,
        "week_numbering": {
            "mode": "dual",
            "mapping_field": "teaching_week",
            "display_field": "date_planner_selected",
            "available_display_fields": ["calendar_week", "teaching_week"],
            "legacy_week_field": "teaching_week",
            "rule": (
                "Mappings use teaching_week. File refresh Date Planner selects calendar_week or teaching_week "
                "for each visible Week column from the source table semantics."
            ),
        },
        "weeks": sorted(weeks, key=lambda item: int(item["week"])),
        "calendar_policy": {
            "teaching_dates_exclude": "calendar.json non_teaching_days with classes_suspended=true and teaching_suspensions with classes_suspended=true",
        },
        "mapping_rules": {
            "same_week_teaching_date": (
                "target date must be one of the target actual teaching_week teaching_dates"
            ),
            "same_weekday_deadline": (
                "target date must preserve the old actual teaching_week number, weekday, and time"
            ),
        },
    }


def normalize_schedule_payload(
    payload: dict[str, Any],
    *,
    course_code: str,
    semester: str,
    source_basis: list[str],
    reference_payloads_value: dict[str, Any],
) -> dict[str, Any] | None:
    if isinstance(payload.get("weeks"), list):
        rows = payload["weeks"]
    elif isinstance(payload.get("rows"), list):
        rows = payload["rows"]
    elif isinstance(payload.get("week_calendar"), list):
        rows = payload["week_calendar"]
    else:
        return None
    normalized_rows: list[dict[str, Any]] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        normalized_rows.append({normalized_header(key): value for key, value in row.items()})
    schedule = schedule_from_rows(
        normalized_rows,
        course_code=course_code,
        semester=str(payload.get("semester") or semester),
        source_basis=source_basis,
        reference_payloads_value=reference_payloads_value,
    )
    if schedule:
        schedule["course_code"] = str(payload.get("course_code") or course_code)
        schedule["timezone"] = str(payload.get("timezone") or DEFAULT_TIMEZONE)
        if isinstance(payload.get("mapping_rules"), dict):
            schedule["mapping_rules"].update(payload["mapping_rules"])
    return schedule


def parse_schedule_file(
    path: Path,
    *,
    course_code: str,
    semester: str,
    reference_payloads_value: dict[str, Any],
) -> tuple[dict[str, Any] | None, str]:
    source_basis = [repo_rel(path)]
    suffix = path.suffix.lower()
    if suffix == ".json":
        payload = load_json(path)
        if isinstance(payload, dict):
            schedule = normalize_schedule_payload(
                payload,
                course_code=course_code,
                semester=semester,
                source_basis=source_basis,
                reference_payloads_value=reference_payloads_value,
            )
            if schedule:
                return schedule, ""
        if isinstance(payload, list):
            schedule = schedule_from_rows(
                [item for item in payload if isinstance(item, dict)],
                course_code=course_code,
                semester=semester,
                source_basis=source_basis,
                reference_payloads_value=reference_payloads_value,
            )
            if schedule:
                return schedule, ""
        return None, path.read_text(encoding="utf-8", errors="replace")
    if suffix in {".csv", ".tsv"}:
        rows = rows_from_csv(path) if suffix == ".csv" else rows_from_tsv(path)
        return (
            schedule_from_rows(
                rows,
                course_code=course_code,
                semester=semester,
                source_basis=source_basis,
                reference_payloads_value=reference_payloads_value,
            ),
            path.read_text(encoding="utf-8", errors="replace"),
        )
    if suffix == ".xlsx":
        rows = rows_from_xlsx(path)
        return (
            schedule_from_rows(
                rows,
                course_code=course_code,
                semester=semester,
                source_basis=source_basis,
                reference_payloads_value=reference_payloads_value,
            ),
            "\n".join(" ".join(str(value) for value in row.values()) for row in rows),
        )
    if suffix == ".docx":
        rows, text = rows_and_text_from_docx(path)
        schedule = schedule_from_rows(
            rows,
            course_code=course_code,
            semester=semester,
            source_basis=source_basis,
            reference_payloads_value=reference_payloads_value,
        )
        return schedule, text
    text = path.read_text(encoding="utf-8", errors="replace")
    return (
        schedule_from_natural_text(
            text,
            course_code=course_code,
            semester=semester,
            source_basis=source_basis,
            reference_payloads_value=reference_payloads_value,
        ),
        text,
    )


def rows_from_tsv(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig")
    reader = csv.DictReader(text.splitlines(), delimiter="\t")
    return [{normalized_header(key): value for key, value in row.items()} for row in reader]


def validate_schedule_payload(payload: dict[str, Any], reference_payloads_value: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    if str(payload.get("merge_mode") or "replace_all") != "replace_all":
        errors.append("merge_mode must be replace_all")
    if not str(payload.get("course_code") or "").strip():
        errors.append("course_code is required")
    if not str(payload.get("semester") or "").strip():
        errors.append("semester is required")
    weeks = payload.get("weeks")
    if not isinstance(weeks, list) or not weeks:
        return errors + ["weeks must be a non-empty list"]
    lower_bound, upper_bound = semester_bounds(reference_payloads_value, str(payload.get("semester") or ""))
    blocked_dates = calendar_blocked_teaching_dates(reference_payloads_value)
    seen_weeks: set[int] = set()
    for index, week_payload in enumerate(weeks):
        if not isinstance(week_payload, dict):
            errors.append(f"weeks[{index}] must be an object")
            continue
        try:
            week = int(week_payload.get("week"))
        except (TypeError, ValueError):
            errors.append(f"weeks[{index}].week must be an integer")
            continue
        if week in seen_weeks:
            errors.append(f"weeks[{index}].week duplicates Week {week}")
        seen_weeks.add(week)
        dates = week_payload.get("teaching_dates")
        if not isinstance(dates, list) or not dates:
            errors.append(f"weeks[{index}].teaching_dates must be a non-empty list")
            continue
        weekdays = week_payload.get("allowed_session_weekdays")
        if not isinstance(weekdays, list) or not weekdays:
            errors.append(f"weeks[{index}].allowed_session_weekdays must be a non-empty list")
            weekdays = []
        canonical_weekdays: list[str] = []
        for raw_weekday in weekdays:
            try:
                canonical_weekdays.append(canonical_weekday(str(raw_weekday)))
            except ValueError as exc:
                errors.append(f"weeks[{index}].allowed_session_weekdays has {exc}")
        for raw_date in dates:
            try:
                parsed = parse_iso_date(str(raw_date))
            except ValueError:
                errors.append(f"weeks[{index}].teaching_dates contains non-ISO date {raw_date!r}")
                continue
            if parsed < lower_bound or parsed > upper_bound:
                errors.append(
                    f"weeks[{index}].teaching_dates date {parsed.isoformat()} is outside target semester range "
                    f"{lower_bound.isoformat()}..{upper_bound.isoformat()}"
                )
            if parsed in blocked_dates:
                errors.append(
                    f"weeks[{index}].teaching_dates date {parsed.isoformat()} is blocked by calendar.json: "
                    f"{blocked_dates[parsed]}"
                )
            actual_weekday = WEEKDAY_NAMES[parsed.weekday()]
            if canonical_weekdays and actual_weekday not in canonical_weekdays:
                errors.append(
                    f"weeks[{index}].teaching_dates date {parsed.isoformat()} is {actual_weekday}, "
                    "but allowed_session_weekdays says " + ", ".join(canonical_weekdays)
                )
        parsed_dates = []
        for raw_date in dates:
            try:
                parsed_dates.append(parse_iso_date(str(raw_date)))
            except ValueError:
                continue
        calendar_weeks = {
            value
            for parsed in parsed_dates
            for value in [calendar_week_for_date(reference_payloads_value, parsed)]
            if value is not None
        }
        if len(calendar_weeks) > 1:
            errors.append(f"weeks[{index}].teaching_dates span multiple calendar weeks")
        derived_calendar_week = next(iter(calendar_weeks), None)
        supplied_calendar_week = week_payload.get("calendar_week")
        if supplied_calendar_week not in (None, "") and derived_calendar_week is not None:
            try:
                if int(supplied_calendar_week) != derived_calendar_week:
                    errors.append(
                        f"weeks[{index}].calendar_week says {supplied_calendar_week}, "
                        f"but its teaching_dates fall in Calendar Week {derived_calendar_week}"
                    )
            except (TypeError, ValueError):
                errors.append(f"weeks[{index}].calendar_week must be an integer")
        week_payload["teaching_week"] = week
        if derived_calendar_week is not None:
            week_payload["calendar_week"] = derived_calendar_week
        week_payload["allowed_session_weekdays"] = sorted(ordered_unique(canonical_weekdays), key=WEEKDAY_NAMES.index)
        week_payload["teaching_dates"] = sorted(ordered_unique([str(item) for item in dates]))
        week_payload.setdefault("label", f"Week {week}")
        week_payload.setdefault("notes", "")
    payload["weeks"] = sorted(weeks, key=lambda item: int(item.get("week") or 0) if isinstance(item, dict) else 0)
    return errors


def merge_schedule_payloads(payloads: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not payloads:
        return None
    if len(payloads) == 1:
        return payloads[0]
    merged = dict(payloads[-1])
    weeks_by_number: dict[int, dict[str, Any]] = {}
    source_basis: list[str] = []
    for payload in payloads:
        for item in payload.get("source_basis") or []:
            if item not in source_basis:
                source_basis.append(str(item))
        for week in payload.get("weeks") or []:
            if isinstance(week, dict):
                weeks_by_number[int(week["week"])] = dict(week)
    merged["source_basis"] = source_basis
    merged["weeks"] = [weeks_by_number[key] for key in sorted(weeks_by_number)]
    return merged


def schedule_planning_task(
    *,
    report_dir: Path,
    raw_task: str,
    schedule_text: str,
    course_code: str,
    semester: str,
    reference_payloads_value: dict[str, Any],
) -> dict[str, Any]:
    return {
        "agent_task_path": str(report_dir / "schedule_planning_agent_task.md"),
        "agent_output_path": str(report_dir / "schedule_planning_agent_output.json"),
        "raw_task": raw_task,
        "schedule_text": schedule_text,
        "course_code": course_code,
        "semester": semester,
        "reference_payloads": reference_payloads_value,
    }


def schedule_from_agent_output(
    output: dict[str, Any],
    *,
    course_code: str,
    semester: str,
    source_basis: list[str],
    reference_payloads_value: dict[str, Any],
) -> tuple[dict[str, Any] | None, list[str]]:
    if str(output.get("agent") or "").strip() != SCHEDULE_PLANNER_AGENT_NAME:
        return None, ["schedule plan was not produced by the Schedule Planner Agent"]
    status = str(output.get("status") or "").strip().lower()
    arrangement_mode = str(output.get("arrangement_mode") or "").strip().lower()
    if arrangement_mode == "recurring_weekdays":
        recurring_weekdays = parse_weekdays(output.get("recurring_weekdays") or [])
        raw_range = output.get("requested_week_range")
        if isinstance(raw_range, dict):
            try:
                start_week = int(raw_range.get("start") or 1)
                end_week = int(raw_range.get("end") or start_week)
            except (TypeError, ValueError):
                start_week, end_week = 1, 0
        else:
            start_week, end_week = 1, 0
        if recurring_weekdays and end_week >= start_week >= 1:
            week_dates = recurring_teaching_week_dates(
                reference_payloads_value,
                recurring_weekdays,
                through_week=end_week,
            )
            missing_weeks = [week for week in range(start_week, end_week + 1) if not week_dates.get(week)]
            if not missing_weeks:
                weeks = [
                    {
                        "week": week,
                        "label": f"Week {week}",
                        "teaching_dates": week_dates[week],
                        "allowed_session_weekdays": recurring_weekdays,
                        "notes": "Derived from the staff recurring-weekday arrangement; non-teaching dates do not consume teaching-week numbers.",
                    }
                    for week in range(start_week, end_week + 1)
                ]
                return base_schedule_payload(course_code, semester, source_basis, weeks), []
            available_count = len([week for week in week_dates if start_week <= week <= end_week])
            calendar = (
                reference_payloads_value.get("calendar.json")
                if isinstance(reference_payloads_value.get("calendar.json"), dict)
                else {}
            )
            facts = first_semester_facts(calendar)
            last_teaching = str(facts.get("last_teaching") or facts.get("revision_start") or facts.get("end") or "").strip()
            weekday_text = ", ".join(recurring_weekdays)
            return None, [
                f"The recurring {weekday_text} arrangement needs teaching Weeks {start_week}-{end_week}, "
                f"but only {available_count} eligible teaching week(s) fit before the semester teaching period ends"
                + (f" ({last_teaching})" if last_teaching else "")
                + f"; Week {missing_weeks[0]} has no valid class date after holidays and class-suspension periods are excluded."
            ]
    if status != "ready":
        blockers = output.get("blockers") if isinstance(output.get("blockers"), list) else []
        formatted = [schedule_agent_blocker_text(item) for item in blockers]
        return None, [item for item in formatted if item] or ["schedule planner did not return ready"]
    schedule = normalize_schedule_payload(
        output,
        course_code=course_code,
        semester=semester,
        source_basis=source_basis,
        reference_payloads_value=reference_payloads_value,
    )
    return schedule, [] if schedule else ["schedule planner output did not contain usable weeks"]


def schedule_agent_blocker_text(value: Any) -> str:
    if not isinstance(value, dict):
        return str(value or "").strip()
    reason = str(value.get("reason") or value.get("message") or "").strip()
    week = value.get("week")
    weekday = str(value.get("requested_weekday") or "").strip()
    blocked_date = str(value.get("blocked_date") or "").strip()
    subject_parts: list[str] = []
    if week not in (None, ""):
        subject_parts.append(f"Week {week}")
    if weekday:
        subject_parts.append(weekday)
    if blocked_date:
        subject_parts.append(blocked_date)
    subject = " / ".join(subject_parts)
    if subject and reason:
        return f"{subject}: {reason}"
    return reason or subject


def recurring_weekdays_for_overflow_draft(
    payload: dict[str, Any] | None,
    agent_output: dict[str, Any],
    *,
    expected_week_count: int,
) -> list[str]:
    if str(agent_output.get("arrangement_mode") or "").strip().lower() == "recurring_weekdays":
        raw_range = agent_output.get("requested_week_range")
        try:
            start = int(raw_range.get("start")) if isinstance(raw_range, dict) else 0
            end = int(raw_range.get("end")) if isinstance(raw_range, dict) else 0
        except (TypeError, ValueError):
            start, end = 0, 0
        weekdays = parse_weekdays(agent_output.get("recurring_weekdays") or [])
        if start == 1 and end == expected_week_count and weekdays:
            return weekdays
    if not isinstance(payload, dict):
        return []
    weeks = [item for item in payload.get("weeks") or [] if isinstance(item, dict)]
    week_numbers = [
        int(item.get("teaching_week") or item.get("week") or 0)
        for item in weeks
    ]
    if week_numbers != list(range(1, expected_week_count + 1)):
        return []
    weekday_sets = {
        tuple(parse_weekdays(item.get("allowed_session_weekdays") or []))
        for item in weeks
    }
    if len(weekday_sets) != 1:
        return []
    return list(next(iter(weekday_sets), ()))


def maybe_create_overflow_timeline_draft(
    *,
    payload: dict[str, Any] | None,
    agent_output: dict[str, Any],
    reference_payloads_value: dict[str, Any],
    target_root: Path,
    staged_sources: list[dict[str, Any]],
    env: dict[str, str],
    run_id: str,
    raw_task: str,
    course_code: str,
    semester: str,
) -> dict[str, Any] | None:
    try:
        from reference_pack_rebuild import teaching_content_rows
        from timeline_draft_review import create_timeline_draft
    except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
        from course_agent_app.reference_pack_rebuild import teaching_content_rows
        from course_agent_app.timeline_draft_review import create_timeline_draft

    timeline_payload = reference_payloads_value.get("timeline.json")
    calendar_payload = reference_payloads_value.get("calendar.json")
    if not isinstance(timeline_payload, dict) or not isinstance(calendar_payload, dict):
        return None
    expected_week_count = len(teaching_content_rows(timeline_payload))
    weekdays = recurring_weekdays_for_overflow_draft(
        payload,
        agent_output,
        expected_week_count=expected_week_count,
    )
    if not weekdays:
        return None
    result = create_timeline_draft(
        origin=OPERATION_KIND,
        run_id=run_id,
        raw_task=raw_task,
        course_code=course_code,
        semester=semester,
        class_weekdays=weekdays,
        reference_pack_path=target_root,
        source_timeline=timeline_payload,
        calendar_payload=calendar_payload,
        candidate_payloads={
            "source_manifest.json": reference_payloads_value.get("source_manifest.json") or {},
            "syllabus_rules.json": reference_payloads_value.get("syllabus_rules.json") or {},
        },
        staged_sources=staged_sources,
        env=env,
    )
    return result if result.get("ok") else None


def execute_update_semester_schedule_operation(
    operation: dict[str, Any],
    env: dict[str, str],
    run_id: str,
    raw_task: str,
) -> dict[str, Any]:
    errors = validate_update_semester_schedule_shape(operation)
    if errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "blocker": "; ".join(errors),
            "attachments": [],
        }

    params = _params(operation)
    course_code = str(params.get("course_code") or env.get("COURSE_AGENT_DEFAULT_COURSE_CODE") or "DASE7137")
    semester = str(params.get("semester") or env.get("COURSE_AGENT_DEFAULT_SEMESTER") or "2026-2027 first semester")
    target_root = reference_pack_path(params, env)
    target_path = target_root / SCHEDULE_FILENAME
    report_dir = REPORT_ROOT / f"{now_stamp()}_{safe_name(run_id)}"
    report_dir.mkdir(parents=True, exist_ok=True)
    target_root.mkdir(parents=True, exist_ok=True)

    reference_payloads_value = reference_payloads(target_root)
    parsed_payloads: list[dict[str, Any]] = []
    collected_text: list[str] = []
    staged_sources: list[dict[str, Any]] = []
    blockers: list[str] = []
    agent_output: dict[str, Any] = {}

    for record in schedule_file_records(operation):
        source_path = resolve_local_path(record["path"])
        staged_sources.append(
            {
                "path": repo_rel(source_path),
                "filename": record.get("filename") or source_path.name,
                "sha256": sha256(source_path),
            }
        )
        try:
            schedule, extracted_text = parse_schedule_file(
                source_path,
                course_code=course_code,
                semester=semester,
                reference_payloads_value=reference_payloads_value,
            )
        except Exception as exc:
            schedule = None
            extracted_text = ""
            blockers.append(f"could not parse {source_path.name}: {exc}")
        if schedule:
            parsed_payloads.append(schedule)
        if extracted_text:
            collected_text.append(f"\n\n# {source_path.name}\n{extracted_text}")

    direct_text = schedule_text_from_operation(operation)
    if direct_text:
        schedule = schedule_from_natural_text(
            direct_text,
            course_code=course_code,
            semester=semester,
            source_basis=["current WhatsApp prompt text"],
            reference_payloads_value=reference_payloads_value,
        )
        if schedule:
            parsed_payloads.append(schedule)
        collected_text.append(f"\n\n# current WhatsApp prompt text\n{direct_text}")

    agent_rejections: list[str] = []
    if not parsed_payloads and collected_text:
        task = schedule_planning_task(
            report_dir=report_dir,
            raw_task=raw_task,
            schedule_text="\n".join(collected_text),
            course_code=course_code,
            semester=semester,
            reference_payloads_value=reference_payloads_value,
        )
        agent_output = run_schedule_planning_agent(env, task)
        schedule, agent_rejections = schedule_from_agent_output(
            agent_output,
            course_code=course_code,
            semester=semester,
            source_basis=["Schedule Planner Agent"],
            reference_payloads_value=reference_payloads_value,
        )
        if schedule:
            parsed_payloads.append(schedule)

    payload = merge_schedule_payloads(parsed_payloads)
    if not payload:
        draft_result = maybe_create_overflow_timeline_draft(
            payload=None,
            agent_output=agent_output,
            reference_payloads_value=reference_payloads_value,
            target_root=target_root,
            staged_sources=staged_sources,
            env=env,
            run_id=run_id,
            raw_task=raw_task,
            course_code=course_code,
            semester=semester,
        )
        if draft_result:
            draft_result["raw_task"] = raw_task
            draft_result["staged_sources"] = staged_sources
            result_path = report_dir / "semester_schedule_update_result.json"
            write_json(result_path, draft_result)
            draft_result["result_path"] = repo_rel(result_path)
            return draft_result
        reason = "; ".join(blockers + agent_rejections) or "No parseable semester schedule rows were found."
        result = {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "raw_task": raw_task,
            "reference_pack_path": repo_rel(target_root),
            "blocker": reason,
            "staged_sources": staged_sources,
            "attachments": [],
        }
        result_path = report_dir / "semester_schedule_update_result.json"
        write_json(result_path, result)
        result["result_path"] = repo_rel(result_path)
        return result

    payload["provenance"] = {
        "origin": "staff_update",
        "user_supplied": True,
        "input_boundary": "current WhatsApp prompt attachments or text",
        "merge_mode": "replace_all",
    }
    payload["mapping_rules"].update(
        {
            "lab_visit": "must_map_to_teaching_dates_in_same_teaching_week",
            "lecture": "must_map_to_teaching_dates_in_same_teaching_week",
            "tutorial": "must_map_to_teaching_dates_in_same_teaching_week",
            "session_date": "must_map_to_teaching_dates_in_same_teaching_week",
            "class_visit": "must_map_to_teaching_dates_in_same_teaching_week",
            "assignment_deadline": "preserve_old_teaching_week_weekday_and_time",
            "submission_deadline": "preserve_old_teaching_week_weekday_and_time",
        }
    )
    validation_errors = validate_schedule_payload(payload, reference_payloads_value)
    if validation_errors:
        draft_result = maybe_create_overflow_timeline_draft(
            payload=payload,
            agent_output=agent_output,
            reference_payloads_value=reference_payloads_value,
            target_root=target_root,
            staged_sources=staged_sources,
            env=env,
            run_id=run_id,
            raw_task=raw_task,
            course_code=course_code,
            semester=semester,
        )
        if draft_result:
            draft_result["raw_task"] = raw_task
            draft_result["staged_sources"] = staged_sources
            result_path = report_dir / "semester_schedule_update_result.json"
            write_json(result_path, draft_result)
            draft_result["result_path"] = repo_rel(result_path)
            return draft_result
        result = {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "raw_task": raw_task,
            "reference_pack_path": repo_rel(target_root),
            "blocker": "; ".join(validation_errors),
            "staged_sources": staged_sources,
            "attachments": [],
        }
        result_path = report_dir / "semester_schedule_update_result.json"
        write_json(result_path, result)
        result["result_path"] = repo_rel(result_path)
        return result

    timeline_path = target_root / "timeline.json"
    manifest_path = target_root / "source_manifest.json"
    timeline_payload = reference_payloads_value.get("timeline.json")
    calendar_payload = reference_payloads_value.get("calendar.json")
    manifest_payload = reference_payloads_value.get("source_manifest.json")
    missing_references = [
        name
        for name, value in (
            ("timeline.json", timeline_payload),
            ("calendar.json", calendar_payload),
            ("source_manifest.json", manifest_payload),
        )
        if not isinstance(value, dict)
    ]
    if missing_references:
        result = {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "raw_task": raw_task,
            "reference_pack_path": repo_rel(target_root),
            "blocker": (
                "The schedule update requires a complete reference pack so it can reflow timeline.json; missing "
                + ", ".join(missing_references)
            ),
            "staged_sources": staged_sources,
            "attachments": [],
        }
        result_path = report_dir / "semester_schedule_update_result.json"
        write_json(result_path, result)
        result["result_path"] = repo_rel(result_path)
        return result

    try:
        from reference_pack_rebuild import reflow_timeline_to_schedule, teaching_content_rows
    except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
        from course_agent_app.reference_pack_rebuild import (
            reflow_timeline_to_schedule,
            teaching_content_rows,
        )

    assert isinstance(timeline_payload, dict)
    assert isinstance(calendar_payload, dict)
    assert isinstance(manifest_payload, dict)
    expected_week_count = len(teaching_content_rows(timeline_payload))
    schedule_week_numbers = [
        int(item.get("teaching_week") or item.get("week"))
        for item in payload.get("weeks") or []
        if isinstance(item, dict)
    ]
    if schedule_week_numbers != list(range(1, expected_week_count + 1)):
        result = {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "raw_task": raw_task,
            "reference_pack_path": repo_rel(target_root),
            "blocker": (
                "replace_all must provide every actual teaching week needed by timeline.json; "
                f"expected Teaching Weeks 1-{expected_week_count}, got {schedule_week_numbers}"
            ),
            "staged_sources": staged_sources,
            "attachments": [],
        }
        result_path = report_dir / "semester_schedule_update_result.json"
        write_json(result_path, result)
        result["result_path"] = repo_rel(result_path)
        return result
    try:
        updated_timeline = reflow_timeline_to_schedule(
            timeline_payload,
            payload,
            calendar_payload,
        )
    except (KeyError, TypeError, ValueError) as exc:
        result = {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "raw_task": raw_task,
            "reference_pack_path": repo_rel(target_root),
            "blocker": f"Could not reflow timeline.json from the complete teaching schedule: {exc}",
            "staged_sources": staged_sources,
            "attachments": [],
        }
        result_path = report_dir / "semester_schedule_update_result.json"
        write_json(result_path, result)
        result["result_path"] = repo_rel(result_path)
        return result

    backup_paths: dict[str, str] = {}
    for path in (target_path, timeline_path, manifest_path):
        if not path.exists():
            continue
        backup = report_dir / f"{path.name}.previous"
        shutil.copy2(path, backup)
        backup_paths[path.name] = repo_rel(backup)

    temp_schedule = target_root / f".{SCHEDULE_FILENAME}.{safe_name(run_id)}.tmp"
    temp_timeline = target_root / f".timeline.json.{safe_name(run_id)}.tmp"
    temp_manifest = target_root / f".source_manifest.json.{safe_name(run_id)}.tmp"
    try:
        write_json(temp_schedule, payload)
        write_json(temp_timeline, updated_timeline)
        updated_manifest = source_manifest_after_schedule_update(
            manifest_payload,
            timeline_path=timeline_path,
            timeline_sha256=sha256(temp_timeline),
            schedule_path=target_path,
            schedule_sha256=sha256(temp_schedule),
            run_id=run_id,
        )
        write_json(temp_manifest, updated_manifest)
        os.replace(temp_timeline, timeline_path)
        os.replace(temp_schedule, target_path)
        os.replace(temp_manifest, manifest_path)
    except Exception as exc:
        for path in (temp_schedule, temp_timeline, temp_manifest):
            path.unlink(missing_ok=True)
        for target in (target_path, timeline_path, manifest_path):
            backup_text = backup_paths.get(target.name)
            if backup_text:
                shutil.copy2(resolve_local_path(backup_text), target)
        result = {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": OPERATION_KIND,
            "run_id": run_id,
            "raw_task": raw_task,
            "reference_pack_path": repo_rel(target_root),
            "blocker": f"Could not atomically replace the schedule references: {exc}",
            "staged_sources": staged_sources,
            "attachments": [],
        }
        result_path = report_dir / "semester_schedule_update_result.json"
        write_json(result_path, result)
        result["result_path"] = repo_rel(result_path)
        return result

    result = {
        "ok": True,
        "status": "READY",
        "operation_kind": OPERATION_KIND,
        "run_id": run_id,
        "raw_task": raw_task,
        "course_code": course_code,
        "semester": semester,
        "reference_pack_path": repo_rel(target_root),
        "schedule_path": repo_rel(target_path),
        "timeline_path": repo_rel(timeline_path),
        "source_manifest_path": repo_rel(manifest_path),
        "backup_path": backup_paths.get(SCHEDULE_FILENAME, ""),
        "backup_paths": backup_paths,
        "week_count": len(payload.get("weeks") or []),
        "schedule_update_scope": "date_only_reflow_preserving_timeline_content_mappings",
        "preserved_content_mapping_count": len(updated_timeline.get("content_week_mappings") or []),
        "preserved_content_item_count": sum(
            len(item.get("content_item_ids") or [])
            for item in updated_timeline.get("week_calendar") or []
            if isinstance(item, dict)
        ),
        "generated_files": [
            repo_rel(target_path),
            repo_rel(timeline_path),
            repo_rel(manifest_path),
        ],
        "staged_sources": staged_sources,
        "attachments": [],
    }
    result_path = report_dir / "semester_schedule_update_result.json"
    write_json(result_path, result)
    result["result_path"] = repo_rel(result_path)
    return result
