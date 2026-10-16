import copy
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from chrome_file_library import FileLibrary, PAGE_SIZE
from run_shared_chrome_agent import Runtime, load_config
from chrome_planner import schema


class LibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.sources = self.root / 'files'
        self.sources.mkdir()
        self.library = FileLibrary(self.sources, self.root / 'v2/index.json')

    def tearDown(self):
        self.temp.cleanup()

    def test_refresh_preserves_versions_and_originals(self):
        (self.sources / 'syllabus 25-26.txt').write_text('old dates')
        first = self.library.snapshot()
        (self.sources / 'syllabus 26-27.txt').write_text('current dates')
        second = self.library.snapshot()
        self.assertEqual(len(first['files']), 1)
        self.assertEqual(len(second['files']), 2)
        self.assertNotEqual(first['snapshot_id'], second['snapshot_id'])
        self.assertTrue(self.library.index.with_suffix('.md').is_file())
        self.assertEqual((self.sources / 'syllabus 25-26.txt').read_text(), 'old dates')

    def test_exact_sources_only_and_hash_drift(self):
        source = self.sources / 'course.txt'
        source.write_text('2026 September')
        snapshot = self.library.snapshot()
        fid = snapshot['files'][0]['id']
        read = self.library.inspect(snapshot, fid)
        self.assertIn('2026', read['text'])
        with self.assertRaises(ValueError): self.library.inspect(snapshot, '../outside.txt')
        forged = copy.deepcopy(snapshot)
        forged['files'][0]['relative_path'] = '../outside.txt'
        with self.assertRaises(ValueError): self.library.inspect(forged, fid)
        source.write_text('2026 December')
        with self.assertRaisesRegex(ValueError, 'changed'): self.library.validate_reads(snapshot, [read])

    def test_pagination_and_unknown_format(self):
        (self.sources / 'long.txt').write_text('a' * (PAGE_SIZE + 50))
        (self.sources / 'image.bin').write_bytes(b'not text')
        snapshot = self.library.snapshot()
        by_path = {e['relative_path']:e['id'] for e in snapshot['files']}
        first = self.library.inspect(snapshot, by_path['long.txt'])
        second = self.library.inspect(snapshot, by_path['long.txt'], first['next_offset'])
        self.assertEqual(first['next_offset'], PAGE_SIZE)
        self.assertIsNone(second['next_offset'])
        self.assertEqual(second['text'].strip(), 'a' * 50)
        self.assertEqual(self.library.inspect(snapshot, by_path['image.bin'])['status'], 'reader_unavailable')

    def test_office_content_read_without_modification(self):
        from docx import Document
        from openpyxl import Workbook
        doc = Document(); doc.add_paragraph('Semester 1')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Presentation 7 October'
        doc.save(self.sources / 'source.docx')
        book = Workbook(); book.active.append(['Class A', 'Workshop'])
        book.save(self.sources / 'schedule.xlsx')
        snapshot = self.library.snapshot()
        texts = [self.library.inspect(snapshot, e['id']) for e in snapshot['files']]
        self.assertTrue(any('Presentation 7 October' in r['text'] for r in texts))
        self.assertTrue(any('Workshop' in r['text'] for r in texts))
        self.assertEqual(snapshot['snapshot_id'], self.library.snapshot()['snapshot_id'])

    def test_linked_source_rejected(self):
        path = self.sources / 'linked.txt'; path.write_text('outside')
        with patch.object(Path, 'lstat', return_value=SimpleNamespace(st_mode=0, st_file_attributes=0x400)):
            with self.assertRaisesRegex(ValueError, 'Linked'): self.library.safe_path('linked.txt')

    def test_runtime_supplies_index_then_selected_content(self):
        (self.sources / 'timeline.txt').write_text('Class B presentation on 10 October')
        browser = Mock(); browser.overviews = {}; browser.forms = {}
        runtime = Runtime(None, browser, ['chat'], [], None, self.root / 'state', None, self.library, True)
        runtime.send = Mock(); runtime.submit = Mock()
        runtime.start('chat', {'text':'查一下 files 里面的时间表', 'sender':'staff','id':'one'})
        context = runtime.submit.call_args.args[1]
        self.assertEqual(len(context['file_library']['files']), 1)
        self.assertEqual(context['file_evidence'], [])
        fid = context['file_library']['files'][0]['id']
        runtime.accept_decision('chat', {'decision':'inspect_files','file_reads':[{'file_id':fid,'offset':0}]},
                                {'rounds':0,'repairs':0})
        self.assertIn('10 October', runtime.submit.call_args.args[1]['file_evidence'][0]['text'])
        self.assertEqual(runtime.flows['chat'].state['phase'], 'PLANNING')
        browser.execute.assert_not_called()

    def test_unapproved_planner_context_omits_library_data(self):
        (self.sources / 'private.txt').write_text('private content')
        browser = Mock(); browser.overviews = {}; browser.forms = {}
        runtime = Runtime(None,browser,['chat'],[],None,self.root/'state',None,self.library)
        runtime.send = Mock(); runtime.submit = Mock()
        runtime.start('chat',{'text':'Check course settings','sender':'staff','id':'one'})
        context = runtime.submit.call_args.args[1]
        self.assertEqual(context['file_library']['status'],'planner_access_not_authorized')
        self.assertNotIn('private',str(context))
        self.assertNotIn(str(self.sources),str(context))
        with self.assertRaisesRegex(ValueError,'not authorized'):
            runtime.accept_decision('chat',{'decision':'inspect_files','file_reads':[]},{'rounds':0,'repairs':0})

    def cached_runtime(self,text='Syllabus week 5'):
        (self.sources/'syllabus.txt').write_text(text)
        browser=Mock();browser.overviews={};browser.forms={}
        runtime=Runtime(None,browser,['chat'],[],None,self.root/'state',None,self.library,True)
        runtime.send=Mock();runtime.submit=Mock()
        runtime.start('chat',{'text':'Align the syllabus','sender':'staff','id':'one'})
        runtime.send.reset_mock()
        fid=runtime.file_snapshot['files'][0]['id']
        runtime.file_evidence=[self.library.inspect(runtime.file_snapshot,fid)]
        return runtime,fid

    def test_real_failure_duplicate_source_after_used_repair_is_successful_read(self):
        from concurrent.futures import Future
        runtime,fid=self.cached_runtime()
        future=Future();future.set_result({'decision':'inspect_files','file_reads':[{'file_id':fid,'offset':0}]})
        runtime.jobs['chat']={'future':future,'rounds':3,'repairs':1,
                             'context':{'feedback':'Syllabus row week and proposed destination differ.'}}
        with patch.object(self.library,'inspect',wraps=self.library.inspect) as read:
            runtime.complete_jobs();read.assert_not_called()
        context=runtime.submit.call_args.args[1]
        self.assertEqual(runtime.flows['chat'].state['phase'],'PLANNING')
        self.assertEqual(len(context['file_evidence']),1)
        self.assertIn('Syllabus week 5',context['file_evidence'][0]['text'])
        self.assertIn('proposed destination differ',context['feedback'])
        self.assertIn('returned successfully',context['feedback'])
        self.assertEqual(runtime.submit.call_args.args[2:],(4,1))
        runtime.send.assert_not_called();runtime.browser.execute.assert_not_called()

    def test_cached_source_changed_on_disk_is_not_silently_reused(self):
        runtime,fid=self.cached_runtime()
        (self.sources/'syllabus.txt').write_text('Changed authoritative source')
        with self.assertRaisesRegex(ValueError,'changed'):
            runtime.accept_decision('chat',{'decision':'inspect_files','file_reads':[{'file_id':fid,'offset':0}]}, {'rounds':1,'repairs':1})

    def test_mixed_cached_and_next_page_retains_correction_and_reads_only_next(self):
        runtime,fid=self.cached_runtime('a'*(PAGE_SIZE+50))
        next_page=runtime.file_evidence[0]['next_offset']
        with patch.object(self.library,'inspect',wraps=self.library.inspect) as read:
            runtime.accept_decision('chat',{'decision':'inspect_files','file_reads':[{'file_id':fid,'offset':0},{'file_id':fid,'offset':next_page}]},
                                    {'rounds':1,'repairs':1,'context':{'feedback':'Correct the cited row'}})
            read.assert_called_once_with(runtime.file_snapshot,file_id=fid,offset=next_page)
        self.assertEqual(len(runtime.file_evidence),2)
        self.assertIn('Correct the cited row',runtime.submit.call_args.args[1]['feedback'])

    def test_repeated_reads_remain_bounded_and_unknown_offsets_fail(self):
        runtime,fid=self.cached_runtime()
        request={'decision':'inspect_files','file_reads':[{'file_id':fid,'offset':0}]}
        with self.assertRaisesRegex(ValueError,'budget'):
            runtime.accept_decision('chat',request,{'rounds':8,'repairs':1})
        request['file_reads'][0]['offset']=99
        with self.assertRaisesRegex(ValueError,'next_offset'):
            runtime.accept_decision('chat',request,{'rounds':1,'repairs':1})

    def test_duplicate_forms_reuse_evidence_without_losing_review_feedback(self):
        runtime,_=self.cached_runtime()
        runtime.browser.forms={'a:module:1':{'fields':[]}}
        runtime.accept_decision('chat',{'decision':'inspect','inspect_targets':['a:module:1']},
                                {'rounds':2,'repairs':1,'context':{'feedback':'Fix visibility scope'}})
        runtime.browser.inspect.assert_not_called()
        self.assertIn('Fix visibility scope',runtime.submit.call_args.args[1]['feedback'])
        self.assertEqual(runtime.flows['chat'].state['phase'],'PLANNING')

    def test_eight_source_reads_are_batched_without_spending_plan_repair(self):
        for n in range(8):(self.sources/f'source{n}.txt').write_text('Source '+str(n))
        runtime,_=self.cached_runtime();runtime.poll_safely=Mock()
        requests=[{'file_id':e['id'],'offset':0} for e in runtime.file_snapshot['files'][:8]]
        runtime.accept_decision('chat',{'decision':'inspect_files','file_reads':requests},
                                {'rounds':2,'repairs':0,'context':{'feedback':'Resolve topic ambiguity'}})
        runtime.poll_safely.assert_called_once()
        self.assertEqual(runtime.submit.call_args.args[2:],(4,0))
        self.assertIn('Resolve topic ambiguity',runtime.submit.call_args.args[1]['feedback'])
        self.assertTrue(all(any(e['file_id']==r['file_id'] for e in runtime.file_evidence) for r in requests))
        with self.assertRaisesRegex(ValueError,'budget'):
            runtime.accept_decision('chat',{'decision':'inspect_files','file_reads':requests},{'rounds':7,'repairs':0})

    def test_default_configuration_and_typed_inspection(self):
        config, _ = load_config()
        self.assertEqual(Path(config['file_library_root']).name, 'files')
        self.assertIn('inspect_files', schema()['properties']['decision']['enum'])


if __name__ == '__main__': unittest.main()
