from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import shutil
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

try:
    from timeline_reconciliation_agent import (
        AGENT_NAME as RECONCILIATION_AGENT_NAME,
        run_timeline_reconciliation_agent,
    )
except ModuleNotFoundError:  # pragma: no cover - import shape depends on entrypoint cwd
    from course_agent_app.timeline_reconciliation_agent import (
        AGENT_NAME as RECONCILIATION_AGENT_NAME,
        run_timeline_reconciliation_agent,
    )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DRAFT_ROOT = (
    PROJECT_ROOT
    / "course_files"
    / "evidence"
    / "process"
    / "course_agent_app"
    / "timeline_draft_review"
)
ACTIVE_DRAFT_CONTEXT_PATH = (
    PROJECT_ROOT
    / "course_files"
    / "evidence"
    / "process"
    / "course_agent_app"
    / "latest_timeline_draft_context.json"
)
REVISE_OPERATION_KIND = "revise_timeline_draft"
COMMIT_OPERATION_KIND = "commit_timeline_draft"
EXECUTION_MODE = "local_file_action"
CONTENT_KINDS = ("lecture", "workshop", "submission")
WEEKDAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
REVIEW_WEEK_FIELD = "calendar_week"
REVIEW_WEEK_LABEL = "Calendar Week"
FORBIDDEN_JOB_KEYS = {
    "command",
    "shell",
    "python",
    "python_code",
    "javascript",
    "javascript_code",
    "raw_code",
    "script_text",
    "replacements",
    "text_replacements",
}


def now_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")


def safe_name(value: str, fallback: str = "timeline-draft") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]+", "-", str(value or "")).strip(" .-_")
    cleaned = re.sub(r"\s+", "-", cleaned)
    return cleaned[:120] or fallback


def resolve_local_path(path_text: str | Path) -> Path:
    path = Path(str(path_text or "")).expanduser()
    if path.is_absolute():
        return path
    for base in (PROJECT_ROOT, Path.cwd()):
        candidate = base / path
        if candidate.exists():
            return candidate
    return PROJECT_ROOT / path


def repo_rel(path: str | Path) -> str:
    resolved = resolve_local_path(path)
    try:
        return str(resolved.resolve(strict=False).relative_to(PROJECT_ROOT))
    except ValueError:
        return str(resolved)


def review_week_numbering(draft: dict[str, Any]) -> dict[str, str]:
    configured = (
        draft.get("review_week_numbering")
        if isinstance(draft.get("review_week_numbering"), dict)
        else {}
    )
    display_field = str(configured.get("display_field") or REVIEW_WEEK_FIELD)
    if display_field not in {"calendar_week", "teaching_week"}:
        display_field = REVIEW_WEEK_FIELD
    default_label = "Teaching Week" if display_field == "teaching_week" else REVIEW_WEEK_LABEL
    return {
        "display_field": display_field,
        "display_label": str(configured.get("display_label") or default_label),
        "unqualified_week_reference": "display_field",
    }


def load_json(path: Path) -> dict[str, Any] | None:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    return payload if isinstance(payload, dict) else None


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _params(operation: dict[str, Any]) -> dict[str, Any]:
    return operation.get("params") if isinstance(operation.get("params"), dict) else {}


def is_revise_timeline_draft_operation(operation: dict[str, Any]) -> bool:
    return str(operation.get("operation_kind") or "").strip() == REVISE_OPERATION_KIND


def is_commit_timeline_draft_operation(operation: dict[str, Any]) -> bool:
    return str(operation.get("operation_kind") or "").strip() == COMMIT_OPERATION_KIND


def validate_timeline_draft_operation_shape(
    operation: dict[str, Any],
    *,
    operation_kind: str,
) -> list[str]:
    params = _params(operation)
    errors: list[str] = []
    forbidden = sorted(FORBIDDEN_JOB_KEYS & set(params))
    if forbidden:
        errors.append(f"{operation_kind}.params contains forbidden fields: " + ", ".join(forbidden))
    draft_id = str(params.get("draft_id") or operation.get("draft_id") or "").strip()
    if not draft_id:
        errors.append(f"{operation_kind} requires params.draft_id from the active timeline review")
    raw_version = params.get("draft_version", operation.get("draft_version"))
    try:
        version = int(raw_version)
    except (TypeError, ValueError):
        version = 0
    if version < 1:
        errors.append(f"{operation_kind} requires a positive params.draft_version")
    return errors


def validate_revise_timeline_draft_shape(operation: dict[str, Any]) -> list[str]:
    return validate_timeline_draft_operation_shape(
        operation,
        operation_kind=REVISE_OPERATION_KIND,
    )


def validate_commit_timeline_draft_shape(operation: dict[str, Any]) -> list[str]:
    return validate_timeline_draft_operation_shape(
        operation,
        operation_kind=COMMIT_OPERATION_KIND,
    )


def active_timeline_draft_context() -> dict[str, Any]:
    context = load_json(ACTIVE_DRAFT_CONTEXT_PATH) or {}
    draft_path_text = str(context.get("draft_path") or "").strip()
    if not draft_path_text:
        return {}
    draft_path = resolve_local_path(draft_path_text)
    draft = load_json(draft_path)
    if not draft:
        return {}
    context["draft_path"] = repo_rel(draft_path)
    context["draft"] = draft
    return context


def planner_safe_active_timeline_draft_context(requester: str = "") -> dict[str, Any]:
    context = active_timeline_draft_context()
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    context_status = str(context.get("status") or "")
    if (
        not draft
        or context_status in {"COMMITTED", "SUPERSEDED"}
        or str(draft.get("status") or "") in {"COMMITTED", "SUPERSEDED"}
    ):
        return {}
    bound_requester = str(draft.get("whatsapp_requester") or "").strip()
    if bound_requester and str(requester or "").strip().casefold() != bound_requester.casefold():
        return {}
    items = [
        {
            "item_id": item.get("item_id"),
            "source_teaching_week": item.get("source_teaching_week"),
            "source_visible_week": item.get("source_visible_week"),
            "content_kind": item.get("content_kind"),
            "text": item.get("text"),
            "target_teaching_week": item.get("target_teaching_week"),
            "status": item.get("status"),
        }
        for item in draft.get("content_items") or []
        if isinstance(item, dict)
    ]
    return {
        "draft_id": draft.get("draft_id"),
        "draft_version": draft.get("version"),
        "status": draft.get("status"),
        "course_code": draft.get("course_code"),
        "semester": draft.get("semester"),
        "origin": draft.get("origin"),
        "valid_teaching_week_count": draft.get("valid_teaching_week_count"),
        "overflow_count": draft.get("overflow_count"),
        "content_items": items,
        "review_session_active": True,
    }


def _supersede_active_draft(new_draft_id: str) -> None:
    context = active_timeline_draft_context()
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else {}
    if not draft or str(draft.get("draft_id") or "") == new_draft_id:
        return
    status = str(draft.get("status") or "")
    if status in {"COMMITTED", "SUPERSEDED"}:
        return
    marker = resolve_local_path(str(context["draft_path"])).parent / "superseded.json"
    write_json(
        marker,
        {
            "draft_id": draft.get("draft_id"),
            "version": draft.get("version"),
            "status": "SUPERSEDED",
            "superseded_by": new_draft_id,
            "recorded_at": datetime.now(timezone.utc).isoformat(),
        },
    )


def _write_active_context(draft: dict[str, Any], draft_path: Path, preview_paths: dict[str, str]) -> None:
    write_json(
        ACTIVE_DRAFT_CONTEXT_PATH,
        {
            "draft_id": draft["draft_id"],
            "draft_version": draft["version"],
            "status": draft["status"],
            "course_code": draft["course_code"],
            "semester": draft["semester"],
        "origin": draft["origin"],
        "whatsapp_requester": draft.get("whatsapp_requester", ""),
        "draft_path": repo_rel(draft_path),
            "preview_docx_path": preview_paths.get("docx_path", ""),
            "preview_pdf_path": preview_paths.get("pdf_path", ""),
            "updated_at": datetime.now(timezone.utc).isoformat(),
        },
    )


def _draft_id(run_id: str, course_code: str, semester: str) -> str:
    seed = f"{run_id}|{course_code}|{semester}|{datetime.now(timezone.utc).isoformat()}"
    suffix = hashlib.sha256(seed.encode("utf-8")).hexdigest()[:8].upper()
    return f"TD-{now_stamp()}-{suffix}"


def _first_semester_facts(calendar_payload: dict[str, Any]) -> dict[str, Any]:
    semesters = calendar_payload.get("semesters")
    if isinstance(semesters, dict) and isinstance(semesters.get("first"), dict):
        return semesters["first"]
    if isinstance(semesters, list):
        for item in semesters:
            name = str(item.get("name") or "").strip().lower() if isinstance(item, dict) else ""
            if isinstance(item, dict) and name in {"first", "first semester", "semester 1", "sem 1"}:
                return item
    return {}


def _calendar_week(calendar_payload: dict[str, Any], value: date) -> int:
    facts = _first_semester_facts(calendar_payload)
    origin = date.fromisoformat(str(facts.get("first_teaching") or facts.get("start")))
    return ((value - origin).days // 7) + 1


def _overflow_slots(
    calendar_payload: dict[str, Any],
    class_weekdays: list[str],
    *,
    first_overflow_week: int,
    overflow_count: int,
    last_valid_date: date,
) -> list[dict[str, Any]]:
    try:
        from semester_schedule_update import calendar_blocked_teaching_dates
    except ModuleNotFoundError:  # pragma: no cover
        from course_agent_app.semester_schedule_update import calendar_blocked_teaching_dates

    facts = _first_semester_facts(calendar_payload)
    last_teaching = date.fromisoformat(str(facts.get("last_teaching") or facts.get("end")))
    semester_end = date.fromisoformat(str(facts.get("end") or facts.get("assessment_end") or last_teaching))
    blocked = calendar_blocked_teaching_dates({"calendar.json": calendar_payload})
    weekday_indexes = [WEEKDAY_NAMES.index(day) for day in class_weekdays]
    week_start = last_valid_date - timedelta(days=last_valid_date.weekday()) + timedelta(days=7)
    slots: list[dict[str, Any]] = []
    for offset in range(overflow_count):
        candidates = [week_start + timedelta(days=index) for index in weekday_indexes]
        reasons: list[str] = []
        for candidate in candidates:
            if candidate > last_teaching:
                reasons.append("after last_teaching")
            if candidate > semester_end:
                reasons.append("outside semester range")
            if candidate in blocked:
                reasons.append(blocked[candidate])
        reasons = list(dict.fromkeys(reasons))
        teaching_week = first_overflow_week + offset
        slots.append(
            {
                "teaching_week": teaching_week,
                "calendar_week": _calendar_week(calendar_payload, candidates[0]),
                "teaching_dates": [candidate.isoformat() for candidate in candidates],
                "allowed_session_weekdays": list(class_weekdays),
                "valid": False,
                "status": "Overflow - " + "; ".join(reasons or ["outside valid teaching slots"]),
                "conflicts": reasons or ["outside valid teaching slots"],
            }
        )
        week_start += timedelta(days=7)
    return slots


def _content_items(source_timeline: dict[str, Any], valid_count: int) -> list[dict[str, Any]]:
    try:
        from reference_pack_rebuild import teaching_content_rows
    except ModuleNotFoundError:  # pragma: no cover
        from course_agent_app.reference_pack_rebuild import teaching_content_rows

    items: list[dict[str, Any]] = []
    for fallback_teaching_week, row in enumerate(teaching_content_rows(source_timeline), start=1):
        source_teaching_week = int(
            row.get("source_teaching_week")
            or fallback_teaching_week
        )
        source_visible_week = int(
            row.get("source_visible_week")
            or row.get("source_week")
            or row.get("teaching_week")
            or row.get("week")
            or source_teaching_week
        )
        for content_kind in CONTENT_KINDS:
            text = re.sub(r"\s+", " ", str(row.get(content_kind) or "")).strip()
            if not text:
                continue
            item_id = f"tw{source_teaching_week}-{content_kind}"
            items.append(
                {
                    "item_id": item_id,
                    "source_teaching_week": source_teaching_week,
                    "source_visible_week": source_visible_week,
                    "source_date": str(row.get("source_date") or row.get("date") or ""),
                    "source_date_iso": str(row.get("source_date_iso") or ""),
                    "source_weekday": str(row.get("source_weekday") or ""),
                    "source_calendar_week": row.get("source_calendar_week"),
                    "source_semester": str(row.get("source_semester") or ""),
                    "source_academic_year": str(row.get("source_academic_year") or ""),
                    "content_kind": content_kind,
                    "text": text,
                    "target_teaching_week": source_teaching_week,
                    "status": "scheduled" if source_teaching_week <= valid_count else "overflow",
                }
            )
    return items


def _draft_status(content_items: list[dict[str, Any]], valid_count: int) -> tuple[str, int]:
    overflow = [
        item
        for item in content_items
        if item.get("status") != "deleted"
        and (
            not isinstance(item.get("target_teaching_week"), int)
            or int(item["target_teaching_week"]) < 1
            or int(item["target_teaching_week"]) > valid_count
        )
    ]
    return (
        "READY_FOR_CONFIRMATION" if not overflow else "NEEDS_TIMELINE_REVIEW",
        len(overflow),
    )


def create_timeline_draft(
    *,
    origin: str,
    run_id: str,
    raw_task: str,
    course_code: str,
    semester: str,
    class_weekdays: list[str],
    reference_pack_path: Path,
    source_timeline: dict[str, Any],
    calendar_payload: dict[str, Any],
    candidate_payloads: dict[str, Any] | None = None,
    staged_sources: list[dict[str, Any]] | None = None,
    env: dict[str, str] | None = None,
) -> dict[str, Any]:
    try:
        from semester_schedule_update import (
            base_schedule_payload,
            recurring_teaching_week_dates,
            validate_schedule_payload,
        )
        from reference_pack_rebuild import (
            non_teaching_gaps_from_schedule,
            teaching_content_rows,
        )
    except ModuleNotFoundError:  # pragma: no cover
        from course_agent_app.semester_schedule_update import (
            base_schedule_payload,
            recurring_teaching_week_dates,
            validate_schedule_payload,
        )
        from course_agent_app.reference_pack_rebuild import (
            non_teaching_gaps_from_schedule,
            teaching_content_rows,
        )

    required_count = len(teaching_content_rows(source_timeline))
    if required_count < 1:
        return {
            "ok": False,
            "status": "BLOCKED",
            "blocker": "Timeline draft requires at least one teaching content row.",
            "attachments": [],
        }
    references = {"calendar.json": calendar_payload, "timeline.json": source_timeline}
    dates_by_week = recurring_teaching_week_dates(
        references,
        class_weekdays,
        through_week=required_count,
    )
    valid_count = len(dates_by_week)
    if valid_count >= required_count:
        return {
            "ok": False,
            "status": "BLOCKED",
            "blocker": "Timeline draft is only created when the requested content exceeds valid teaching slots.",
            "attachments": [],
        }
    if valid_count < 1:
        return {
            "ok": False,
            "status": "BLOCKED",
            "blocker": "The calendar does not provide any valid teaching slot for the requested weekdays.",
            "attachments": [],
        }
    schedule_weeks: list[dict[str, Any]] = []
    for week in range(1, valid_count + 1):
        dates = dates_by_week[week]
        first_date = date.fromisoformat(dates[0])
        schedule_weeks.append(
            {
                "week": week,
                "teaching_week": week,
                "calendar_week": _calendar_week(calendar_payload, first_date),
                "label": f"Teaching Week {week}",
                "teaching_dates": dates,
                "allowed_session_weekdays": list(class_weekdays),
                "notes": "Valid teaching slot derived from the staff-provided recurring weekdays.",
            }
        )
    schedule = base_schedule_payload(
        course_code,
        semester,
        ["timeline draft review", "current staff class weekdays"],
        schedule_weeks,
    )
    schedule["provenance"] = {
        "origin": f"{origin}_timeline_draft",
        "user_supplied": True,
        "merge_mode": "replace_all",
    }
    schedule["mapping_rules"].update(
        {
            "lab_visit": "use_content_week_mapping_then_target_teaching_dates",
            "lecture": "use_content_week_mapping_then_target_teaching_dates",
            "tutorial": "use_content_week_mapping_then_target_teaching_dates",
            "session_date": "use_content_week_mapping_then_target_teaching_dates",
            "assignment_deadline": "use_submission_content_week_mapping_then_preserve_weekday_and_time",
            "submission_deadline": "use_submission_content_week_mapping_then_preserve_weekday_and_time",
        }
    )
    validation_errors = validate_schedule_payload(schedule, references)
    if validation_errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "blocker": "; ".join(validation_errors),
            "attachments": [],
        }
    last_valid_date = max(
        date.fromisoformat(value)
        for values in dates_by_week.values()
        for value in values
    )
    overflow_slots = _overflow_slots(
        calendar_payload,
        class_weekdays,
        first_overflow_week=valid_count + 1,
        overflow_count=required_count - valid_count,
        last_valid_date=last_valid_date,
    )
    valid_slots = [
        {
            "teaching_week": int(item["teaching_week"]),
            "calendar_week": int(item["calendar_week"]),
            "teaching_dates": list(item["teaching_dates"]),
            "allowed_session_weekdays": list(item["allowed_session_weekdays"]),
            "valid": True,
            "status": "Valid teaching week",
            "conflicts": [],
        }
        for item in schedule["weeks"]
    ]
    items = _content_items(source_timeline, valid_count)
    status, overflow_count = _draft_status(items, valid_count)
    draft_id = _draft_id(run_id, course_code, semester)
    _supersede_active_draft(draft_id)
    draft = {
        "draft_id": draft_id,
        "version": 1,
        "status": status,
        "origin": origin,
        "run_id": run_id,
        "last_review_run_id": run_id,
        "whatsapp_requester": str((env or {}).get("_WHATSAPP_REQUESTER") or "").strip(),
        "whatsapp_root_message_id": str((env or {}).get("_WHATSAPP_ROOT_MESSAGE_ID") or "").strip(),
        "raw_task": raw_task,
        "course_code": course_code,
        "semester": semester,
        "class_weekdays": list(class_weekdays),
        "reference_pack_path": repo_rel(reference_pack_path),
        "required_teaching_week_count": required_count,
        "valid_teaching_week_count": valid_count,
        "overflow_count": overflow_count,
        "review_week_numbering": {
            "display_field": REVIEW_WEEK_FIELD,
            "display_label": REVIEW_WEEK_LABEL,
            "unqualified_week_reference": "display_field",
        },
        "teaching_slots": [*valid_slots, *overflow_slots],
        "non_teaching_gaps": non_teaching_gaps_from_schedule(schedule, calendar_payload),
        "content_items": items,
        "candidate_schedule": schedule,
        "candidate_payloads": {
            **(candidate_payloads or {}),
            "calendar.json": calendar_payload,
            "source_timeline.json": source_timeline,
        },
        "staged_sources": list(staged_sources or []),
        "revision_history": [],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    draft_path, preview = save_draft_version(draft, env or {})
    _write_active_context(draft, draft_path, preview)
    attachments = [preview["pdf_path"]] if preview.get("pdf_path") else [preview["docx_path"]]
    return {
        "ok": True,
        "status": "NEEDS_TIMELINE_REVIEW",
        "operation_kind": origin,
        "run_id": run_id,
        "draft_id": draft_id,
        "draft_version": 1,
        "draft_status": status,
        "overflow_count": overflow_count,
        "valid_teaching_week_count": valid_count,
        "required_teaching_week_count": required_count,
        "draft_path": repo_rel(draft_path),
        "preview_docx_path": preview["docx_path"],
        "preview_pdf_path": preview.get("pdf_path", ""),
        "attachments": attachments,
        "reference_pack_path": repo_rel(reference_pack_path),
        "canonical_files_changed": False,
    }


def _set_cell_shading(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_width(cell: Any, width_dxa: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def _set_cell_margins(cell: Any, *, top: int = 70, start: int = 90, bottom: int = 70, end: int = 90) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _draft_table_rows(draft: dict[str, Any]) -> list[dict[str, Any]]:
    items = [
        item
        for item in draft.get("content_items") or []
        if isinstance(item, dict) and item.get("status") != "deleted"
    ]
    rows: list[dict[str, Any]] = []
    for slot in draft.get("teaching_slots") or []:
        if not isinstance(slot, dict):
            continue
        week = int(slot["teaching_week"])
        week_items = [item for item in items if item.get("target_teaching_week") == week]
        row = {
            "row_type": "teaching" if slot.get("valid") else "overflow",
            "teaching_week": week,
            "calendar_week": slot.get("calendar_week"),
            "date": "\n".join(str(item) for item in slot.get("teaching_dates") or []),
            "status": slot.get("status"),
            "lecture": "",
            "workshop": "",
            "submission": "",
        }
        for content_kind in CONTENT_KINDS:
            values = [
                str(item.get("text") or "")
                for item in sorted(
                    [item for item in week_items if item.get("content_kind") == content_kind],
                    key=lambda item: (
                        int(item.get("source_teaching_week") or 0),
                        str(item.get("item_id") or ""),
                    ),
                )
            ]
            row[content_kind] = "\n".join(value for value in values if value)
        rows.append(row)
    for gap in draft.get("non_teaching_gaps") or []:
        if not isinstance(gap, dict):
            continue
        rows.append(
            {
                "row_type": "gap",
                "teaching_week": "",
                "calendar_week": gap.get("calendar_week"),
                "date": gap.get("date"),
                "status": gap.get("reason"),
                "lecture": gap.get("reason"),
                "workshop": gap.get("reason"),
                "submission": "",
            }
        )
    return sorted(
        rows,
        key=lambda item: (
            int(item.get("calendar_week") or 9999),
            0 if item.get("row_type") == "gap" else 1,
        ),
    )


def render_draft_docx(draft: dict[str, Any], output_path: Path) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.4)
    section.bottom_margin = Inches(0.4)
    section.left_margin = Inches(0.35)
    section.right_margin = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run(f"{draft.get('course_code')} Timeline Draft")
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(46, 116, 181)
    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(7)
    meta_run = meta.add_run(
        f"{draft.get('semester')}  |  Version {draft.get('version')}  |  "
        f"{draft.get('status')}  |  {draft.get('overflow_count')} overflow item(s)"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(85, 85, 85)

    week_numbering = review_week_numbering(draft)
    headers = [
        week_numbering["display_label"].replace(" ", "\n"),
        "Date",
        "Status",
        "Lecture Theme",
        "Tutorials and Workshop",
        "Submission",
    ]
    widths = [850, 1150, 1600, 2250, 2650, 1820]
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header_row = table.rows[0]
    header_row._tr.get_or_add_trPr().append(
        __import__("docx").oxml.OxmlElement("w:tblHeader")
    )
    for index, (heading, width) in enumerate(zip(headers, widths)):
        cell = header_row.cells[index]
        cell.text = heading
        _set_cell_width(cell, width)
        _set_cell_margins(cell)
        _set_cell_shading(cell, "E8EEF5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(8)

    for row_data in _draft_table_rows(draft):
        cells = table.add_row().cells
        values = [
            row_data.get(week_numbering["display_field"], ""),
            row_data.get("date", ""),
            row_data.get("status", ""),
            row_data.get("lecture", ""),
            row_data.get("workshop", ""),
            row_data.get("submission", ""),
        ]
        fill = "FDE9E7" if row_data["row_type"] == "overflow" else "FFF4CE" if row_data["row_type"] == "gap" else "FFFFFF"
        for index, (value, width) in enumerate(zip(values, widths)):
            cell = cells[index]
            cell.text = str(value or "")
            _set_cell_width(cell, width)
            _set_cell_margins(cell)
            _set_cell_shading(cell, fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if index < 3:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(7.5)
                if row_data["row_type"] == "overflow" and index == 2:
                    run.bold = True
                    run.font.color.rgb = RGBColor(156, 0, 6)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def save_draft_version(
    draft: dict[str, Any],
    env: dict[str, str],
) -> tuple[Path, dict[str, str]]:
    version_dir = DRAFT_ROOT / safe_name(str(draft["draft_id"])) / f"v{int(draft['version'])}"
    version_dir.mkdir(parents=True, exist_ok=True)
    draft_path = version_dir / "timeline_draft.json"
    if draft_path.exists():
        raise FileExistsError(
            f"Timeline draft {draft['draft_id']} version {draft['version']} already exists."
        )
    write_json(draft_path, draft)
    draft_parts = safe_name(str(draft["draft_id"])).split("-")
    preview_identity = "-".join(draft_parts[1:4]) if len(draft_parts) >= 4 else safe_name(str(draft["draft_id"]))
    course_code = safe_name(str(draft.get("course_code") or "Course"))
    docx_path = version_dir / (
        f"{course_code} Timeline Draft {preview_identity} v{int(draft['version'])}.docx"
    )
    render_draft_docx(draft, docx_path)
    pdf_path = ""
    try:
        from file_refresh_preview import convert_to_pdf
    except ModuleNotFoundError:  # pragma: no cover
        from course_agent_app.file_refresh_preview import convert_to_pdf
    conversion = convert_to_pdf(docx_path, version_dir, env)
    if conversion.get("ok"):
        pdf_path = str(conversion.get("pdf_path") or "")
    return draft_path, {
        "docx_path": repo_rel(docx_path),
        "pdf_path": pdf_path,
    }


def _operation_draft_identity(operation: dict[str, Any]) -> tuple[str, int]:
    params = _params(operation)
    draft_id = str(params.get("draft_id") or operation.get("draft_id") or "").strip()
    try:
        version = int(params.get("draft_version", operation.get("draft_version")))
    except (TypeError, ValueError):
        version = 0
    return draft_id, version


def _active_draft_for_operation(
    operation: dict[str, Any],
    *,
    operation_kind: str,
) -> tuple[dict[str, Any] | None, str]:
    context = active_timeline_draft_context()
    draft = context.get("draft") if isinstance(context.get("draft"), dict) else None
    if not draft:
        return None, "There is no active timeline draft review."
    draft_id, version = _operation_draft_identity(operation)
    if draft_id != str(draft.get("draft_id") or "") or version != int(draft.get("version") or 0):
        return None, (
            f"{operation_kind} targets a stale or different draft version; "
            f"the active review is {draft.get('draft_id')} version {draft.get('version')}."
        )
    context_status = str(context.get("status") or "")
    terminal_status = context_status if context_status in {"COMMITTED", "SUPERSEDED"} else str(draft.get("status") or "")
    if terminal_status in {"COMMITTED", "SUPERSEDED"}:
        return None, f"The active timeline draft is already {terminal_status.lower()}."
    return draft, ""


def _workflow_requester_blocker(draft: dict[str, Any], env: dict[str, str]) -> str:
    expected = str(draft.get("whatsapp_requester") or "").strip()
    observed = str(env.get("_WHATSAPP_REQUESTER") or "").strip()
    if expected and observed.casefold() != expected.casefold():
        return "This timeline review is bound to the Staff member who started it."
    return ""


def _planner_blockers(output: dict[str, Any]) -> list[str]:
    blockers = output.get("blockers") if isinstance(output.get("blockers"), list) else []
    return [
        str(item.get("reason") or item.get("message") or item)
        if isinstance(item, dict)
        else str(item)
        for item in blockers
        if str(item).strip()
    ]


def validate_reconciliation_output(
    output: dict[str, Any],
    draft: dict[str, Any],
) -> list[str]:
    errors: list[str] = []
    if str(output.get("agent") or "") != RECONCILIATION_AGENT_NAME:
        return ["Timeline revision was not produced by the Timeline Reconciliation Agent."]
    if str(output.get("status") or "").lower() != "ready":
        return _planner_blockers(output) or ["Timeline Reconciliation Agent did not return ready."]
    if str(output.get("draft_id") or "") != str(draft.get("draft_id") or ""):
        errors.append("Timeline revision draft_id does not match the active draft.")
    try:
        base_version = int(output.get("base_version"))
    except (TypeError, ValueError):
        base_version = 0
    if base_version != int(draft.get("version") or 0):
        errors.append("Timeline revision base_version does not match the active draft.")
    reference_field = str(output.get("week_reference_field") or "").strip()
    if reference_field not in {"calendar_week", "teaching_week", "none"}:
        errors.append(
            "Timeline revision must declare week_reference_field as calendar_week, "
            "teaching_week, or none."
        )
    resolutions = output.get("week_reference_resolution")
    if not isinstance(resolutions, list):
        errors.append("Timeline revision must provide week_reference_resolution as a list.")
        resolutions = []
    if reference_field == "none" and resolutions:
        errors.append("Timeline revision with week_reference_field none cannot resolve numbered weeks.")
    if reference_field in {"calendar_week", "teaching_week"} and not resolutions:
        errors.append(
            "Timeline revision using numbered weeks must record how each visible week resolves "
            "to an internal Teaching Week."
        )
    operations = output.get("operations")
    if not isinstance(operations, list) or not operations:
        errors.append("Timeline revision must contain at least one move_items or delete_items operation.")
        return errors
    items = {
        str(item.get("item_id")): item
        for item in draft.get("content_items") or []
        if isinstance(item, dict)
    }
    valid_weeks = {
        int(slot["teaching_week"])
        for slot in draft.get("teaching_slots") or []
        if isinstance(slot, dict) and slot.get("valid") is True
    }
    slots = [
        slot
        for slot in draft.get("teaching_slots") or []
        if isinstance(slot, dict)
    ]
    resolved_roles: dict[str, int] = {}
    for index, resolution in enumerate(resolutions, start=1):
        if not isinstance(resolution, dict):
            errors.append(f"Timeline week reference {index} must be an object.")
            continue
        role = str(resolution.get("role") or "")
        if role not in {"source", "target"}:
            errors.append(f"Timeline week reference {index} must have role source or target.")
            continue
        try:
            display_week = int(resolution.get("display_week"))
            resolved_week = int(resolution.get("resolved_teaching_week"))
        except (TypeError, ValueError):
            errors.append(f"Timeline week reference {index} requires integer week values.")
            continue
        if reference_field == "none":
            continue
        matching_slots = [
            slot
            for slot in slots
            if int(slot.get(reference_field) or 0) == display_week
        ]
        if len(matching_slots) != 1:
            errors.append(
                f"Timeline week reference {index} does not uniquely match "
                f"{reference_field} {display_week}."
            )
            continue
        expected_week = int(matching_slots[0].get("teaching_week") or 0)
        if resolved_week != expected_week:
            errors.append(
                f"Timeline week reference {index} resolves {reference_field} {display_week} "
                f"to Teaching Week {resolved_week}, expected Teaching Week {expected_week}."
            )
            continue
        resolved_roles[role] = resolved_week
    seen: set[str] = set()
    allowed_keys = {
        "move_items": {"type", "item_ids", "target_teaching_week", "reason"},
        "delete_items": {"type", "item_ids", "reason"},
    }
    for index, action in enumerate(operations, start=1):
        if not isinstance(action, dict):
            errors.append(f"Timeline revision operation {index} must be an object.")
            continue
        action_type = str(action.get("type") or "")
        if action_type not in allowed_keys:
            errors.append(f"Timeline revision operation {index} has unsupported type {action_type!r}.")
            continue
        extra = sorted(set(action) - allowed_keys[action_type])
        if extra:
            errors.append(
                f"Timeline revision operation {index} contains unsupported fields: {', '.join(extra)}."
            )
        item_ids = action.get("item_ids")
        if not isinstance(item_ids, list) or not item_ids:
            errors.append(f"Timeline revision operation {index} requires non-empty item_ids.")
            continue
        for item_id in item_ids:
            item_key = str(item_id)
            if item_key not in items:
                errors.append(f"Timeline revision operation {index} references unknown item_id {item_key!r}.")
            if item_key in seen:
                errors.append(f"Timeline revision item_id {item_key!r} appears more than once.")
            seen.add(item_key)
        if action_type == "move_items":
            try:
                target_week = int(action.get("target_teaching_week"))
            except (TypeError, ValueError):
                target_week = 0
            if target_week not in valid_weeks:
                errors.append(
                    f"Timeline revision operation {index} targets invalid Teaching Week {target_week}."
                )
    first_action = operations[0] if operations and isinstance(operations[0], dict) else {}
    first_item_ids = (
        [str(item) for item in first_action.get("item_ids") or []]
        if isinstance(first_action.get("item_ids"), list)
        else []
    )
    if "source" in resolved_roles and first_item_ids:
        source_week = resolved_roles["source"]
        if any(
            int(items[item_id].get("target_teaching_week") or 0) != source_week
            for item_id in first_item_ids
            if item_id in items
        ):
            errors.append(
                "The primary revision operation does not select content from the resolved "
                f"source Teaching Week {source_week}."
            )
    if "target" in resolved_roles and str(first_action.get("type") or "") == "move_items":
        try:
            first_target = int(first_action.get("target_teaching_week"))
        except (TypeError, ValueError):
            first_target = 0
        if first_target != resolved_roles["target"]:
            errors.append(
                "The primary revision operation does not target the Teaching Week resolved "
                "from the Staff-visible week number."
            )
    return errors


def apply_reconciliation_output(
    draft: dict[str, Any],
    output: dict[str, Any],
    *,
    run_id: str,
    raw_task: str,
) -> dict[str, Any]:
    revised = copy.deepcopy(draft)
    revised.setdefault(
        "review_week_numbering",
        {
            "display_field": REVIEW_WEEK_FIELD,
            "display_label": REVIEW_WEEK_LABEL,
            "unqualified_week_reference": "display_field",
        },
    )
    by_id = {
        str(item["item_id"]): item
        for item in revised.get("content_items") or []
        if isinstance(item, dict)
    }
    for action in output["operations"]:
        action_type = str(action["type"])
        for item_id in action["item_ids"]:
            item = by_id[str(item_id)]
            if action_type == "delete_items":
                item["target_teaching_week"] = None
                item["status"] = "deleted"
            else:
                target_week = int(action["target_teaching_week"])
                item["target_teaching_week"] = target_week
                item["status"] = (
                    "scheduled"
                    if target_week == int(item.get("source_teaching_week") or 0)
                    else "moved"
                )
    revised["version"] = int(draft["version"]) + 1
    revised["last_review_run_id"] = run_id
    revised["revision_history"] = [
        *(draft.get("revision_history") or []),
        {
            "run_id": run_id,
            "from_version": draft["version"],
            "to_version": revised["version"],
            "staff_prompt": raw_task,
            "summary": str(output.get("summary") or ""),
            "operations": copy.deepcopy(output["operations"]),
            "recorded_at": datetime.now(timezone.utc).isoformat(),
        },
    ]
    status, overflow_count = _draft_status(
        revised["content_items"],
        int(revised["valid_teaching_week_count"]),
    )
    revised["status"] = status
    revised["overflow_count"] = overflow_count
    return revised


def execute_revise_timeline_draft_operation(
    operation: dict[str, Any],
    env: dict[str, str],
    run_id: str,
    raw_task: str,
) -> dict[str, Any]:
    errors = validate_revise_timeline_draft_shape(operation)
    if errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": REVISE_OPERATION_KIND,
            "run_id": run_id,
            "blocker": "; ".join(errors),
            "attachments": [],
        }
    draft, blocker = _active_draft_for_operation(
        operation,
        operation_kind=REVISE_OPERATION_KIND,
    )
    if not draft:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": REVISE_OPERATION_KIND,
            "run_id": run_id,
            "blocker": blocker,
            "attachments": [],
        }
    requester_blocker = _workflow_requester_blocker(draft, env)
    if requester_blocker:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": REVISE_OPERATION_KIND,
            "run_id": run_id,
            "blocker": requester_blocker,
            "attachments": [],
        }
    report_dir = DRAFT_ROOT / safe_name(str(draft["draft_id"])) / f"revision-run-{safe_name(run_id)}"
    output = run_timeline_reconciliation_agent(
        env,
        {
            "agent_task_path": str(report_dir / "timeline_reconciliation_task.md"),
            "agent_output_path": str(report_dir / "timeline_reconciliation_output.json"),
            "raw_task": raw_task,
            "draft": draft,
        },
    )
    validation_errors = validate_reconciliation_output(output, draft)
    if validation_errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": REVISE_OPERATION_KIND,
            "run_id": run_id,
            "draft_id": draft["draft_id"],
            "draft_version": draft["version"],
            "blocker": "; ".join(validation_errors),
            "attachments": [],
        }
    revised = apply_reconciliation_output(draft, output, run_id=run_id, raw_task=raw_task)
    draft_path, preview = save_draft_version(revised, env)
    _write_active_context(revised, draft_path, preview)
    attachments = [preview["pdf_path"]] if preview.get("pdf_path") else [preview["docx_path"]]
    result = {
        "ok": True,
        "status": revised["status"],
        "operation_kind": REVISE_OPERATION_KIND,
        "run_id": run_id,
        "draft_id": revised["draft_id"],
        "draft_version": revised["version"],
        "overflow_count": revised["overflow_count"],
        "preview_docx_path": preview["docx_path"],
        "preview_pdf_path": preview.get("pdf_path", ""),
        "attachments": attachments,
        "canonical_files_changed": False,
    }
    result_path = report_dir / "revision_result.json"
    write_json(result_path, result)
    result["result_path"] = repo_rel(result_path)
    return result


def _content_week_mappings(draft: dict[str, Any]) -> list[dict[str, Any]]:
    mappings: list[dict[str, Any]] = []
    for item in draft.get("content_items") or []:
        if not isinstance(item, dict):
            continue
        source_week = int(item.get("source_teaching_week") or 0)
        target_week = item.get("target_teaching_week")
        status = str(item.get("status") or "")
        mappings.append(
            {
                "item_id": item.get("item_id"),
                "content_kind": item.get("content_kind"),
                "content_text": item.get("text"),
                "source_teaching_week": source_week,
                "source_visible_week": item.get("source_visible_week"),
                "source_calendar_week": item.get("source_calendar_week"),
                "source_date": item.get("source_date"),
                "source_date_iso": item.get("source_date_iso"),
                "source_weekday": item.get("source_weekday"),
                "source_semester": item.get("source_semester"),
                "source_academic_year": item.get("source_academic_year"),
                "target_teaching_week": int(target_week) if isinstance(target_week, int) else None,
                "status": (
                    "deleted"
                    if status == "deleted"
                    else "unchanged"
                    if int(target_week) == source_week
                    else "moved"
                ),
            }
        )
    return mappings


def materialize_committed_payloads(draft: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    candidate_payloads = (
        draft.get("candidate_payloads")
        if isinstance(draft.get("candidate_payloads"), dict)
        else {}
    )
    source_timeline = copy.deepcopy(candidate_payloads.get("source_timeline.json") or {})
    calendar_payload = copy.deepcopy(candidate_payloads.get("calendar.json") or {})
    valid_slots = [
        slot
        for slot in draft.get("teaching_slots") or []
        if isinstance(slot, dict) and slot.get("valid") is True
    ]
    valid_slots = sorted(valid_slots, key=lambda slot: int(slot["teaching_week"]))
    valid_count = int(draft.get("valid_teaching_week_count") or 0)
    valid_weeks = [int(slot["teaching_week"]) for slot in valid_slots]
    if valid_weeks != list(range(1, valid_count + 1)):
        raise ValueError("Final timeline draft teaching weeks are not continuous from 1.")
    items = [
        item
        for item in draft.get("content_items") or []
        if isinstance(item, dict) and item.get("status") != "deleted"
    ]
    invalid_items = [
        str(item.get("item_id") or "")
        for item in items
        if item.get("target_teaching_week") not in set(valid_weeks)
    ]
    if invalid_items:
        raise ValueError(
            "Final timeline draft still contains content outside valid teaching weeks: "
            + ", ".join(invalid_items)
        )
    rows: list[dict[str, Any]] = []
    for slot in valid_slots:
        week = int(slot["teaching_week"])
        week_items = sorted(
            [item for item in items if item.get("target_teaching_week") == week],
            key=lambda item: (
                int(item.get("source_teaching_week") or 0),
                CONTENT_KINDS.index(str(item.get("content_kind"))),
                str(item.get("item_id") or ""),
            ),
        )
        by_kind = {
            kind: [item for item in week_items if item.get("content_kind") == kind]
            for kind in CONTENT_KINDS
        }
        source_visible_weeks = list(
            dict.fromkeys(
                int(item.get("source_visible_week") or item.get("source_teaching_week") or week)
                for item in week_items
            )
        )
        source_teaching_weeks = list(
            dict.fromkeys(int(item.get("source_teaching_week") or week) for item in week_items)
        )
        source_dates = list(
            dict.fromkeys(str(item.get("source_date") or "") for item in week_items if item.get("source_date"))
        )
        source_date_isos = list(
            dict.fromkeys(
                str(item.get("source_date_iso") or "")
                for item in week_items
                if item.get("source_date_iso")
            )
        )
        source_weekdays = list(
            dict.fromkeys(
                str(item.get("source_weekday") or "")
                for item in week_items
                if item.get("source_weekday")
            )
        )
        source_calendar_weeks = list(
            dict.fromkeys(
                int(item["source_calendar_week"])
                for item in week_items
                if isinstance(item.get("source_calendar_week"), int)
            )
        )
        dates = [str(value) for value in slot.get("teaching_dates") or []]
        rows.append(
            {
                "week": week,
                "teaching_week": week,
                "calendar_week": int(slot["calendar_week"]),
                "date": dates[0],
                "teaching_dates": dates,
                "allowed_session_weekdays": list(slot.get("allowed_session_weekdays") or []),
                "lecture": "\n".join(str(item["text"]) for item in by_kind["lecture"]),
                "workshop": "\n".join(str(item["text"]) for item in by_kind["workshop"]),
                "submission": "\n".join(str(item["text"]) for item in by_kind["submission"]),
                "source_week": source_visible_weeks[0] if source_visible_weeks else week,
                "source_weeks": source_visible_weeks,
                "source_teaching_weeks": source_teaching_weeks,
                "source_date": source_dates[0] if source_dates else "",
                "source_dates": source_dates,
                "source_date_iso": source_date_isos[0] if source_date_isos else "",
                "source_date_isos": source_date_isos,
                "source_weekday": source_weekdays[0] if len(source_weekdays) == 1 else "",
                "source_weekdays": source_weekdays,
                "source_calendar_week": (
                    source_calendar_weeks[0]
                    if len(source_calendar_weeks) == 1
                    else None
                ),
                "source_calendar_weeks": source_calendar_weeks,
                "counts_as_teaching_week": True,
                "content_item_ids": [str(item["item_id"]) for item in week_items],
            }
        )
    timeline = copy.deepcopy(source_timeline)
    timeline.update(
        {
            "status": "source-confirmed",
            "week_numbering": {
                "mode": "dual",
                "mapping_field": "teaching_week",
                "display_field": "date_planner_selected",
                "available_display_fields": ["calendar_week", "teaching_week"],
                "legacy_week_field": "teaching_week",
                "rule": (
                    "Mappings use actual teaching_week, with content_week_mappings overriding "
                    "same-week mapping only for explicitly moved or deleted content."
                ),
                "source_non_teaching_rows_excluded": True,
                "target_non_teaching_calendar_gaps_excluded": True,
            },
            "week_count": len(rows),
            "teaching_week_count": len(rows),
            "calendar_week_count": max(
                [int(slot.get("calendar_week") or 0) for slot in valid_slots]
                + [
                    int(gap.get("calendar_week") or 0)
                    for gap in draft.get("non_teaching_gaps") or []
                    if isinstance(gap, dict)
                ]
                + [0]
            ),
            "week_calendar": rows,
            "non_teaching_calendar_gaps": copy.deepcopy(draft.get("non_teaching_gaps") or []),
            "class_weekdays": list(draft.get("class_weekdays") or []),
            "content_week_mappings": _content_week_mappings(draft),
            "timeline_draft_resolution": {
                "draft_id": draft["draft_id"],
                "version": draft["version"],
                "committed_after_staff_confirmation": True,
            },
        }
    )
    schedule = copy.deepcopy(draft.get("candidate_schedule") or {})
    schedule["weeks"] = [
        {
            "week": int(slot["teaching_week"]),
            "teaching_week": int(slot["teaching_week"]),
            "calendar_week": int(slot["calendar_week"]),
            "label": f"Teaching Week {int(slot['teaching_week'])}",
            "teaching_dates": list(slot.get("teaching_dates") or []),
            "allowed_session_weekdays": list(slot.get("allowed_session_weekdays") or []),
            "notes": "Committed from the staff-reviewed timeline draft.",
        }
        for slot in valid_slots
    ]
    schedule["content_week_mappings"] = _content_week_mappings(draft)
    schedule["provenance"] = {
        "origin": "staff_confirmed_timeline_draft",
        "draft_id": draft["draft_id"],
        "draft_version": draft["version"],
        "merge_mode": "replace_all",
    }
    return schedule, timeline


def _replace_paths(value: Any, replacements: dict[str, str]) -> Any:
    if isinstance(value, dict):
        return {key: _replace_paths(item, replacements) for key, item in value.items()}
    if isinstance(value, list):
        return [_replace_paths(item, replacements) for item in value]
    if isinstance(value, str):
        return replacements.get(value, value)
    return value


def _copy_staged_sources(
    draft: dict[str, Any],
    target_root: Path,
) -> tuple[list[dict[str, Any]], list[Path], dict[str, str]]:
    try:
        from reference_pack_rebuild import copy_to_sources
    except ModuleNotFoundError:  # pragma: no cover
        from course_agent_app.reference_pack_rebuild import copy_to_sources

    committed: list[dict[str, Any]] = []
    created: list[Path] = []
    replacements: dict[str, str] = {}
    for record in draft.get("staged_sources") or []:
        if not isinstance(record, dict):
            continue
        source = resolve_local_path(str(record.get("path") or ""))
        if not source.is_file():
            raise FileNotFoundError(f"Draft source is no longer available: {source.name}")
        before = {
            path.resolve(strict=False)
            for path in (target_root / "sources").glob("*")
            if path.is_file()
        } if (target_root / "sources").exists() else set()
        target = copy_to_sources(source, target_root / "sources")
        if target.resolve(strict=False) not in before:
            created.append(target)
        committed_record = dict(record)
        replacements[str(record.get("path") or "")] = repo_rel(target)
        committed_record["path"] = repo_rel(target)
        committed.append(committed_record)
    return committed, created, replacements


def _manifest_for_commit(
    draft: dict[str, Any],
    temp_paths: dict[str, Path],
    committed_sources: list[dict[str, Any]],
    base_manifest: dict[str, Any],
) -> dict[str, Any]:
    manifest = copy.deepcopy(base_manifest)
    generated_references = [
        dict(item)
        for item in manifest.get("generated_references") or []
        if isinstance(item, dict)
        and Path(str(item.get("path") or "")).name not in temp_paths
    ]
    for filename, path in temp_paths.items():
        if filename == "source_manifest.json":
            continue
        generated_references.append(
            {
                "path": repo_rel(resolve_local_path(str(draft["reference_pack_path"])) / filename),
                "sha256": sha256(path),
            }
        )
    manifest.update(
        {
            "course_code": draft.get("course_code"),
            "semester": draft.get("semester"),
            "status": "ready",
            "generated_references": generated_references,
            "source_files": committed_sources or manifest.get("source_files") or [],
            "missing_inputs": [],
            "last_timeline_draft_commit": {
                "draft_id": draft["draft_id"],
                "version": draft["version"],
                "week_coordinate": "teaching_week",
                "content_week_mappings": True,
            },
        }
    )
    return manifest


def commit_draft_payloads(
    draft: dict[str, Any],
    *,
    run_id: str,
) -> dict[str, Any]:
    target_root = resolve_local_path(str(draft["reference_pack_path"]))
    target_root.mkdir(parents=True, exist_ok=True)
    schedule, timeline = materialize_committed_payloads(draft)
    candidate_payloads = (
        draft.get("candidate_payloads")
        if isinstance(draft.get("candidate_payloads"), dict)
        else {}
    )
    payloads: dict[str, dict[str, Any]] = {
        "timeline.json": timeline,
        "semester_schedule.json": schedule,
    }
    if draft.get("origin") == "rebuild_reference_pack":
        calendar = candidate_payloads.get("calendar.json")
        source_calendar = candidate_payloads.get("source_calendar.json")
        syllabus = candidate_payloads.get("syllabus_rules.json")
        if not isinstance(calendar, dict) or not isinstance(syllabus, dict):
            raise ValueError("Rebuild timeline draft is missing candidate calendar or syllabus rules.")
        payloads["calendar.json"] = copy.deepcopy(calendar)
        if isinstance(source_calendar, dict):
            payloads["source_calendar.json"] = copy.deepcopy(source_calendar)
        payloads["syllabus_rules.json"] = copy.deepcopy(syllabus)

    commit_dir = DRAFT_ROOT / safe_name(str(draft["draft_id"])) / f"commit-{safe_name(run_id)}"
    backup_dir = commit_dir / "previous_reference_json"
    temp_paths: dict[str, Path] = {}
    backup_paths: dict[str, str] = {}
    created_sources: list[Path] = []
    committed_sources: list[dict[str, Any]] = []
    path_replacements: dict[str, str] = {}
    try:
        if draft.get("origin") == "rebuild_reference_pack":
            committed_sources, created_sources, path_replacements = _copy_staged_sources(
                draft,
                target_root,
            )
            payloads = {
                name: _replace_paths(payload, path_replacements)
                for name, payload in payloads.items()
            }
        for filename, payload in payloads.items():
            temp_path = commit_dir / f"{filename}.new"
            write_json(temp_path, payload)
            temp_paths[filename] = temp_path
        base_manifest = candidate_payloads.get("source_manifest.json")
        if not isinstance(base_manifest, dict):
            base_manifest = {}
        manifest_payload = _manifest_for_commit(
            draft,
            temp_paths,
            committed_sources,
            base_manifest,
        )
        manifest_temp = commit_dir / "source_manifest.json.new"
        write_json(manifest_temp, manifest_payload)
        temp_paths["source_manifest.json"] = manifest_temp

        for filename in temp_paths:
            target = target_root / filename
            if target.exists():
                backup_dir.mkdir(parents=True, exist_ok=True)
                backup = backup_dir / filename
                shutil.copy2(target, backup)
                backup_paths[filename] = repo_rel(backup)
        replaced: list[str] = []
        try:
            for filename, temp_path in temp_paths.items():
                target = target_root / filename
                replace_temp = target_root / f".{filename}.{safe_name(run_id)}.tmp"
                shutil.copy2(temp_path, replace_temp)
                expected_checksum = sha256(temp_path)
                if sha256(replace_temp) != expected_checksum:
                    raise OSError(f"Checksum mismatch while staging {filename}.")
                os.replace(replace_temp, target)
                replaced.append(filename)
                if sha256(target) != expected_checksum:
                    raise OSError(f"Checksum mismatch after writing {filename}.")
        except Exception:
            for filename in replaced:
                target = target_root / filename
                backup_text = backup_paths.get(filename)
                if backup_text:
                    shutil.copy2(resolve_local_path(backup_text), target)
                else:
                    target.unlink(missing_ok=True)
            raise
    except Exception:
        for path in created_sources:
            path.unlink(missing_ok=True)
        raise
    return {
        "reference_pack_path": repo_rel(target_root),
        "generated_files": [repo_rel(target_root / name) for name in temp_paths],
        "backup_paths": backup_paths,
        "content_week_mapping_count": len(timeline.get("content_week_mappings") or []),
    }


def acquire_commit_lock(draft: dict[str, Any], run_id: str) -> Path:
    lock_path = DRAFT_ROOT / safe_name(str(draft["draft_id"])) / "commit.lock"
    lock_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        descriptor = os.open(lock_path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    except FileExistsError as exc:
        raise RuntimeError("Another commit is already in progress for this timeline draft.") from exc
    with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
        handle.write(
            json.dumps(
                {
                    "draft_id": draft["draft_id"],
                    "version": draft["version"],
                    "run_id": run_id,
                    "created_at": datetime.now(timezone.utc).isoformat(),
                }
            )
            + "\n"
        )
    return lock_path


def execute_commit_timeline_draft_operation(
    operation: dict[str, Any],
    env: dict[str, str],
    run_id: str,
    raw_task: str,
) -> dict[str, Any]:
    errors = validate_commit_timeline_draft_shape(operation)
    if errors:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "blocker": "; ".join(errors),
            "attachments": [],
        }
    draft, blocker = _active_draft_for_operation(
        operation,
        operation_kind=COMMIT_OPERATION_KIND,
    )
    if not draft:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "blocker": blocker,
            "attachments": [],
        }
    requester_blocker = _workflow_requester_blocker(draft, env)
    if requester_blocker:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "blocker": requester_blocker,
            "attachments": [],
        }
    if str(draft.get("status") or "") != "READY_FOR_CONFIRMATION":
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "blocker": "The latest timeline draft still has unresolved overflow content.",
            "attachments": [],
        }
    if str(draft.get("last_review_run_id") or "") == run_id:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "blocker": "Timeline revisions and final confirmation must be separate staff prompts.",
            "attachments": [],
        }
    status, overflow_count = _draft_status(
        draft.get("content_items") or [],
        int(draft.get("valid_teaching_week_count") or 0),
    )
    if status != "READY_FOR_CONFIRMATION" or overflow_count:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "blocker": "The timeline draft failed final overflow validation.",
            "attachments": [],
        }
    lock_path: Path | None = None
    try:
        lock_path = acquire_commit_lock(draft, run_id)
        locked_draft, locked_blocker = _active_draft_for_operation(
            operation,
            operation_kind=COMMIT_OPERATION_KIND,
        )
        if not locked_draft:
            raise RuntimeError(locked_blocker)
        committed = commit_draft_payloads(locked_draft, run_id=run_id)
        context = active_timeline_draft_context()
        write_json(
            ACTIVE_DRAFT_CONTEXT_PATH,
            {
                **{key: value for key, value in context.items() if key != "draft"},
                "status": "COMMITTED",
                "committed_run_id": run_id,
                "committed_at": datetime.now(timezone.utc).isoformat(),
            },
        )
    except Exception as exc:
        return {
            "ok": False,
            "status": "BLOCKED",
            "operation_kind": COMMIT_OPERATION_KIND,
            "run_id": run_id,
            "draft_id": draft["draft_id"],
            "draft_version": draft["version"],
            "blocker": f"Could not atomically commit the reviewed timeline draft: {exc}",
            "attachments": [],
        }
    finally:
        if lock_path is not None:
            lock_path.unlink(missing_ok=True)
    commit_result_path = (
        DRAFT_ROOT
        / safe_name(str(draft["draft_id"]))
        / f"commit-{safe_name(run_id)}"
        / "commit_result.json"
    )
    result = {
        "ok": True,
        "status": "COMMITTED",
        "operation_kind": COMMIT_OPERATION_KIND,
        "run_id": run_id,
        "raw_task": raw_task,
        "draft_id": draft["draft_id"],
        "draft_version": draft["version"],
        "week_count": draft["valid_teaching_week_count"],
        "attachments": [],
        **committed,
    }
    write_json(commit_result_path, result)
    result["result_path"] = repo_rel(commit_result_path)
    return result
